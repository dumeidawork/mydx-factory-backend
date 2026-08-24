"""仓库管理 API — 辅料闭环 + 成品骨架 + 缺口补齐（审批/附件/价格/对账/导出）。"""
from __future__ import annotations

import io
from datetime import date, datetime
from decimal import Decimal
from typing import Any, Optional
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from openpyxl import Workbook
from pydantic import BaseModel, Field
from typing_extensions import Annotated

from app.api.deps import CurrentUser, get_current_user
from app.core.database import execute, fetch_all, fetch_one, get_db
from app.services.warehouse_audit import (
    audit_tuple,
    diff_fields,
    gen_no,
    serialize_row,
    serialize_rows,
    update_audit_tuple,
    write_op_log,
)
from app.services.warehouse_gap import (
    CASH_REJECT_TO,
    attachment_abs_path,
    count_po_attachments,
    invoice_rows,
    list_po_attachments,
    record_cash_approval,
    record_price_and_alerts,
    refresh_named_alerts,
    require_cash_role,
    save_purchase_attachment,
    upsert_invoice,
)
from app.services.warehouse_ux import (
    add_dict_option,
    ensure_ux_schema,
    get_setting,
    list_dict_options,
    maybe_archive_old_receipts,
    receipts_archives_root,
    refresh_cash_warn,
    search_receipts,
    set_setting,
)

router = APIRouter(prefix="/warehouse", tags=["仓库管理"])
UserDep = Annotated[CurrentUser, Depends(get_current_user)]
_QTY_TOL = Decimal("0.001")
_POSTED = "COALESCE(status,'posted') <> 'voided'"

router = APIRouter(prefix="/warehouse", tags=["仓库管理"])
UserDep = Annotated[CurrentUser, Depends(get_current_user)]
_QTY_TOL = Decimal("0.001")


def _month_bounds(month: str) -> tuple[str, str, str]:
    if not month:
        month = date.today().strftime("%Y-%m")
    y, m = map(int, month.split("-"))
    start = f"{month}-01"
    end = f"{y + 1}-01-01" if m == 12 else f"{y}-{m + 1:02d}-01"
    return month, start, end


def _excel_response(wb: Workbook, file_name: str, user: CurrentUser, export_type: str, filters: dict | None = None):
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    write_op_log(
        domain="aux",
        action="export",
        entity_type=export_type,
        user=user,
        change_summary={"op": "export", "export_type": export_type, "filters": filters or {}, "file_name": file_name},
    )
    return StreamingResponse(
        bio,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(file_name)}"},
    )


def _d(v: Any, default: float = 0.0) -> Decimal:
    if v is None:
        return Decimal(str(default))
    return Decimal(str(v))


def _stock_status(qty: Decimal, qty_min: Decimal | None, qty_max: Decimal | None) -> str:
    if qty_min is not None and qty < qty_min:
        return "low"
    if qty_max is not None and qty > qty_max:
        return "high"
    return "normal"


def _refresh_item_alerts(item_id: int, user: CurrentUser) -> None:
    """按当前余额重建该辅料的低/高库存预警。"""
    item = fetch_one("SELECT * FROM wh_aux_items WHERE id=%s AND is_active=1", (item_id,))
    if not item:
        return
    bal = fetch_one(
        "SELECT COALESCE(SUM(qty_on_hand),0) AS qty FROM wh_aux_balances WHERE item_id=%s",
        (item_id,),
    )
    qty = _d(bal["qty"] if bal else 0)
    qty_min = item.get("qty_min")
    qty_max = item.get("qty_max")
    status = _stock_status(qty, _d(qty_min) if qty_min is not None else None, _d(qty_max) if qty_max is not None else None)

    execute(
        """
        UPDATE wh_alerts SET is_resolved=1, updated_by_user_id=%s, updated_by_name=%s
        WHERE domain='aux' AND ref_type='aux_item' AND ref_id=%s
          AND alert_type IN ('low_stock','high_stock') AND is_resolved=0
        """,
        (*update_audit_tuple(user), item_id),
    )
    if status == "normal":
        return
    title = (
        f"{item['name_spec']} 低于下限（{float(qty)} < {float(qty_min)}）"
        if status == "low"
        else f"{item['name_spec']} 高于上限（{float(qty)} > {float(qty_max)}）"
    )
    level = "error" if status == "low" else "warning"
    alert_type = "low_stock" if status == "low" else "high_stock"
    c_uid, c_name, u_uid, u_name = audit_tuple(user)
    execute(
        """
        INSERT INTO wh_alerts
          (domain, alert_type, title, level, ref_type, ref_id, is_resolved,
           created_by_user_id, created_by_name, updated_by_user_id, updated_by_name)
        VALUES ('aux', %s, %s, %s, 'aux_item', %s, 0, %s, %s, %s, %s)
        """,
        (alert_type, title, level, item_id, c_uid, c_name, u_uid, u_name),
    )


def _adjust_balance(
    *,
    conn,
    item_id: int,
    location_id: int | None,
    delta_qty: Decimal,
    unit_price: Decimal | None,
    user: CurrentUser,
) -> None:
    cur = conn.cursor(dictionary=True)
    loc = location_id
    cur.execute(
        "SELECT id, qty_on_hand, last_unit_price FROM wh_aux_balances WHERE item_id=%s AND "
        + ("location_id=%s" if loc is not None else "location_id IS NULL"),
        (item_id, loc) if loc is not None else (item_id,),
    )
    row = cur.fetchone()
    u_uid, u_name = update_audit_tuple(user)
    if row:
        new_qty = _d(row["qty_on_hand"]) + delta_qty
        if new_qty < 0:
            cur.close()
            raise HTTPException(status_code=400, detail="库存不足，无法出库")
        price = unit_price if unit_price is not None else row.get("last_unit_price")
        cur.execute(
            """
            UPDATE wh_aux_balances
            SET qty_on_hand=%s, last_unit_price=%s,
                updated_by_user_id=%s, updated_by_name=%s
            WHERE id=%s
            """,
            (new_qty, price, u_uid, u_name, row["id"]),
        )
    else:
        if delta_qty < 0:
            cur.close()
            raise HTTPException(status_code=400, detail="库存不足，无法出库")
        c_uid, c_name, u_uid2, u_name2 = audit_tuple(user)
        cur.execute(
            """
            INSERT INTO wh_aux_balances
              (item_id, location_id, qty_on_hand, last_unit_price,
               created_by_user_id, created_by_name, updated_by_user_id, updated_by_name)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
            """,
            (item_id, loc, delta_qty, unit_price, c_uid, c_name, u_uid2, u_name2),
        )
    cur.close()


def _ensure() -> None:
    ensure_ux_schema()
    maybe_archive_old_receipts()


def _is_posted(row: dict | None) -> bool:
    if not row:
        return False
    return str(row.get("status") or "posted") != "voided"


def _active_inbound_for_po(po_id: int, po_no: str | None) -> dict | None:
    return fetch_one(
        f"SELECT id, inbound_no, status FROM wh_inbound_orders WHERE (ref_po_id=%s OR ref_po_no=%s) AND COALESCE(status,'posted') <> 'voided' LIMIT 1",
        (po_id, po_no),
    )


def _slug_code(text: str, prefix: str) -> str:
    raw = "".join(ch if ch.isalnum() else "_" for ch in (text or "").strip())[:24].strip("_")
    return (raw or prefix)[:32]


# ---------- meta ----------
class CostCenterCreate(BaseModel):
    code: Optional[str] = None
    name: str
    sort_no: int = 0


class DictOptionCreate(BaseModel):
    kind: str
    value: str
    code: Optional[str] = None


@router.get("/meta/cost-centers")
def list_cost_centers(_user: UserDep):
    _ensure()
    rows = fetch_all(
        "SELECT id, code, name, sort_no FROM wh_cost_centers WHERE is_active=1 ORDER BY sort_no, id"
    )
    return {"items": serialize_rows(rows)}


@router.post("/meta/cost-centers")
def create_cost_center(req: CostCenterCreate, user: UserDep):
    _ensure()
    name = req.name.strip()
    if not name:
        raise HTTPException(status_code=400, detail="成本中心名称不能为空")
    code = (req.code or "").strip() or _slug_code(name, "CC")
    existed = fetch_one("SELECT * FROM wh_cost_centers WHERE code=%s OR name=%s", (code, name))
    if existed:
        return {"item": serialize_row(existed)}
    c_uid, c_name, u_uid, u_name = audit_tuple(user)
    cid = execute(
        """
        INSERT INTO wh_cost_centers
          (code, name, sort_no, created_by_user_id, created_by_name, updated_by_user_id, updated_by_name)
        VALUES (%s,%s,%s,%s,%s,%s,%s)
        """,
        (code, name, req.sort_no, c_uid, c_name, u_uid, u_name),
    )
    write_op_log(
        domain="aux",
        action="create",
        entity_type="cost_center",
        entity_id=cid,
        entity_no=code,
        user=user,
        change_summary={"op": "create", "name": name},
    )
    return {"item": serialize_row(fetch_one("SELECT * FROM wh_cost_centers WHERE id=%s", (cid,)))}


@router.get("/meta/options")
def meta_options(kind: str, _user: UserDep):
    _ensure()
    k = (kind or "").strip()
    if k in ("cost_center", "cost-center"):
        rows = fetch_all(
            "SELECT id, code, name FROM wh_cost_centers WHERE is_active=1 ORDER BY sort_no, id"
        )
        return {"items": [{"value": r["code"], "label": r["name"], "id": r["id"]} for r in rows]}
    if k in ("location", "shelf"):
        rows = fetch_all("SELECT id, code, name FROM wh_aux_locations ORDER BY sort_no, id")
        return {"items": [{"value": r["id"], "label": r["name"], "code": r["code"]} for r in rows]}
    if k in ("item", "aux"):
        rows = fetch_all(
            "SELECT id, name_spec, item_code FROM wh_aux_items WHERE is_active=1 ORDER BY item_code"
        )
        return {"items": [{"value": r["id"], "label": r["name_spec"], "code": r["item_code"]} for r in rows]}
    values = list_dict_options(k)
    return {"items": [{"value": v, "label": v} for v in values]}


@router.post("/meta/options")
def add_meta_option(req: DictOptionCreate, user: UserDep):
    _ensure()
    kind = req.kind.strip()
    value = req.value.strip()
    if not value:
        raise HTTPException(status_code=400, detail="选项不能为空")
    if kind in ("cost_center", "cost-center"):
        return create_cost_center(CostCenterCreate(code=req.code, name=value), user)
    if kind in ("location", "shelf"):
        code = (req.code or "").strip() or _slug_code(value, "LOC")
        existed = fetch_one("SELECT * FROM wh_aux_locations WHERE code=%s OR name=%s", (code, value))
        if existed:
            return {"item": serialize_row(existed)}
        return create_location(LocationCreate(code=code, name=value), user)
    if kind in ("item", "aux"):
        raise HTTPException(status_code=400, detail="辅料请到库存台账建档")
    val = add_dict_option(kind, value, user)
    return {"value": val, "item": {"value": val, "label": val}}


@router.get("/meta/uom-options")
def uom_options(_user: UserDep):
    return {"items": ["个", "条", "盒", "套", "米", "卷", "kg", "升", "桶", "双", "付", "件"]}


# ---------- locations ----------
class LocationCreate(BaseModel):
    code: str
    name: str
    sort_no: int = 0


@router.get("/aux/locations")
def list_locations(_user: UserDep):
    rows = fetch_all("SELECT * FROM wh_aux_locations ORDER BY sort_no, id")
    return {"items": serialize_rows(rows)}


@router.post("/aux/locations")
def create_location(req: LocationCreate, user: UserDep):
    if fetch_one("SELECT id FROM wh_aux_locations WHERE code=%s", (req.code.strip(),)):
        raise HTTPException(status_code=400, detail="货架编码已存在")
    c_uid, c_name, u_uid, u_name = audit_tuple(user)
    lid = execute(
        """
        INSERT INTO wh_aux_locations
          (code, name, sort_no, created_by_user_id, created_by_name, updated_by_user_id, updated_by_name)
        VALUES (%s, %s, %s, %s, %s, %s, %s)
        """,
        (req.code.strip(), req.name.strip(), req.sort_no, c_uid, c_name, u_uid, u_name),
    )
    write_op_log(
        domain="aux",
        action="create",
        entity_type="aux_location",
        entity_id=lid,
        entity_no=req.code,
        user=user,
        change_summary={"op": "create", "name": req.name},
    )
    row = fetch_one("SELECT * FROM wh_aux_locations WHERE id=%s", (lid,))
    return {"item": serialize_row(row)}


# ---------- items / balances ----------
class AuxItemCreate(BaseModel):
    item_code: str
    name_spec: str
    uom: str
    uom_desc: Optional[str] = None
    category: Optional[str] = None
    default_location_id: Optional[int] = None
    default_unit_price: Optional[float] = 0
    qty_min: Optional[float] = 0
    qty_max: Optional[float] = 0
    remark: Optional[str] = None
    qty_on_hand: Optional[float] = 0


class AuxItemUpdate(BaseModel):
    name_spec: Optional[str] = None
    uom: Optional[str] = None
    uom_desc: Optional[str] = None
    category: Optional[str] = None
    default_location_id: Optional[int] = None
    default_unit_price: Optional[float] = None
    qty_min: Optional[float] = None
    qty_max: Optional[float] = None
    remark: Optional[str] = None
    is_active: Optional[int] = None


@router.get("/aux/balances")
def list_balances(
    user: UserDep,
    keyword: str = "",
    location_id: Optional[int] = None,
    only_alert: bool = False,
):
    _ensure()
    sql = """
        SELECT i.id, i.item_code, i.name_spec, i.uom, i.uom_desc, i.category,
               i.qty_min, i.qty_max, i.default_unit_price, i.default_location_id, i.remark,
               i.created_by_name, i.updated_by_name, i.created_at, i.updated_at,
               COALESCE(SUM(b.qty_on_hand), 0) AS qty_on_hand,
               MAX(b.last_unit_price) AS last_unit_price,
               GROUP_CONCAT(DISTINCT loc.name ORDER BY loc.sort_no SEPARATOR ',') AS location_name
        FROM wh_aux_items i
        LEFT JOIN wh_aux_balances b ON b.item_id = i.id
        LEFT JOIN wh_aux_locations loc ON loc.id = b.location_id
        WHERE i.is_active = 1
    """
    params: list[Any] = []
    if keyword.strip():
        sql += " AND (i.item_code LIKE %s OR i.name_spec LIKE %s)"
        kw = f"%{keyword.strip()}%"
        params.extend([kw, kw])
    if location_id is not None:
        sql += " AND b.location_id = %s"
        params.append(location_id)
    sql += " GROUP BY i.id ORDER BY i.item_code"
    rows = fetch_all(sql, tuple(params))
    items = []
    for r in rows:
        qty = _d(r.get("qty_on_hand"))
        price = r.get("last_unit_price")
        if price is None:
            price = r.get("default_unit_price") or 0
        status = _stock_status(
            qty,
            _d(r["qty_min"]) if r.get("qty_min") is not None else None,
            _d(r["qty_max"]) if r.get("qty_max") is not None else None,
        )
        if only_alert and status == "normal":
            continue
        item = serialize_row(r) or {}
        item["qty_on_hand"] = float(qty)
        item["last_unit_price"] = float(_d(price))
        item["stock_amount"] = float(qty * _d(price))
        item["stock_status"] = status
        item["location_name"] = r.get("location_name") or ""
        items.append(item)
    return {"items": items}


@router.post("/aux/items")
def create_item(req: AuxItemCreate, user: UserDep):
    code = req.item_code.strip()
    if fetch_one("SELECT id FROM wh_aux_items WHERE item_code=%s", (code,)):
        raise HTTPException(status_code=400, detail="辅料编码已存在")
    c_uid, c_name, u_uid, u_name = audit_tuple(user)
    item_id = execute(
        """
        INSERT INTO wh_aux_items
          (item_code, name_spec, uom, uom_desc, category, default_location_id,
           default_unit_price, qty_min, qty_max, remark,
           created_by_user_id, created_by_name, updated_by_user_id, updated_by_name)
        VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
        """,
        (
            code,
            req.name_spec.strip(),
            req.uom.strip(),
            req.uom_desc,
            req.category,
            req.default_location_id,
            req.default_unit_price or 0,
            req.qty_min or 0,
            req.qty_max or 0,
            req.remark,
            c_uid,
            c_name,
            u_uid,
            u_name,
        ),
    )
    qty0 = _d(req.qty_on_hand or 0)
    if qty0 > 0:
        with get_db() as conn:
            _adjust_balance(
                conn=conn,
                item_id=item_id,
                location_id=req.default_location_id,
                delta_qty=qty0,
                unit_price=_d(req.default_unit_price or 0),
                user=user,
            )
            cur = conn.cursor()
            txn_no = gen_no("TXN")
            cur.execute(
                """
                INSERT INTO wh_aux_txns
                  (txn_no, txn_type, biz_date, item_id, location_id, qty, unit_price, amount,
                   operator_name, remark,
                   created_by_user_id, created_by_name, updated_by_user_id, updated_by_name)
                VALUES (%s,'adjust',%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
                """,
                (
                    txn_no,
                    date.today(),
                    item_id,
                    req.default_location_id,
                    qty0,
                    req.default_unit_price or 0,
                    float(qty0 * _d(req.default_unit_price or 0)),
                    user.name,
                    "期初入库",
                    c_uid,
                    c_name,
                    u_uid,
                    u_name,
                ),
            )
            cur.close()
    write_op_log(
        domain="aux",
        action="create",
        entity_type="aux_item",
        entity_id=item_id,
        entity_no=code,
        user=user,
        change_summary={"op": "create", "name_spec": req.name_spec, "uom": req.uom, "qty_on_hand": float(qty0)},
    )
    _refresh_item_alerts(item_id, user)
    if req.category:
        add_dict_option("category", req.category, user)
    return {"id": item_id}


@router.put("/aux/items/{item_id}")
def update_item(item_id: int, req: AuxItemUpdate, user: UserDep):
    _ensure()
    before = fetch_one("SELECT * FROM wh_aux_items WHERE id=%s", (item_id,))
    if not before:
        raise HTTPException(status_code=404, detail="辅料不存在")
    fields = req.model_dump(exclude_unset=True)
    if not fields:
        return {"status": "ok"}
    if fields.get("is_active") == 0:
        bal = fetch_one(
            "SELECT COALESCE(SUM(qty_on_hand),0) AS qty FROM wh_aux_balances WHERE item_id=%s",
            (item_id,),
        )
        if _d(bal["qty"] if bal else 0) > 0:
            raise HTTPException(status_code=400, detail="仍有库存，禁止停用")
    sets = []
    params: list[Any] = []
    for k, v in fields.items():
        sets.append(f"{k}=%s")
        params.append(v)
    u_uid, u_name = update_audit_tuple(user)
    sets.extend(["updated_by_user_id=%s", "updated_by_name=%s"])
    params.extend([u_uid, u_name, item_id])
    execute(f"UPDATE wh_aux_items SET {', '.join(sets)} WHERE id=%s", tuple(params))
    after = fetch_one("SELECT * FROM wh_aux_items WHERE id=%s", (item_id,)) or {}
    write_op_log(
        domain="aux",
        action="update",
        entity_type="aux_item",
        entity_id=item_id,
        entity_no=str(before.get("item_code") or ""),
        user=user,
        change_summary={"op": "update", "fields": diff_fields(before, after, list(fields.keys()))},
    )
    _refresh_item_alerts(item_id, user)
    if fields.get("category"):
        add_dict_option("category", str(fields["category"]), user)
    return {"item": serialize_row(after)}


# ---------- inbound ----------
class InboundLine(BaseModel):
    item_id: int
    location_id: Optional[int] = None
    qty: float
    unit_price: Optional[float] = 0
    uom: Optional[str] = None


class InboundCreate(BaseModel):
    biz_date: date
    source_type: str = "manual"  # purchase|manual
    supplier_name: Optional[str] = None
    ref_po_no: Optional[str] = None
    ref_po_id: Optional[int] = None
    operator_name: Optional[str] = None
    remark: Optional[str] = None
    lines: list[InboundLine] = Field(min_length=1)


@router.get("/aux/inbound")
def list_inbound(_user: UserDep, limit: int = Query(100, ge=1, le=500)):
    _ensure()
    rows = fetch_all(
        "SELECT * FROM wh_inbound_orders ORDER BY biz_date DESC, id DESC LIMIT %s",
        (limit,),
    )
    return {"items": serialize_rows(rows)}


@router.get("/aux/inbound/{inbound_id}")
def get_inbound(inbound_id: int, _user: UserDep):
    _ensure()
    row = fetch_one("SELECT * FROM wh_inbound_orders WHERE id=%s", (inbound_id,))
    if not row:
        raise HTTPException(status_code=404, detail="入库单不存在")
    lines = fetch_all(
        """
        SELECT l.*, i.item_code, i.name_spec
        FROM wh_inbound_order_lines l
        LEFT JOIN wh_aux_items i ON i.id = l.item_id
        WHERE l.inbound_id=%s ORDER BY l.line_no, l.id
        """,
        (inbound_id,),
    )
    return {"item": serialize_row(row), "lines": serialize_rows(lines)}


@router.post("/aux/inbound/{inbound_id}/void")
def void_inbound(inbound_id: int, user: UserDep):
    _ensure()
    row = fetch_one("SELECT * FROM wh_inbound_orders WHERE id=%s", (inbound_id,))
    if not row:
        raise HTTPException(status_code=404, detail="入库单不存在")
    if not _is_posted(row):
        raise HTTPException(status_code=400, detail="该入库单已作废")
    lines = fetch_all("SELECT * FROM wh_inbound_order_lines WHERE inbound_id=%s", (inbound_id,))
    c_uid, c_name, u_uid, u_name = audit_tuple(user)
    with get_db() as conn:
        for ln in lines:
            qty = _d(ln["qty"])
            price = _d(ln.get("unit_price") or 0)
            loc = ln.get("location_id")
            try:
                _adjust_balance(
                    conn=conn,
                    item_id=int(ln["item_id"]),
                    location_id=int(loc) if loc is not None else None,
                    delta_qty=-qty,
                    unit_price=price,
                    user=user,
                )
            except HTTPException as exc:
                if exc.status_code == 400:
                    raise HTTPException(status_code=400, detail="库存已领用，请先作废对应领用单后再作废入库") from exc
                raise
            cur = conn.cursor()
            txn_no = gen_no("TXN")
            cur.execute(
                """
                INSERT INTO wh_aux_txns
                  (txn_no, txn_type, biz_date, item_id, location_id, qty, unit_price, amount,
                   ref_type, ref_id, operator_name, remark,
                   created_by_user_id, created_by_name, updated_by_user_id, updated_by_name)
                VALUES (%s,'inbound',%s,%s,%s,%s,%s,%s,'inbound_order',%s,%s,%s,%s,%s,%s,%s)
                """,
                (
                    txn_no,
                    row["biz_date"],
                    ln["item_id"],
                    loc,
                    -qty,
                    price,
                    -(_d(ln.get("amount") or 0)),
                    inbound_id,
                    user.name,
                    f"作废冲回 {row.get('inbound_no')}",
                    c_uid,
                    c_name,
                    u_uid,
                    u_name,
                ),
            )
            cur.close()
        cur = conn.cursor()
        cur.execute(
            "UPDATE wh_inbound_orders SET status='voided', updated_by_user_id=%s, updated_by_name=%s WHERE id=%s",
            (u_uid, u_name, inbound_id),
        )
        cur.close()
    for ln in lines:
        _refresh_item_alerts(int(ln["item_id"]), user)
    write_op_log(
        domain="aux",
        action="void",
        entity_type="inbound_order",
        entity_id=inbound_id,
        entity_no=str(row.get("inbound_no") or ""),
        user=user,
        change_summary={"op": "void", "inbound_no": row.get("inbound_no")},
    )
    return {"item": serialize_row(fetch_one("SELECT * FROM wh_inbound_orders WHERE id=%s", (inbound_id,)))}


@router.post("/aux/inbound")
def create_inbound(req: InboundCreate, user: UserDep):
    _ensure()
    c_uid, c_name, u_uid, u_name = audit_tuple(user)
    inbound_no = gen_no("IN")
    total_qty = Decimal("0")
    total_amt = Decimal("0")
    prepared: list[tuple] = []
    for idx, line in enumerate(req.lines, start=1):
        item = fetch_one("SELECT * FROM wh_aux_items WHERE id=%s AND is_active=1", (line.item_id,))
        if not item:
            raise HTTPException(status_code=400, detail=f"辅料不存在: {line.item_id}")
        qty = _d(line.qty)
        if qty <= 0:
            raise HTTPException(status_code=400, detail="入库数量必须大于0")
        price = _d(line.unit_price or item.get("default_unit_price") or 0)
        amt = qty * price
        total_qty += qty
        total_amt += amt
        loc = line.location_id or item.get("default_location_id")
        prepared.append((item, loc, qty, price, amt, line.uom or item["uom"], idx))

    with get_db() as conn:
        cur = conn.cursor()
        cur.execute(
            """
            INSERT INTO wh_inbound_orders
              (inbound_no, biz_date, source_type, supplier_name, ref_po_no, ref_po_id,
               total_qty, total_amount, operator_name, remark, status,
               created_by_user_id, created_by_name, updated_by_user_id, updated_by_name)
            VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,'posted',%s,%s,%s,%s)
            """,
            (
                inbound_no,
                req.biz_date,
                req.source_type,
                req.supplier_name,
                req.ref_po_no,
                req.ref_po_id,
                total_qty,
                total_amt,
                req.operator_name or user.name,
                req.remark,
                c_uid,
                c_name,
                u_uid,
                u_name,
            ),
        )
        inbound_id = int(cur.lastrowid)
        for item, loc, qty, price, amt, uom, idx in prepared:
            cur.execute(
                """
                INSERT INTO wh_inbound_order_lines
                  (inbound_id, item_id, location_id, uom, qty, unit_price, amount, line_no,
                   created_by_user_id, created_by_name, updated_by_user_id, updated_by_name)
                VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
                """,
                (inbound_id, item["id"], loc, uom, qty, price, amt, idx, c_uid, c_name, u_uid, u_name),
            )
            _adjust_balance(conn=conn, item_id=int(item["id"]), location_id=loc, delta_qty=qty, unit_price=price, user=user)
            txn_no = gen_no("TXN")
            cur.execute(
                """
                INSERT INTO wh_aux_txns
                  (txn_no, txn_type, biz_date, item_id, location_id, qty, unit_price, amount,
                   ref_type, ref_id, operator_name,
                   created_by_user_id, created_by_name, updated_by_user_id, updated_by_name)
                VALUES (%s,'inbound',%s,%s,%s,%s,%s,%s,'inbound_order',%s,%s,%s,%s,%s,%s)
                """,
                (
                    txn_no,
                    req.biz_date,
                    item["id"],
                    loc,
                    qty,
                    price,
                    amt,
                    inbound_id,
                    req.operator_name or user.name,
                    c_uid,
                    c_name,
                    u_uid,
                    u_name,
                ),
            )
        cur.close()

    for item, *_rest in prepared:
        _refresh_item_alerts(int(item["id"]), user)

    write_op_log(
        domain="aux",
        action="inbound",
        entity_type="inbound_order",
        entity_id=inbound_id,
        entity_no=inbound_no,
        user=user,
        change_summary={"op": "inbound", "total_qty": float(total_qty), "total_amount": float(total_amt)},
    )
    row = fetch_one("SELECT * FROM wh_inbound_orders WHERE id=%s", (inbound_id,))
    return {"item": serialize_row(row)}


# ---------- issue ----------
class IssueLine(BaseModel):
    item_id: int
    location_id: Optional[int] = None
    qty: float
    unit_price: Optional[float] = None
    consumable_type: Optional[str] = None


class IssueCreate(BaseModel):
    biz_date: date
    cost_center_code: str
    receiver_name: str
    keeper_name: Optional[str] = None
    issue_mode: str = "normal"
    remark: Optional[str] = None
    lines: list[IssueLine] = Field(min_length=1)


@router.get("/aux/issues")
def list_issues(
    _user: UserDep,
    cost_center: str = "",
    date_from: str = "",
    date_to: str = "",
    keyword: str = "",
    limit: int = Query(100, ge=1, le=500),
):
    _ensure()
    sql = "SELECT * FROM wh_issue_orders WHERE 1=1"
    params: list[Any] = []
    if cost_center.strip():
        sql += " AND (cost_center_code=%s OR cost_center_name=%s)"
        params.extend([cost_center.strip(), cost_center.strip()])
    if date_from.strip():
        sql += " AND biz_date>=%s"
        params.append(date_from.strip())
    if date_to.strip():
        sql += " AND biz_date<=%s"
        params.append(date_to.strip())
    if keyword.strip():
        sql += " AND (receiver_name LIKE %s OR summary LIKE %s OR issue_no LIKE %s)"
        kw = f"%{keyword.strip()}%"
        params.extend([kw, kw, kw])
    sql += " ORDER BY biz_date DESC, id DESC LIMIT %s"
    params.append(limit)
    return {"items": serialize_rows(fetch_all(sql, tuple(params)))}


@router.get("/aux/issues/{issue_id}")
def get_issue(issue_id: int, _user: UserDep):
    _ensure()
    row = fetch_one("SELECT * FROM wh_issue_orders WHERE id=%s", (issue_id,))
    if not row:
        raise HTTPException(status_code=404, detail="领用单不存在")
    lines = fetch_all(
        """
        SELECT l.*, i.item_code, i.name_spec
        FROM wh_issue_order_lines l
        LEFT JOIN wh_aux_items i ON i.id = l.item_id
        WHERE l.issue_id=%s ORDER BY l.id
        """,
        (issue_id,),
    )
    return {"item": serialize_row(row), "lines": serialize_rows(lines)}


@router.post("/aux/issues/{issue_id}/void")
def void_issue(issue_id: int, user: UserDep):
    _ensure()
    row = fetch_one("SELECT * FROM wh_issue_orders WHERE id=%s", (issue_id,))
    if not row:
        raise HTTPException(status_code=404, detail="领用单不存在")
    if not _is_posted(row):
        raise HTTPException(status_code=400, detail="该领用单已作废")
    lines = fetch_all("SELECT * FROM wh_issue_order_lines WHERE issue_id=%s", (issue_id,))
    c_uid, c_name, u_uid, u_name = audit_tuple(user)
    with get_db() as conn:
        for ln in lines:
            qty = _d(ln["qty"])
            price = _d(ln.get("unit_price") or 0)
            loc = ln.get("location_id")
            _adjust_balance(
                conn=conn,
                item_id=int(ln["item_id"]),
                location_id=int(loc) if loc is not None else None,
                delta_qty=qty,
                unit_price=price,
                user=user,
            )
            cur = conn.cursor()
            txn_no = gen_no("TXN")
            cur.execute(
                """
                INSERT INTO wh_aux_txns
                  (txn_no, txn_type, biz_date, item_id, location_id, qty, unit_price, amount,
                   cost_center_code, ref_type, ref_id, operator_name, remark,
                   created_by_user_id, created_by_name, updated_by_user_id, updated_by_name)
                VALUES (%s,'issue',%s,%s,%s,%s,%s,%s,%s,'issue_order',%s,%s,%s,%s,%s,%s,%s)
                """,
                (
                    txn_no,
                    row["biz_date"],
                    ln["item_id"],
                    loc,
                    qty,
                    price,
                    _d(ln.get("amount") or 0),
                    row.get("cost_center_code"),
                    issue_id,
                    user.name,
                    f"作废冲回 {row.get('issue_no')}",
                    c_uid,
                    c_name,
                    u_uid,
                    u_name,
                ),
            )
            cur.close()
        cur = conn.cursor()
        cur.execute(
            "UPDATE wh_issue_orders SET status='voided', updated_by_user_id=%s, updated_by_name=%s WHERE id=%s",
            (u_uid, u_name, issue_id),
        )
        cur.close()
    for ln in lines:
        _refresh_item_alerts(int(ln["item_id"]), user)
    write_op_log(
        domain="aux",
        action="void",
        entity_type="issue_order",
        entity_id=issue_id,
        entity_no=str(row.get("issue_no") or ""),
        user=user,
        change_summary={"op": "void", "issue_no": row.get("issue_no")},
    )
    return {"item": serialize_row(fetch_one("SELECT * FROM wh_issue_orders WHERE id=%s", (issue_id,)))}


@router.post("/aux/issues")
def create_issue(req: IssueCreate, user: UserDep):
    _ensure()
    cc = fetch_one(
        "SELECT code, name FROM wh_cost_centers WHERE code=%s OR name=%s",
        (req.cost_center_code, req.cost_center_code),
    )
    if not cc:
        created = create_cost_center(CostCenterCreate(name=req.cost_center_code), user)
        cc = created.get("item") or fetch_one(
            "SELECT code, name FROM wh_cost_centers WHERE code=%s OR name=%s",
            (req.cost_center_code, req.cost_center_code),
        )
    if not cc:
        raise HTTPException(status_code=400, detail="成本中心不存在")
    c_uid, c_name, u_uid, u_name = audit_tuple(user)
    issue_no = gen_no("IS")
    total_qty = Decimal("0")
    total_amt = Decimal("0")
    prepared: list[tuple] = []
    summaries: list[str] = []
    for line in req.lines:
        item = fetch_one("SELECT * FROM wh_aux_items WHERE id=%s AND is_active=1", (line.item_id,))
        if not item:
            raise HTTPException(status_code=400, detail=f"辅料不存在: {line.item_id}")
        qty = _d(line.qty)
        if qty <= 0:
            raise HTTPException(status_code=400, detail="领用数量必须大于0")
        bal = fetch_one(
            "SELECT COALESCE(SUM(qty_on_hand),0) AS qty, MAX(last_unit_price) AS price FROM wh_aux_balances WHERE item_id=%s",
            (line.item_id,),
        )
        price = _d(line.unit_price if line.unit_price is not None else (bal.get("price") if bal else None) or item.get("default_unit_price") or 0)
        amt = qty * price
        total_qty += qty
        total_amt += amt
        loc = line.location_id or item.get("default_location_id")
        summaries.append(f"{item['name_spec']}×{float(qty)}")
        prepared.append((item, loc, qty, price, amt, line.consumable_type))

    with get_db() as conn:
        cur = conn.cursor()
        cur.execute(
            """
            INSERT INTO wh_issue_orders
              (issue_no, biz_date, cost_center_code, cost_center_name, receiver_name, keeper_name,
               issue_mode, total_qty, total_amount, summary, status,
               created_by_user_id, created_by_name, updated_by_user_id, updated_by_name)
            VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,'posted',%s,%s,%s,%s)
            """,
            (
                issue_no,
                req.biz_date,
                cc["code"],
                cc["name"],
                req.receiver_name,
                req.keeper_name or user.name,
                req.issue_mode,
                total_qty,
                total_amt,
                "、".join(summaries)[:240],
                c_uid,
                c_name,
                u_uid,
                u_name,
            ),
        )
        issue_id = int(cur.lastrowid)
        for item, loc, qty, price, amt, consumable_type in prepared:
            # 若未指定库位，从有库存的库位扣减（优先默认库位）
            use_loc = loc
            if use_loc is None:
                cur2 = conn.cursor(dictionary=True)
                cur2.execute(
                    """
                    SELECT location_id FROM wh_aux_balances
                    WHERE item_id=%s AND qty_on_hand>0
                    ORDER BY qty_on_hand DESC LIMIT 1
                    """,
                    (item["id"],),
                )
                br = cur2.fetchone()
                cur2.close()
                use_loc = br["location_id"] if br else None
            cur.execute(
                """
                INSERT INTO wh_issue_order_lines
                  (issue_id, item_id, location_id, qty, uom, unit_price, amount, consumable_type,
                   created_by_user_id, created_by_name, updated_by_user_id, updated_by_name)
                VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
                """,
                (
                    issue_id,
                    item["id"],
                    use_loc,
                    qty,
                    item["uom"],
                    price,
                    amt,
                    consumable_type,
                    c_uid,
                    c_name,
                    u_uid,
                    u_name,
                ),
            )
            _adjust_balance(conn=conn, item_id=int(item["id"]), location_id=use_loc, delta_qty=-qty, unit_price=price, user=user)
            txn_no = gen_no("TXN")
            cur.execute(
                """
                INSERT INTO wh_aux_txns
                  (txn_no, txn_type, biz_date, item_id, location_id, qty, unit_price, amount,
                   cost_center_code, ref_type, ref_id, operator_name, keeper_name,
                   created_by_user_id, created_by_name, updated_by_user_id, updated_by_name)
                VALUES (%s,'issue',%s,%s,%s,%s,%s,%s,%s,'issue_order',%s,%s,%s,%s,%s,%s,%s)
                """,
                (
                    txn_no,
                    req.biz_date,
                    item["id"],
                    use_loc,
                    -qty,
                    price,
                    amt,
                    cc["code"],
                    issue_id,
                    req.receiver_name,
                    req.keeper_name or user.name,
                    c_uid,
                    c_name,
                    u_uid,
                    u_name,
                ),
            )
        cur.close()

    for item, *_rest in prepared:
        _refresh_item_alerts(int(item["id"]), user)

    write_op_log(
        domain="aux",
        action="issue",
        entity_type="issue_order",
        entity_id=issue_id,
        entity_no=issue_no,
        user=user,
        change_summary={
            "op": "issue",
            "cost_center": cc["name"],
            "receiver": req.receiver_name,
            "total_amount": float(total_amt),
        },
    )
    row = fetch_one("SELECT * FROM wh_issue_orders WHERE id=%s", (issue_id,))
    return {"item": serialize_row(row)}


@router.get("/aux/exports/issues")
def export_issues(
    user: UserDep,
    cost_center: str = "",
    date_from: str = "",
    date_to: str = "",
    keyword: str = "",
):
    data = list_issues(user, cost_center=cost_center, date_from=date_from, date_to=date_to, keyword=keyword, limit=500)
    wb = Workbook()
    ws = wb.active
    ws.title = "领用表"
    ws.append(["单号", "日期", "成本中心", "领用人", "物品摘要", "数量", "金额", "保管员", "类型", "录入人"])
    for r in data["items"]:
        ws.append(
            [
                r.get("issue_no"),
                r.get("biz_date"),
                r.get("cost_center_name") or r.get("cost_center_code"),
                r.get("receiver_name"),
                r.get("summary"),
                r.get("total_qty"),
                r.get("total_amount"),
                r.get("keeper_name"),
                r.get("issue_mode"),
                r.get("created_by_name"),
            ]
        )
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    file_name = f"领用表_{cost_center or '全部'}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    write_op_log(
        domain="aux",
        action="export",
        entity_type="issue_order",
        user=user,
        change_summary={
            "op": "export",
            "export_type": "issue_my06",
            "filters": {"cost_center": cost_center, "date_from": date_from, "date_to": date_to, "keyword": keyword},
            "file_name": file_name,
        },
    )
    return StreamingResponse(
        bio,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(file_name)}"},
    )


@router.get("/aux/exports/balances")
def export_balances(user: UserDep):
    data = list_balances(user)
    wb = Workbook()
    ws = wb.active
    ws.title = "库存台账"
    ws.append(["编码", "品名规格", "单位", "货架", "品类", "库存", "下限", "上限", "单价", "金额", "状态", "录入人"])
    for r in data["items"]:
        ws.append(
            [
                r.get("item_code"),
                r.get("name_spec"),
                r.get("uom"),
                r.get("location_name"),
                r.get("category"),
                r.get("qty_on_hand"),
                r.get("qty_min"),
                r.get("qty_max"),
                r.get("last_unit_price"),
                r.get("stock_amount"),
                r.get("stock_status"),
                r.get("created_by_name"),
            ]
        )
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    file_name = f"辅料库存_{datetime.now().strftime('%Y%m%d')}.xlsx"
    write_op_log(
        domain="aux",
        action="export",
        entity_type="aux_item",
        user=user,
        change_summary={"op": "export", "export_type": "balances", "file_name": file_name},
    )
    return StreamingResponse(
        bio,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(file_name)}"},
    )


@router.get("/aux/exports/reconcile")
def export_reconcile(user: UserDep, month: str = ""):
    month, start, end = _month_bounds(month)
    lines = fetch_all(
        """
        SELECT p.biz_date, l.name_spec, l.qty, l.unit_price, l.amount, p.purchaser_name, p.supplier_name, p.po_no
        FROM wh_purchase_order_lines l
        JOIN wh_purchase_orders p ON p.id = l.po_id
        WHERE p.pay_type='monthly' AND p.biz_date >= %s AND p.biz_date < %s
          AND p.status NOT IN ('cancelled','rejected','draft','voided')
        ORDER BY p.supplier_name, p.biz_date, p.id, l.line_no
        """,
        (start, end),
    )
    wb = Workbook()
    ws = wb.active
    ws.title = "对账单"
    ws.append(["日期", "物品及规格", "数量", "单价/元", "总价/元", "采购人", "供应商", "采购单号"])
    for r in lines:
        ws.append(
            [
                r.get("biz_date"),
                r.get("name_spec"),
                r.get("qty"),
                r.get("unit_price"),
                r.get("amount"),
                r.get("purchaser_name"),
                r.get("supplier_name"),
                r.get("po_no"),
            ]
        )
    file_name = f"供应商对账单_{month}.xlsx"
    return _excel_response(wb, file_name, user, "reconcile_gpxj11", {"month": month})


@router.get("/aux/exports/cash-ledger")
def export_cash_ledger(user: UserDep):
    rows = fetch_all("SELECT * FROM wh_cash_ledger ORDER BY biz_date, id")
    wb = Workbook()
    ws = wb.active
    ws.title = "现金明细"
    ws.append(["日期", "转款/元", "上次对账余/元", "合计", "花费/元", "下次对账剩/元", "备注", "录入人"])
    for r in rows:
        transfer = float(r.get("transfer_in") or 0)
        prev = float(r.get("prev_balance") or 0)
        ws.append(
            [
                r.get("biz_date"),
                transfer,
                prev,
                transfer + prev,
                r.get("spent"),
                r.get("next_balance"),
                r.get("remark"),
                r.get("created_by_name"),
            ]
        )
    file_name = f"现金明细_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return _excel_response(wb, file_name, user, "cash_ledger")


@router.get("/aux/exports/monthly-consume")
def export_monthly_consume(user: UserDep, month: str = ""):
    data = monthly_summary(user, month=month)
    wb = Workbook()
    ws = wb.active
    ws.title = "分类消耗"
    ws.append(["成本中心", "金额(元)"])
    for r in data.get("consume") or []:
        ws.append([r.get("category"), r.get("amount")])
    file_name = f"每月总消耗_{data.get('month')}.xlsx"
    return _excel_response(wb, file_name, user, "monthly_consume", {"month": data.get("month")})


@router.get("/aux/exports/alerts")
def export_alerts(user: UserDep):
    data = list_alerts(user, include_resolved=False)
    wb = Workbook()
    ws = wb.active
    ws.title = "预警清单"
    ws.append(["级别", "类型", "内容", "时间"])
    for r in data.get("items") or []:
        ws.append([r.get("level"), r.get("alert_type"), r.get("title"), r.get("created_at")])
    file_name = f"预警异常清单_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return _excel_response(wb, file_name, user, "alerts")
@router.get("/aux/alerts")
def list_alerts(_user: UserDep, include_resolved: bool = False):
    _ensure()
    if include_resolved:
        rows = fetch_all("SELECT * FROM wh_alerts WHERE domain='aux' ORDER BY id DESC LIMIT 200")
    else:
        rows = fetch_all(
            "SELECT * FROM wh_alerts WHERE domain='aux' AND is_resolved=0 ORDER BY id DESC LIMIT 200"
        )
    return {"items": serialize_rows(rows)}


@router.get("/aux/monthly-summary")
def monthly_summary(user: UserDep, month: str = ""):
    """month: YYYY-MM，默认当月。"""
    _ensure()
    month, start, end = _month_bounds(month)

    purchase = fetch_one(
        """
        SELECT COALESCE(SUM(total_amount),0) AS amt, COALESCE(SUM(total_qty),0) AS qty
        FROM wh_purchase_orders
        WHERE biz_date >= %s AND biz_date < %s AND status NOT IN ('cancelled','rejected','draft','voided')
        """,
        (start, end),
    )
    inbound = fetch_one(
        """
        SELECT COALESCE(SUM(total_amount),0) AS amt, COALESCE(SUM(total_qty),0) AS qty
        FROM wh_inbound_orders WHERE biz_date >= %s AND biz_date < %s AND COALESCE(status,'posted') <> 'voided'
        """,
        (start, end),
    )
    issue = fetch_one(
        """
        SELECT COALESCE(SUM(total_amount),0) AS amt, COALESCE(SUM(total_qty),0) AS qty
        FROM wh_issue_orders WHERE biz_date >= %s AND biz_date < %s AND COALESCE(status,'posted') <> 'voided'
        """,
        (start, end),
    )
    stock = fetch_one(
        """
        SELECT COALESCE(SUM(b.qty_on_hand * COALESCE(b.last_unit_price, i.default_unit_price, 0)),0) AS amt,
               COALESCE(SUM(b.qty_on_hand),0) AS qty
        FROM wh_aux_balances b
        JOIN wh_aux_items i ON i.id = b.item_id
        WHERE i.is_active=1
        """
    )
    consume_by_cc = fetch_all(
        """
        SELECT COALESCE(cost_center_name, cost_center_code) AS category,
               COALESCE(SUM(total_amount),0) AS amount
        FROM wh_issue_orders
        WHERE biz_date >= %s AND biz_date < %s AND COALESCE(status,'posted') <> 'voided'
        GROUP BY COALESCE(cost_center_name, cost_center_code)
        ORDER BY amount DESC
        """,
        (start, end),
    )
    triple = fetch_all(
        """
        SELECT
          COALESCE(p.supplier_name, i.supplier_name, '未指定') AS `key`,
          COALESCE(p.purchase_qty,0) AS purchase_qty,
          COALESCE(i.inbound_qty,0) AS inbound_qty,
          COALESCE(p.purchase_amt,0) AS purchase_amt,
          COALESCE(i.inbound_amt,0) AS inbound_amt
        FROM (
          SELECT supplier_name, SUM(total_qty) purchase_qty, SUM(total_amount) purchase_amt
          FROM wh_purchase_orders
          WHERE biz_date >= %s AND biz_date < %s AND pay_type='monthly'
            AND status NOT IN ('cancelled','rejected','draft','voided')
          GROUP BY supplier_name
        ) p
        LEFT JOIN (
          SELECT supplier_name, SUM(total_qty) inbound_qty, SUM(total_amount) inbound_amt
          FROM wh_inbound_orders
          WHERE biz_date >= %s AND biz_date < %s AND COALESCE(status,'posted') <> 'voided'
          GROUP BY supplier_name
        ) i ON COALESCE(p.supplier_name,'') = COALESCE(i.supplier_name,'')
        """,
        (start, end, start, end),
    )
    triple_out = []
    recon_titles: list[str] = []
    for r in serialize_rows(triple):
        ok = abs(float(r.get("purchase_amt") or 0) - float(r.get("inbound_amt") or 0)) < 0.01
        r["issue_qty"] = 0
        r["issue_amt"] = 0
        r["ok"] = ok
        triple_out.append(r)
        if not ok:
            diff = float(r.get("purchase_amt") or 0) - float(r.get("inbound_amt") or 0)
            recon_titles.append(f"{month} {r.get('key')} 采购入库金额差 {diff:.2f} 元")

    end_qty = _d(stock["qty"] if stock else 0)
    in_qty = _d(inbound["qty"] if inbound else 0)
    out_qty = _d(issue["qty"] if issue else 0)
    month_net = fetch_one(
        """
        SELECT COALESCE(SUM(qty),0) AS qty FROM wh_aux_txns
        WHERE biz_date >= %s AND biz_date < %s
        """,
        (start, end),
    )
    begin_qty = end_qty - _d(month_net["qty"] if month_net else 0)
    identity_left = in_qty - out_qty
    identity_right = end_qty - begin_qty
    identity_ok = abs(identity_left - identity_right) < _QTY_TOL
    identity_row = {
        "key": "全仓数量恒等式",
        "purchase_qty": 0,
        "inbound_qty": float(in_qty),
        "issue_qty": float(out_qty),
        "purchase_amt": 0,
        "inbound_amt": float(_d(inbound["amt"] if inbound else 0)),
        "issue_amt": float(_d(issue["amt"] if issue else 0)),
        "ok": identity_ok,
        "begin_qty": float(begin_qty),
        "end_qty": float(end_qty),
        "delta": float(identity_left - identity_right),
    }
    triple_out.insert(0, identity_row)

    refresh_named_alerts(
        alert_type="reconcile_diff",
        title_prefix=f"{month} ",
        titles=recon_titles,
        level="warning",
        user=user,
    )
    ident_titles = [] if identity_ok else [
        f"{month} 全仓数量恒等式异常：入库{float(in_qty)}-领用{float(out_qty)} vs 期末{float(end_qty)}-期初{float(begin_qty)}"
    ]
    refresh_named_alerts(
        alert_type="qty_identity",
        title_prefix=f"{month} 全仓数量恒等式",
        titles=ident_titles,
        level="warning",
        user=user,
    )

    alerts = list_alerts(user, include_resolved=False)["items"]
    return {
        "month": month,
        "purchase_total": float(_d(purchase["amt"] if purchase else 0)),
        "payable_total": float(_d(purchase["amt"] if purchase else 0)),
        "stock_value": float(_d(stock["amt"] if stock else 0)),
        "inbound_total": float(_d(inbound["amt"] if inbound else 0)),
        "issue_total": float(_d(issue["amt"] if issue else 0)),
        "begin_qty": float(begin_qty),
        "end_qty": float(end_qty),
        "identity_ok": identity_ok,
        "consume": serialize_rows(consume_by_cc),
        "triple_check": triple_out,
        "alerts": alerts,
    }


# ---------- purchases (basic) ----------
class PurchaseLine(BaseModel):
    item_id: Optional[int] = None
    name_spec: str
    uom: str = "个"
    qty: float
    unit_price: float


class PurchaseCreate(BaseModel):
    pay_type: str = "monthly"  # monthly|cash
    biz_date: date
    supplier_name: Optional[str] = None
    purchaser_name: Optional[str] = None
    remark: Optional[str] = None
    lines: list[PurchaseLine] = Field(min_length=1)


@router.get("/aux/purchases")
def list_purchases(_user: UserDep, pay_type: str = ""):
    _ensure()
    if pay_type:
        rows = fetch_all(
            "SELECT * FROM wh_purchase_orders WHERE pay_type=%s ORDER BY biz_date DESC, id DESC LIMIT 200",
            (pay_type,),
        )
    else:
        rows = fetch_all("SELECT * FROM wh_purchase_orders ORDER BY biz_date DESC, id DESC LIMIT 200")
    items = serialize_rows(rows)
    for it in items:
        pid = it.get("id")
        lines = fetch_all("SELECT item_id FROM wh_purchase_order_lines WHERE po_id=%s", (pid,))
        it["line_count"] = len(lines)
        it["linked_item_count"] = sum(1 for x in lines if x.get("item_id"))
        inbound = _active_inbound_for_po(int(pid), it.get("po_no"))
        it["already_inbound"] = bool(inbound)
        it["inbound_id"] = inbound.get("id") if inbound else None
        it["inbound_no"] = inbound.get("inbound_no") if inbound else None
        it["attachment_count"] = count_po_attachments(int(pid))
        it["can_inbound"] = (
            it["linked_item_count"] > 0
            and it["linked_item_count"] == it["line_count"]
            and not it["already_inbound"]
            and it.get("status") not in ("voided", "cancelled", "rejected", "draft")
            and (
                (it.get("pay_type") == "monthly" and it.get("status") == "confirmed")
                or (it.get("pay_type") == "cash" and it.get("status") in ("paid", "posted"))
            )
        )
    return {"items": items}


@router.get("/aux/purchases/{po_id}")
def get_purchase(po_id: int, _user: UserDep):
    _ensure()
    row = fetch_one("SELECT * FROM wh_purchase_orders WHERE id=%s", (po_id,))
    if not row:
        raise HTTPException(status_code=404, detail="采购单不存在")
    lines = fetch_all("SELECT * FROM wh_purchase_order_lines WHERE po_id=%s ORDER BY line_no, id", (po_id,))
    atts = serialize_rows(list_po_attachments(po_id))
    item = serialize_row(row) or {}
    inbound = _active_inbound_for_po(po_id, row.get("po_no"))
    item["already_inbound"] = bool(inbound)
    item["inbound_id"] = inbound.get("id") if inbound else None
    item["inbound_no"] = inbound.get("inbound_no") if inbound else None
    return {"item": item, "lines": serialize_rows(lines), "attachments": atts}


@router.post("/aux/purchases")
def create_purchase(req: PurchaseCreate, user: UserDep):
    if req.pay_type not in ("monthly", "cash"):
        raise HTTPException(status_code=400, detail="pay_type 无效")
    c_uid, c_name, u_uid, u_name = audit_tuple(user)
    po_no = gen_no("PO-M-" if req.pay_type == "monthly" else "PO-C-")
    total_qty = sum((_d(x.qty) for x in req.lines), Decimal("0"))
    total_amt = sum((_d(x.qty) * _d(x.unit_price) for x in req.lines), Decimal("0"))
    summary = "、".join(f"{x.name_spec}×{x.qty}" for x in req.lines)[:240]
    status = "confirmed" if req.pay_type == "monthly" else "draft"
    supplier_id = None
    if req.supplier_name:
        s = fetch_one("SELECT id FROM wh_suppliers WHERE name=%s", (req.supplier_name.strip(),))
        if s:
            supplier_id = s["id"]
        else:
            supplier_id = execute(
                """
                INSERT INTO wh_suppliers (name, settle_type, created_by_user_id, created_by_name, updated_by_user_id, updated_by_name)
                VALUES (%s,%s,%s,%s,%s,%s)
                """,
                (req.supplier_name.strip(), req.pay_type if req.pay_type == "monthly" else "cash", c_uid, c_name, u_uid, u_name),
            )
    po_id = execute(
        """
        INSERT INTO wh_purchase_orders
          (po_no, pay_type, supplier_id, supplier_name, biz_date, status,
           purchaser_user_id, purchaser_name, total_qty, total_amount, remark, summary,
           created_by_user_id, created_by_name, updated_by_user_id, updated_by_name)
        VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
        """,
        (
            po_no,
            req.pay_type,
            supplier_id,
            req.supplier_name,
            req.biz_date,
            status,
            user.id,
            req.purchaser_name or user.name,
            total_qty,
            total_amt,
            req.remark,
            summary,
            c_uid,
            c_name,
            u_uid,
            u_name,
        ),
    )
    for idx, line in enumerate(req.lines, start=1):
        execute(
            """
            INSERT INTO wh_purchase_order_lines
              (po_id, item_id, name_spec, uom, qty, unit_price, amount, line_no,
               created_by_user_id, created_by_name, updated_by_user_id, updated_by_name)
            VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
            """,
            (
                po_id,
                line.item_id,
                line.name_spec,
                line.uom,
                line.qty,
                line.unit_price,
                float(_d(line.qty) * _d(line.unit_price)),
                idx,
                c_uid,
                c_name,
                u_uid,
                u_name,
            ),
        )
    write_op_log(
        domain="aux",
        action="create",
        entity_type="purchase_order",
        entity_id=po_id,
        entity_no=po_no,
        user=user,
        change_summary={"op": "create", "pay_type": req.pay_type, "amount": float(total_amt)},
    )
    warnings: list[dict] = []
    for line in req.lines:
        w = record_price_and_alerts(
            item_id=line.item_id,
            supplier_id=supplier_id,
            unit_price=_d(line.unit_price),
            biz_date=req.biz_date,
            po_id=po_id,
            name_spec=line.name_spec,
            user=user,
        )
        if w:
            warnings.append(w)
    row = fetch_one("SELECT * FROM wh_purchase_orders WHERE id=%s", (po_id,))
    return {"item": serialize_row(row), "price_warnings": warnings}


class PurchaseUpdate(BaseModel):
    biz_date: Optional[date] = None
    supplier_name: Optional[str] = None
    purchaser_name: Optional[str] = None
    remark: Optional[str] = None
    lines: Optional[list[PurchaseLine]] = None


@router.put("/aux/purchases/{po_id}")
def update_purchase(po_id: int, req: PurchaseUpdate, user: UserDep):
    _ensure()
    row = fetch_one("SELECT * FROM wh_purchase_orders WHERE id=%s", (po_id,))
    if not row:
        raise HTTPException(status_code=404, detail="采购单不存在")
    if row["status"] != "draft":
        raise HTTPException(status_code=400, detail="仅草稿可修改")
    fields = req.model_dump(exclude_unset=True)
    lines = fields.pop("lines", None)
    u_uid, u_name = update_audit_tuple(user)
    if fields:
        sets = [f"{k}=%s" for k in fields]
        params = list(fields.values())
        sets.extend(["updated_by_user_id=%s", "updated_by_name=%s"])
        params.extend([u_uid, u_name, po_id])
        execute(f"UPDATE wh_purchase_orders SET {', '.join(sets)} WHERE id=%s", tuple(params))
    if lines is not None:
        if not lines:
            raise HTTPException(status_code=400, detail="明细不能为空")
        execute("DELETE FROM wh_purchase_order_lines WHERE po_id=%s", (po_id,))
        c_uid, c_name, uu, un = audit_tuple(user)
        total_qty = sum((_d(x.qty) for x in lines), Decimal("0"))
        total_amt = sum((_d(x.qty) * _d(x.unit_price) for x in lines), Decimal("0"))
        summary = "、".join(f"{x.name_spec}×{x.qty}" for x in lines)[:240]
        for idx, line in enumerate(lines, start=1):
            execute(
                """
                INSERT INTO wh_purchase_order_lines
                  (po_id, item_id, name_spec, uom, qty, unit_price, amount, line_no,
                   created_by_user_id, created_by_name, updated_by_user_id, updated_by_name)
                VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
                """,
                (
                    po_id,
                    line.item_id,
                    line.name_spec,
                    line.uom,
                    line.qty,
                    line.unit_price,
                    float(_d(line.qty) * _d(line.unit_price)),
                    idx,
                    c_uid,
                    c_name,
                    uu,
                    un,
                ),
            )
        execute(
            """
            UPDATE wh_purchase_orders
            SET total_qty=%s, total_amount=%s, summary=%s, updated_by_user_id=%s, updated_by_name=%s
            WHERE id=%s
            """,
            (total_qty, total_amt, summary, u_uid, u_name, po_id),
        )
    write_op_log(
        domain="aux",
        action="update",
        entity_type="purchase_order",
        entity_id=po_id,
        entity_no=str(row.get("po_no") or ""),
        user=user,
        change_summary={"op": "update"},
    )
    return get_purchase(po_id, user)


@router.delete("/aux/purchases/{po_id}")
def delete_purchase(po_id: int, user: UserDep):
    _ensure()
    row = fetch_one("SELECT * FROM wh_purchase_orders WHERE id=%s", (po_id,))
    if not row:
        raise HTTPException(status_code=404, detail="采购单不存在")
    if row["status"] != "draft":
        raise HTTPException(status_code=400, detail="仅草稿可删除")
    execute("DELETE FROM wh_purchase_order_lines WHERE po_id=%s", (po_id,))
    execute("DELETE FROM wh_purchase_orders WHERE id=%s", (po_id,))
    write_op_log(
        domain="aux",
        action="delete",
        entity_type="purchase_order",
        entity_id=po_id,
        entity_no=str(row.get("po_no") or ""),
        user=user,
        change_summary={"op": "delete"},
    )
    return {"status": "ok"}


@router.post("/aux/purchases/{po_id}/void")
def void_purchase(po_id: int, user: UserDep):
    _ensure()
    row = fetch_one("SELECT * FROM wh_purchase_orders WHERE id=%s", (po_id,))
    if not row:
        raise HTTPException(status_code=404, detail="采购单不存在")
    if row["status"] in ("voided", "cancelled"):
        raise HTTPException(status_code=400, detail="该采购单已作废")
    if row["status"] == "draft":
        raise HTTPException(status_code=400, detail="草稿请直接删除")
    inbound = _active_inbound_for_po(po_id, row.get("po_no"))
    if inbound:
        raise HTTPException(status_code=400, detail=f"已入库（{inbound['inbound_no']}），请先作废入库单")
    u_uid, u_name = update_audit_tuple(user)
    execute(
        "UPDATE wh_purchase_orders SET status='voided', updated_by_user_id=%s, updated_by_name=%s WHERE id=%s",
        (u_uid, u_name, po_id),
    )
    write_op_log(
        domain="aux",
        action="void",
        entity_type="purchase_order",
        entity_id=po_id,
        entity_no=str(row.get("po_no") or ""),
        user=user,
        change_summary={"op": "void", "from": row.get("status")},
    )
    return {"item": serialize_row(fetch_one("SELECT * FROM wh_purchase_orders WHERE id=%s", (po_id,)))}


class CashStatusUpdate(BaseModel):
    to_status: str
    comment: Optional[str] = None


_CASH_FLOW = {
    "draft": "pending_verify",
    "pending_verify": "pending_approve",
    "pending_approve": "pending_pay",
    "pending_pay": "paid",
    "paid": "posted",
}


@router.post("/aux/purchases/{po_id}/advance")
def advance_cash_purchase(po_id: int, req: CashStatusUpdate, user: UserDep):
    _ensure()
    row = fetch_one("SELECT * FROM wh_purchase_orders WHERE id=%s", (po_id,))
    if not row:
        raise HTTPException(status_code=404, detail="采购单不存在")
    if row["pay_type"] != "cash":
        raise HTTPException(status_code=400, detail="仅现金采购可推进状态")
    cur_status = row["status"]
    expected = _CASH_FLOW.get(cur_status)
    if req.to_status != expected:
        raise HTTPException(status_code=400, detail=f"当前状态 {cur_status} 不可转到 {req.to_status}")
    require_cash_role(user, cur_status)
    if cur_status == "draft" and count_po_attachments(po_id) < 1:
        raise HTTPException(status_code=400, detail="提交核实前请至少上传 1 张收据/二联单")
    u_uid, u_name = update_audit_tuple(user)
    execute(
        "UPDATE wh_purchase_orders SET status=%s, updated_by_user_id=%s, updated_by_name=%s WHERE id=%s",
        (req.to_status, u_uid, u_name, po_id),
    )
    action = {"pending_verify": "update", "pending_approve": "verify", "pending_pay": "approve", "paid": "pay", "posted": "update"}.get(
        req.to_status, "update"
    )
    record_cash_approval(po_id=po_id, from_status=cur_status, to_status=req.to_status, user=user, comment=req.comment)
    write_op_log(
        domain="aux",
        action=action,
        entity_type="purchase_order",
        entity_id=po_id,
        entity_no=row["po_no"],
        user=user,
        change_summary={"op": action, "from": cur_status, "to": req.to_status, "comment": req.comment},
    )
    if req.to_status == "paid":
        prev = _latest_open_cash()
        prev_bal = _d(prev["next_balance"] if prev else 0)
        spent = _d(row.get("total_amount") or 0)
        c_uid, c_name, uu, un = audit_tuple(user)
        execute(
            """
            INSERT INTO wh_cash_ledger
              (biz_date, transfer_in, prev_balance, spent, next_balance, remark, source, status,
               created_by_user_id, created_by_name, updated_by_user_id, updated_by_name)
            VALUES (%s,0,%s,%s,%s,%s,'purchase_pay','posted',%s,%s,%s,%s)
            """,
            (
                date.today(),
                prev_bal,
                spent,
                prev_bal - spent,
                f"现金采购付款 {row['po_no']}",
                c_uid,
                c_name,
                uu,
                un,
            ),
        )
        refresh_cash_warn(user)
    return {"item": serialize_row(fetch_one("SELECT * FROM wh_purchase_orders WHERE id=%s", (po_id,)))}


class CashReject(BaseModel):
    comment: Optional[str] = None


@router.post("/aux/purchases/{po_id}/reject")
def reject_cash_purchase(po_id: int, req: CashReject, user: UserDep):
    _ensure()
    row = fetch_one("SELECT * FROM wh_purchase_orders WHERE id=%s", (po_id,))
    if not row:
        raise HTTPException(status_code=404, detail="采购单不存在")
    if row["pay_type"] != "cash":
        raise HTTPException(status_code=400, detail="仅现金采购可驳回")
    cur_status = row["status"]
    to_status = CASH_REJECT_TO.get(cur_status)
    if not to_status:
        raise HTTPException(status_code=400, detail=f"当前状态 {cur_status} 不可驳回")
    require_cash_role(user, cur_status)
    u_uid, u_name = update_audit_tuple(user)
    execute(
        "UPDATE wh_purchase_orders SET status=%s, updated_by_user_id=%s, updated_by_name=%s WHERE id=%s",
        (to_status, u_uid, u_name, po_id),
    )
    record_cash_approval(po_id=po_id, from_status=cur_status, to_status=to_status, user=user, comment=req.comment or "驳回")
    write_op_log(
        domain="aux",
        action="update",
        entity_type="purchase_order",
        entity_id=po_id,
        entity_no=row["po_no"],
        user=user,
        change_summary={"op": "reject", "from": cur_status, "to": to_status, "comment": req.comment},
    )
    return {"item": serialize_row(fetch_one("SELECT * FROM wh_purchase_orders WHERE id=%s", (po_id,)))}


@router.post("/aux/purchases/{po_id}/attachments")
async def upload_purchase_attachment(po_id: int, user: UserDep, file: UploadFile = File(...)):
    row = fetch_one("SELECT * FROM wh_purchase_orders WHERE id=%s", (po_id,))
    if not row:
        raise HTTPException(status_code=404, detail="采购单不存在")
    raw = await file.read()
    att = save_purchase_attachment(
        po=row,
        original_name=file.filename or "receipt.bin",
        raw=raw,
        content_type=file.content_type or "",
        user=user,
    )
    return {"item": serialize_row(att)}


@router.get("/aux/purchases/{po_id}/attachments")
def get_purchase_attachments(po_id: int, _user: UserDep):
    if not fetch_one("SELECT id FROM wh_purchase_orders WHERE id=%s", (po_id,)):
        raise HTTPException(status_code=404, detail="采购单不存在")
    return {"items": serialize_rows(list_po_attachments(po_id))}


@router.get("/aux/attachments/{att_id}/file")
def download_attachment(att_id: int, user: UserDep):
    _ensure()
    row = fetch_one("SELECT * FROM wh_attachments WHERE id=%s", (att_id,))
    if not row:
        raise HTTPException(status_code=404, detail="附件不存在")
    ym = row.get("archived_ym")
    if ym:
        zip_path = receipts_archives_root() / f"{ym}.zip"
        if zip_path.is_file():
            write_op_log(
                domain="aux",
                action="export",
                entity_type="attachment",
                entity_id=att_id,
                user=user,
                change_summary={"op": "download_archive", "ym": ym},
            )
            return FileResponse(zip_path, filename=f"{ym}.zip", media_type="application/zip")
        raise HTTPException(status_code=404, detail="该月收据已归档但压缩包不存在")
    path = attachment_abs_path(row)
    write_op_log(
        domain="aux",
        action="export",
        entity_type="attachment",
        entity_id=att_id,
        user=user,
        change_summary={"op": "download", "file_name": row.get("file_name")},
    )
    return FileResponse(path, filename=str(row.get("file_name") or path.name), media_type=row.get("content_type") or None)


class PurchaseInboundCreate(BaseModel):
    biz_date: date
    location_id: Optional[int] = None
    operator_name: Optional[str] = None


@router.post("/aux/purchases/{po_id}/inbound")
def inbound_from_purchase(po_id: int, req: PurchaseInboundCreate, user: UserDep):
    _ensure()
    po = fetch_one("SELECT * FROM wh_purchase_orders WHERE id=%s", (po_id,))
    if not po:
        raise HTTPException(status_code=404, detail="采购单不存在")
    if po["pay_type"] == "monthly" and po["status"] != "confirmed":
        raise HTTPException(status_code=400, detail="月结采购需已确认才能入库")
    if po["pay_type"] == "cash" and po["status"] not in ("paid", "posted"):
        raise HTTPException(status_code=400, detail="现金采购需已付款后才能入库")
    existed = _active_inbound_for_po(po_id, po.get("po_no"))
    if existed:
        raise HTTPException(status_code=400, detail=f"该采购单已入库（{existed['inbound_no']}）")
    lines = fetch_all("SELECT * FROM wh_purchase_order_lines WHERE po_id=%s ORDER BY line_no, id", (po_id,))
    if not lines:
        raise HTTPException(status_code=400, detail="采购单无明细")
    inbound_lines: list[InboundLine] = []
    for ln in lines:
        if not ln.get("item_id"):
            raise HTTPException(status_code=400, detail="采购明细未关联辅料主数据，无法一键入库")
        inbound_lines.append(
            InboundLine(
                item_id=int(ln["item_id"]),
                location_id=req.location_id,
                qty=float(ln["qty"]),
                unit_price=float(ln["unit_price"] or 0),
                uom=ln.get("uom"),
            )
        )
    return create_inbound(
        InboundCreate(
            biz_date=req.biz_date,
            source_type="purchase",
            supplier_name=po.get("supplier_name"),
            ref_po_no=po.get("po_no"),
            ref_po_id=po_id,
            operator_name=req.operator_name or user.name,
            lines=inbound_lines,
        ),
        user,
    )


# ---------- borrow / cash ----------
class BorrowCreate(BaseModel):
    item_name_spec: str
    qty: float = 1
    department: Optional[str] = None
    borrow_date: date
    due_date: Optional[date] = None
    borrower_name: str
    keeper_name: Optional[str] = None
    purpose: Optional[str] = None


@router.get("/aux/borrows")
def list_borrows(_user: UserDep, date_from: str = "", date_to: str = ""):
    _ensure()
    sql = "SELECT * FROM wh_borrow_records WHERE 1=1"
    params: list[Any] = []
    if date_from.strip():
        sql += " AND borrow_date>=%s"
        params.append(date_from.strip())
    if date_to.strip():
        sql += " AND borrow_date<=%s"
        params.append(date_to.strip())
    sql += " ORDER BY borrow_date DESC, id DESC LIMIT 200"
    rows = fetch_all(sql, tuple(params))
    today = date.today()
    out = []
    for r in rows:
        item = serialize_row(r) or {}
        if item.get("status") == "borrowed" and r.get("due_date") and r["due_date"] < today:
            item["status"] = "overdue"
        out.append(item)
    return {"items": out}


@router.get("/aux/borrows/{borrow_id}")
def get_borrow(borrow_id: int, _user: UserDep):
    row = fetch_one("SELECT * FROM wh_borrow_records WHERE id=%s", (borrow_id,))
    if not row:
        raise HTTPException(status_code=404, detail="借用记录不存在")
    item = serialize_row(row) or {}
    if item.get("status") == "borrowed" and row.get("due_date") and row["due_date"] < date.today():
        item["status"] = "overdue"
    return {"item": item}


class BorrowUpdate(BaseModel):
    item_name_spec: Optional[str] = None
    qty: Optional[float] = None
    department: Optional[str] = None
    borrow_date: Optional[date] = None
    due_date: Optional[date] = None
    borrower_name: Optional[str] = None
    keeper_name: Optional[str] = None
    purpose: Optional[str] = None


@router.put("/aux/borrows/{borrow_id}")
def update_borrow(borrow_id: int, req: BorrowUpdate, user: UserDep):
    _ensure()
    row = fetch_one("SELECT * FROM wh_borrow_records WHERE id=%s", (borrow_id,))
    if not row:
        raise HTTPException(status_code=404, detail="借用记录不存在")
    if row.get("status") == "returned":
        raise HTTPException(status_code=400, detail="已归还记录不可修改")
    fields = req.model_dump(exclude_unset=True)
    if not fields:
        return {"item": serialize_row(row)}
    u_uid, u_name = update_audit_tuple(user)
    sets = [f"{k}=%s" for k in fields]
    params = list(fields.values())
    sets.extend(["updated_by_user_id=%s", "updated_by_name=%s"])
    params.extend([u_uid, u_name, borrow_id])
    execute(f"UPDATE wh_borrow_records SET {', '.join(sets)} WHERE id=%s", tuple(params))
    if req.department:
        add_dict_option("department", req.department, user)
    if req.purpose:
        add_dict_option("purpose", req.purpose, user)
    write_op_log(
        domain="aux",
        action="update",
        entity_type="borrow",
        entity_id=borrow_id,
        user=user,
        change_summary={"op": "update", "fields": list(fields.keys())},
    )
    return get_borrow(borrow_id, user)


@router.delete("/aux/borrows/{borrow_id}")
def delete_borrow(borrow_id: int, user: UserDep):
    _ensure()
    row = fetch_one("SELECT * FROM wh_borrow_records WHERE id=%s", (borrow_id,))
    if not row:
        raise HTTPException(status_code=404, detail="借用记录不存在")
    if row.get("status") == "returned":
        raise HTTPException(status_code=400, detail="已归还记录不可删除")
    execute("DELETE FROM wh_borrow_records WHERE id=%s", (borrow_id,))
    write_op_log(
        domain="aux",
        action="delete",
        entity_type="borrow",
        entity_id=borrow_id,
        user=user,
        change_summary={"op": "delete", "item": row.get("item_name_spec")},
    )
    return {"status": "ok"}


@router.get("/aux/exports/borrows")
def export_borrows(user: UserDep, date_from: str = "", date_to: str = ""):
    data = list_borrows(user, date_from=date_from, date_to=date_to)
    wb = Workbook()
    ws = wb.active
    ws.title = "借用登记"
    ws.append(["物品", "数量", "部门", "借用日期", "约定归还", "实际归还", "借用人", "库管员", "用途", "状态", "录入人"])
    for r in data["items"]:
        ws.append(
            [
                r.get("item_name_spec"),
                r.get("qty"),
                r.get("department"),
                r.get("borrow_date"),
                r.get("due_date"),
                r.get("return_date"),
                r.get("borrower_name"),
                r.get("keeper_name"),
                r.get("purpose"),
                r.get("status"),
                r.get("created_by_name"),
            ]
        )
    file_name = f"借用登记_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return _excel_response(wb, file_name, user, "borrows", {"date_from": date_from, "date_to": date_to})


@router.post("/aux/borrows")
def create_borrow(req: BorrowCreate, user: UserDep):
    _ensure()
    if req.department:
        add_dict_option("department", req.department, user)
    if req.purpose:
        add_dict_option("purpose", req.purpose, user)
    c_uid, c_name, u_uid, u_name = audit_tuple(user)
    bid = execute(
        """
        INSERT INTO wh_borrow_records
          (item_name_spec, qty, department, borrow_date, due_date, borrower_name, keeper_name, purpose, status,
           created_by_user_id, created_by_name, updated_by_user_id, updated_by_name)
        VALUES (%s,%s,%s,%s,%s,%s,%s,%s,'borrowed',%s,%s,%s,%s)
        """,
        (
            req.item_name_spec,
            req.qty,
            req.department,
            req.borrow_date,
            req.due_date,
            req.borrower_name,
            req.keeper_name or user.name,
            req.purpose,
            c_uid,
            c_name,
            u_uid,
            u_name,
        ),
    )
    write_op_log(
        domain="aux",
        action="create",
        entity_type="borrow",
        entity_id=bid,
        user=user,
        change_summary={"op": "create", "item": req.item_name_spec, "borrower": req.borrower_name},
    )
    return {"item": serialize_row(fetch_one("SELECT * FROM wh_borrow_records WHERE id=%s", (bid,)))}


@router.post("/aux/borrows/{borrow_id}/return")
def return_borrow(borrow_id: int, user: UserDep):
    row = fetch_one("SELECT * FROM wh_borrow_records WHERE id=%s", (borrow_id,))
    if not row:
        raise HTTPException(status_code=404, detail="借用记录不存在")
    u_uid, u_name = update_audit_tuple(user)
    execute(
        """
        UPDATE wh_borrow_records
        SET status='returned', return_date=%s, updated_by_user_id=%s, updated_by_name=%s
        WHERE id=%s
        """,
        (date.today(), u_uid, u_name, borrow_id),
    )
    write_op_log(
        domain="aux",
        action="return",
        entity_type="borrow",
        entity_id=borrow_id,
        user=user,
        change_summary={"op": "return"},
    )
    return {"item": serialize_row(fetch_one("SELECT * FROM wh_borrow_records WHERE id=%s", (borrow_id,)))}


class CashLedgerCreate(BaseModel):
    biz_date: date
    transfer_in: float = 0
    spent: float = 0
    remark: Optional[str] = None


class CashLedgerUpdate(BaseModel):
    remark: Optional[str] = None


class AuxSettingsUpdate(BaseModel):
    cash_warn_amount: Optional[str] = None


def _latest_open_cash() -> dict | None:
    return fetch_one(
        f"SELECT * FROM wh_cash_ledger WHERE {_POSTED} ORDER BY id DESC LIMIT 1"
    )


def _latest_manual_cash() -> dict | None:
    return fetch_one(
        f"""
        SELECT * FROM wh_cash_ledger
        WHERE {_POSTED} AND COALESCE(source,'manual')='manual'
        ORDER BY id DESC LIMIT 1
        """
    )


@router.get("/aux/settings")
def get_aux_settings(_user: UserDep):
    _ensure()
    return {"cash_warn_amount": get_setting("cash_warn_amount", "")}


@router.put("/aux/settings")
def put_aux_settings(req: AuxSettingsUpdate, user: UserDep):
    _ensure()
    if req.cash_warn_amount is not None:
        set_setting("cash_warn_amount", str(req.cash_warn_amount), user)
        refresh_cash_warn(user)
    return get_aux_settings(user)


@router.get("/aux/cash-ledger")
def list_cash_ledger(_user: UserDep):
    _ensure()
    rows = fetch_all("SELECT * FROM wh_cash_ledger ORDER BY biz_date DESC, id DESC LIMIT 200")
    latest_manual = _latest_manual_cash()
    latest_id = latest_manual["id"] if latest_manual else None
    items = []
    for r in serialize_rows(rows) or []:
        src = str(r.get("source") or "manual")
        if src == "manual" and str(r.get("remark") or "").startswith("现金采购付款"):
            src = "purchase_pay"
        r["source"] = src
        r["can_void"] = bool(latest_id and r.get("id") == latest_id and str(r.get("status") or "posted") != "voided")
        items.append(r)
    return {"items": items, "cash_warn_amount": get_setting("cash_warn_amount", "")}


@router.post("/aux/cash-ledger")
def create_cash_ledger(req: CashLedgerCreate, user: UserDep):
    _ensure()
    prev = _latest_open_cash()
    prev_bal = _d(prev["next_balance"] if prev else 0)
    transfer = _d(req.transfer_in)
    spent = _d(req.spent)
    next_bal = prev_bal + transfer - spent
    c_uid, c_name, u_uid, u_name = audit_tuple(user)
    lid = execute(
        """
        INSERT INTO wh_cash_ledger
          (biz_date, transfer_in, prev_balance, spent, next_balance, remark, source, status,
           created_by_user_id, created_by_name, updated_by_user_id, updated_by_name)
        VALUES (%s,%s,%s,%s,%s,%s,'manual','posted',%s,%s,%s,%s)
        """,
        (req.biz_date, transfer, prev_bal, spent, next_bal, req.remark, c_uid, c_name, u_uid, u_name),
    )
    write_op_log(
        domain="aux",
        action="create",
        entity_type="cash_ledger",
        entity_id=lid,
        user=user,
        change_summary={"op": "create", "transfer_in": float(transfer), "spent": float(spent), "remark": req.remark},
    )
    refresh_cash_warn(user)
    return {"item": serialize_row(fetch_one("SELECT * FROM wh_cash_ledger WHERE id=%s", (lid,)))}


@router.put("/aux/cash-ledger/{ledger_id}")
def update_cash_ledger(ledger_id: int, req: CashLedgerUpdate, user: UserDep):
    _ensure()
    row = fetch_one("SELECT * FROM wh_cash_ledger WHERE id=%s", (ledger_id,))
    if not row:
        raise HTTPException(status_code=404, detail="账本记录不存在")
    u_uid, u_name = update_audit_tuple(user)
    execute(
        "UPDATE wh_cash_ledger SET remark=%s, updated_by_user_id=%s, updated_by_name=%s WHERE id=%s",
        (req.remark, u_uid, u_name, ledger_id),
    )
    write_op_log(
        domain="aux",
        action="update",
        entity_type="cash_ledger",
        entity_id=ledger_id,
        user=user,
        change_summary={"op": "update", "remark": req.remark},
    )
    return {"item": serialize_row(fetch_one("SELECT * FROM wh_cash_ledger WHERE id=%s", (ledger_id,)))}


@router.post("/aux/cash-ledger/{ledger_id}/void")
def void_cash_ledger(ledger_id: int, user: UserDep):
    _ensure()
    row = fetch_one("SELECT * FROM wh_cash_ledger WHERE id=%s", (ledger_id,))
    if not row:
        raise HTTPException(status_code=404, detail="账本记录不存在")
    if str(row.get("status") or "posted") == "voided":
        raise HTTPException(status_code=400, detail="该记录已作废")
    src = str(row.get("source") or "manual")
    if src == "purchase_pay" or str(row.get("remark") or "").startswith("现金采购付款"):
        raise HTTPException(status_code=400, detail="采购付款自动行不可作废，仅可改备注")
    latest = _latest_manual_cash()
    if not latest or int(latest["id"]) != ledger_id:
        raise HTTPException(status_code=400, detail="仅可作废最新一条手工登记")
    u_uid, u_name = update_audit_tuple(user)
    execute(
        "UPDATE wh_cash_ledger SET status='voided', updated_by_user_id=%s, updated_by_name=%s WHERE id=%s",
        (u_uid, u_name, ledger_id),
    )
    write_op_log(
        domain="aux",
        action="void",
        entity_type="cash_ledger",
        entity_id=ledger_id,
        user=user,
        change_summary={"op": "void"},
    )
    refresh_cash_warn(user)
    return {"item": serialize_row(fetch_one("SELECT * FROM wh_cash_ledger WHERE id=%s", (ledger_id,)))}


@router.get("/aux/reconcile")
def list_reconcile(user: UserDep, month: str = ""):
    _ensure()
    month, start, end = _month_bounds(month)
    rows = fetch_all(
        """
        SELECT
          COALESCE(p.supplier_name, i.supplier_name, '未指定') AS supplier,
          %s AS month,
          COALESCE(p.purchase_amt,0) AS purchase_amt,
          COALESCE(i.inbound_amt,0) AS inbound_amt,
          COALESCE(p.purchase_amt,0) - COALESCE(i.inbound_amt,0) AS diff
        FROM (
          SELECT supplier_name, SUM(total_amount) purchase_amt
          FROM wh_purchase_orders
          WHERE biz_date >= %s AND biz_date < %s AND pay_type='monthly'
            AND status NOT IN ('cancelled','rejected','draft','voided')
          GROUP BY supplier_name
        ) p
        LEFT JOIN (
          SELECT supplier_name, SUM(total_amount) inbound_amt
          FROM wh_inbound_orders
          WHERE biz_date >= %s AND biz_date < %s AND COALESCE(status,'posted') <> 'voided'
          GROUP BY supplier_name
        ) i ON COALESCE(p.supplier_name,'') = COALESCE(i.supplier_name,'')
        """,
        (month, start, end, start, end),
    )
    imap = invoice_rows(month)
    items = serialize_rows(rows)
    recon_titles: list[str] = []
    for it in items:
        rec = imap.get(str(it.get("supplier") or ""), {})
        st = str(rec.get("status") or "unissued")
        it["invoice"] = "已开" if st == "issued" else "未开"
        it["invoice_status"] = st
        it["diff_note"] = rec.get("diff_note") or ""
        if abs(float(it.get("diff") or 0)) >= 0.01:
            recon_titles.append(f"{month} {it.get('supplier')} 采购入库金额差 {float(it.get('diff') or 0):.2f} 元")
    refresh_named_alerts(
        alert_type="reconcile_diff",
        title_prefix=f"{month} ",
        titles=recon_titles,
        level="warning",
        user=user,
    )
    return {"items": items, "month": month}


class InvoiceUpdate(BaseModel):
    supplier: str
    month: str
    status: str  # issued|unissued
    diff_note: Optional[str] = None


@router.put("/aux/reconcile/invoice")
def update_invoice(req: InvoiceUpdate, user: UserDep):
    _ensure()
    row = upsert_invoice(req.supplier, req.month, req.status, user, diff_note=req.diff_note)
    return {"item": serialize_row(row)}


@router.get("/aux/reconcile/detail")
def reconcile_detail(_user: UserDep, supplier: str, month: str = ""):
    _ensure()
    month, start, end = _month_bounds(month)
    name = (supplier or "").strip() or "未指定"
    purchases = fetch_all(
        """
        SELECT id, po_no, biz_date, status, total_qty, total_amount, purchaser_name, summary
        FROM wh_purchase_orders
        WHERE pay_type='monthly' AND COALESCE(supplier_name,'未指定')=%s
          AND biz_date >= %s AND biz_date < %s
          AND status NOT IN ('cancelled','rejected','draft','voided')
        ORDER BY biz_date, id
        """,
        (name, start, end),
    )
    inbounds = fetch_all(
        f"""
        SELECT id, inbound_no, biz_date, status, total_qty, total_amount, ref_po_no, operator_name
        FROM wh_inbound_orders
        WHERE COALESCE(supplier_name,'未指定')=%s AND biz_date >= %s AND biz_date < %s AND COALESCE(status,'posted') <> 'voided'
        ORDER BY biz_date, id
        """,
        (name, start, end),
    )
    purchase_amt = sum(float(r.get("total_amount") or 0) for r in purchases)
    inbound_amt = sum(float(r.get("total_amount") or 0) for r in inbounds)
    rec = invoice_rows(month).get(name, {})
    return {
        "supplier": name,
        "month": month,
        "purchase_amt": purchase_amt,
        "inbound_amt": inbound_amt,
        "diff": purchase_amt - inbound_amt,
        "diff_note": rec.get("diff_note") or "",
        "invoice_status": rec.get("status") or "unissued",
        "purchases": serialize_rows(purchases),
        "inbounds": serialize_rows(inbounds),
    }


@router.get("/aux/receipts")
def list_receipts(_user: UserDep, keyword: str = ""):
    _ensure()
    return {"items": search_receipts(keyword)}


@router.get("/aux/receipts/archives/{ym}/file")
def download_receipt_archive(ym: str, user: UserDep):
    _ensure()
    safe = "".join(ch for ch in ym if ch.isdigit())[:6]
    path = receipts_archives_root() / f"{safe}.zip"
    if not path.is_file():
        raise HTTPException(status_code=404, detail="该月归档包不存在")
    write_op_log(
        domain="aux",
        action="export",
        entity_type="receipt_archive",
        user=user,
        change_summary={"op": "download", "ym": safe},
    )
    return FileResponse(path, filename=f"{safe}.zip", media_type="application/zip")


# ---------- finished goods skeleton ----------
@router.get("/fg/warehouses")
def list_fg_warehouses(_user: UserDep):
    return {"items": serialize_rows(fetch_all("SELECT * FROM wh_fg_warehouses ORDER BY id"))}


class FgItemCreate(BaseModel):
    item_code: str
    name_spec: str
    material_grade: Optional[str] = None
    uom: str = "件"
    remark: Optional[str] = None


class FgTxnCreate(BaseModel):
    biz_date: date
    txn_type: str  # inbound_finish|outbound_ship|inbound_outsource|manual
    warehouse_id: int
    item_id: int
    heat_no: Optional[str] = ""
    order_no: Optional[str] = ""
    location_name: Optional[str] = None
    qty: float
    operator_name: Optional[str] = None
    remark: Optional[str] = None


@router.get("/fg/balances")
def list_fg_balances(_user: UserDep, warehouse_id: Optional[int] = None):
    sql = """
        SELECT b.*, w.name AS warehouse, i.item_code, i.name_spec, i.material_grade, i.uom, i.remark
        FROM wh_fg_balances b
        JOIN wh_fg_warehouses w ON w.id = b.warehouse_id
        JOIN wh_fg_items i ON i.id = b.item_id
        WHERE 1=1
    """
    params: list[Any] = []
    if warehouse_id is not None:
        sql += " AND b.warehouse_id=%s"
        params.append(warehouse_id)
    sql += " ORDER BY b.id DESC"
    return {"items": serialize_rows(fetch_all(sql, tuple(params)))}


@router.post("/fg/items")
def create_fg_item(req: FgItemCreate, user: UserDep):
    if fetch_one("SELECT id FROM wh_fg_items WHERE item_code=%s", (req.item_code.strip(),)):
        raise HTTPException(status_code=400, detail="成品编码已存在")
    c_uid, c_name, u_uid, u_name = audit_tuple(user)
    iid = execute(
        """
        INSERT INTO wh_fg_items
          (item_code, name_spec, material_grade, uom, remark,
           created_by_user_id, created_by_name, updated_by_user_id, updated_by_name)
        VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s)
        """,
        (req.item_code.strip(), req.name_spec.strip(), req.material_grade, req.uom, req.remark, c_uid, c_name, u_uid, u_name),
    )
    write_op_log(
        domain="fg",
        action="create",
        entity_type="fg_item",
        entity_id=iid,
        entity_no=req.item_code,
        user=user,
        change_summary={"op": "create", "name_spec": req.name_spec},
    )
    return {"item": serialize_row(fetch_one("SELECT * FROM wh_fg_items WHERE id=%s", (iid,)))}


@router.get("/fg/items")
def list_fg_items(_user: UserDep):
    return {"items": serialize_rows(fetch_all("SELECT * FROM wh_fg_items WHERE is_active=1 ORDER BY id DESC"))}


@router.get("/fg/items/{item_id}")
def get_fg_item(item_id: int, _user: UserDep):
    row = fetch_one("SELECT * FROM wh_fg_items WHERE id=%s", (item_id,))
    if not row:
        raise HTTPException(status_code=404, detail="成品档案不存在")
    return {"item": serialize_row(row)}


class FgItemUpdate(BaseModel):
    name_spec: Optional[str] = None
    material_grade: Optional[str] = None
    uom: Optional[str] = None
    remark: Optional[str] = None
    is_active: Optional[int] = None


@router.put("/fg/items/{item_id}")
def update_fg_item(item_id: int, req: FgItemUpdate, user: UserDep):
    before = fetch_one("SELECT * FROM wh_fg_items WHERE id=%s", (item_id,))
    if not before:
        raise HTTPException(status_code=404, detail="成品档案不存在")
    fields = req.model_dump(exclude_unset=True)
    if not fields:
        return {"item": serialize_row(before)}
    if fields.get("is_active") == 0:
        bal = fetch_one(
            "SELECT COALESCE(SUM(qty_on_hand),0) AS qty FROM wh_fg_balances WHERE item_id=%s",
            (item_id,),
        )
        if _d(bal["qty"] if bal else 0) > 0:
            raise HTTPException(status_code=400, detail="仍有库存，禁止停用")
    u_uid, u_name = update_audit_tuple(user)
    sets = [f"{k}=%s" for k in fields]
    params = list(fields.values())
    sets.extend(["updated_by_user_id=%s", "updated_by_name=%s"])
    params.extend([u_uid, u_name, item_id])
    execute(f"UPDATE wh_fg_items SET {', '.join(sets)} WHERE id=%s", tuple(params))
    after = fetch_one("SELECT * FROM wh_fg_items WHERE id=%s", (item_id,)) or {}
    write_op_log(
        domain="fg",
        action="update",
        entity_type="fg_item",
        entity_id=item_id,
        entity_no=str(before.get("item_code") or ""),
        user=user,
        change_summary={"op": "update", "fields": diff_fields(before, after, list(fields.keys()))},
    )
    return {"item": serialize_row(after)}


@router.get("/fg/exports/balances")
def export_fg_balances(user: UserDep, warehouse_id: Optional[int] = None):
    data = list_fg_balances(user, warehouse_id=warehouse_id)
    wb = Workbook()
    ws = wb.active
    ws.title = "成品库存"
    ws.append(["仓库类型", "编码", "规格型号", "材质", "炉号", "订单号", "数量", "单位", "库位", "录入人"])
    for r in data["items"]:
        ws.append(
            [
                r.get("warehouse"),
                r.get("item_code"),
                r.get("name_spec"),
                r.get("material_grade"),
                r.get("heat_no"),
                r.get("order_no"),
                r.get("qty_on_hand"),
                r.get("uom"),
                r.get("location_name"),
                r.get("created_by_name"),
            ]
        )
    file_name = f"成品库存_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return _excel_response(wb, file_name, user, "fg_balances", {"warehouse_id": warehouse_id})


@router.get("/fg/txns")
def list_fg_txns(_user: UserDep):
    rows = fetch_all(
        """
        SELECT t.*, w.name AS warehouse, i.name_spec, i.item_code
        FROM wh_fg_txns t
        JOIN wh_fg_warehouses w ON w.id = t.warehouse_id
        JOIN wh_fg_items i ON i.id = t.item_id
        ORDER BY t.biz_date DESC, t.id DESC LIMIT 200
        """
    )
    return {"items": serialize_rows(rows)}


@router.post("/fg/txns")
def create_fg_txn(req: FgTxnCreate, user: UserDep):
    wh = fetch_one("SELECT * FROM wh_fg_warehouses WHERE id=%s", (req.warehouse_id,))
    item = fetch_one("SELECT * FROM wh_fg_items WHERE id=%s AND is_active=1", (req.item_id,))
    if not wh or not item:
        raise HTTPException(status_code=400, detail="仓库或物料不存在")
    qty = _d(req.qty)
    if qty == 0:
        raise HTTPException(status_code=400, detail="数量不能为0")
    # 出库类型数量可为负；若正数且类型为出库则取负
    signed = qty
    if req.txn_type in ("outbound_ship",) and qty > 0:
        signed = -qty
    heat = req.heat_no or ""
    order_no = req.order_no or ""
    c_uid, c_name, u_uid, u_name = audit_tuple(user)
    txn_no = gen_no("FG")
    with get_db() as conn:
        cur = conn.cursor(dictionary=True)
        cur.execute(
            """
            SELECT id, qty_on_hand FROM wh_fg_balances
            WHERE warehouse_id=%s AND item_id=%s AND heat_no=%s AND order_no=%s
            """,
            (req.warehouse_id, req.item_id, heat, order_no),
        )
        bal = cur.fetchone()
        if bal:
            new_qty = _d(bal["qty_on_hand"]) + signed
            if new_qty < 0:
                cur.close()
                raise HTTPException(status_code=400, detail="成品库存不足")
            cur.execute(
                """
                UPDATE wh_fg_balances SET qty_on_hand=%s, location_name=COALESCE(%s, location_name),
                  updated_by_user_id=%s, updated_by_name=%s WHERE id=%s
                """,
                (new_qty, req.location_name, u_uid, u_name, bal["id"]),
            )
        else:
            if signed < 0:
                cur.close()
                raise HTTPException(status_code=400, detail="成品库存不足")
            cur.execute(
                """
                INSERT INTO wh_fg_balances
                  (warehouse_id, item_id, heat_no, order_no, location_name, qty_on_hand,
                   created_by_user_id, created_by_name, updated_by_user_id, updated_by_name)
                VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
                """,
                (
                    req.warehouse_id,
                    req.item_id,
                    heat,
                    order_no,
                    req.location_name,
                    signed,
                    c_uid,
                    c_name,
                    u_uid,
                    u_name,
                ),
            )
        cur.execute(
            """
            INSERT INTO wh_fg_txns
              (txn_no, txn_type, biz_date, warehouse_id, item_id, heat_no, order_no, qty, operator_name, remark,
               created_by_user_id, created_by_name, updated_by_user_id, updated_by_name)
            VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
            """,
            (
                txn_no,
                req.txn_type,
                req.biz_date,
                req.warehouse_id,
                req.item_id,
                heat,
                order_no,
                signed,
                req.operator_name or user.name,
                req.remark,
                c_uid,
                c_name,
                u_uid,
                u_name,
            ),
        )
        txn_id = int(cur.lastrowid)
        cur.close()
    write_op_log(
        domain="fg",
        action="create",
        entity_type="fg_txn",
        entity_id=txn_id,
        entity_no=txn_no,
        user=user,
        change_summary={"op": "create", "txn_type": req.txn_type, "qty": float(signed)},
    )
    return {"txn_no": txn_no, "id": txn_id}


@router.get("/fg/alerts")
def list_fg_alerts(_user: UserDep):
    # Phase1：低于 5 件的成品作为预警
    rows = fetch_all(
        """
        SELECT b.id, w.name AS warehouse, i.name_spec, i.item_code, b.qty_on_hand, b.heat_no, b.order_no
        FROM wh_fg_balances b
        JOIN wh_fg_warehouses w ON w.id=b.warehouse_id
        JOIN wh_fg_items i ON i.id=b.item_id
        WHERE b.qty_on_hand < 5
        ORDER BY b.qty_on_hand ASC
        """
    )
    alerts = []
    for r in serialize_rows(rows):
        alerts.append(
            {
                "id": r["id"],
                "level": "warning",
                "title": f"{r['warehouse']} {r['name_spec']} 库存偏低（{r['qty_on_hand']}）",
                "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            }
        )
    return {"items": alerts}


# ---------- logs ----------
@router.get("/logs")
def list_logs(
    _user: UserDep,
    domain: str = "",
    entity_type: str = "",
    entity_id: Optional[int] = None,
    limit: int = Query(100, ge=1, le=500),
):
    sql = "SELECT * FROM wh_operation_logs WHERE 1=1"
    params: list[Any] = []
    if domain:
        sql += " AND domain=%s"
        params.append(domain)
    if entity_type:
        sql += " AND entity_type=%s"
        params.append(entity_type)
    if entity_id is not None:
        sql += " AND entity_id=%s"
        params.append(entity_id)
    sql += " ORDER BY id DESC LIMIT %s"
    params.append(limit)
    return {"items": serialize_rows(fetch_all(sql, tuple(params)))}
