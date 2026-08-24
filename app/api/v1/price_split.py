"""价格拆分表建档 API：法国 / 大连分表，Excel 导入与逐条 CRUD。"""
from __future__ import annotations

import io
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Any
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from typing_extensions import Annotated

from app.api.deps import CurrentUser, get_current_user
from app.core.database import execute, fetch_all, fetch_one
from app.core.paths import get_storage_dir
from app.services.price_split_audit import diff_fields, snapshot_keys, write_op_log
from app.services.price_split_calc import apply_derived
from app.services.price_split_parser import parse_price_split_workbook
from app.services.price_split_schema import (
    ALL_DATA_COLUMNS,
    AUDIT_COLUMN_ALTERS,
    CREATE_LOG_TABLE_SQL,
    CREATE_TABLE_SQL,
    INT_FIELDS,
    MATERIAL_GROUPS,
    NUMERIC_FIELDS,
    PRICE_COMPARE_FIELDS,
    REGION_TABLES,
    SELECT_SQL,
    STRING_FIELDS,
    UNIQUE_INDEX_SQL,
    VAT_COLUMN_ALTERS,
    VERSION_COLUMN_ALTERS,
    identity_clause,
    identity_values,
    sanitize_inactive_fields,
)

router = APIRouter(prefix="/price-split", tags=["价格拆分表建档"])

UserDep = Annotated[CurrentUser, Depends(get_current_user)]

_TABLES_READY = False


def _now_compact() -> str:
    return datetime.now().strftime("%Y%m%d%H%M%S")


def _normalize_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _safe_float(value: Any, default: float = 0.0) -> float:
    if value is None or value == "":
        return default
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _safe_int(value: Any, default: int | None = 0) -> int | None:
    if value is None or value == "":
        return default
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return default


def _jsonable(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    return value


def _column_exists(table: str, column: str) -> bool:
    row = fetch_one(
        """
        SELECT COUNT(*) AS n FROM information_schema.COLUMNS
        WHERE TABLE_SCHEMA=DATABASE() AND TABLE_NAME=%s AND COLUMN_NAME=%s
        """,
        (table, column),
    )
    return int((row or {}).get("n") or 0) > 0


def _index_exists(table: str, index_name: str) -> bool:
    row = fetch_one(
        """
        SELECT COUNT(*) AS n FROM information_schema.STATISTICS
        WHERE TABLE_SCHEMA=DATABASE() AND TABLE_NAME=%s AND INDEX_NAME=%s
        """,
        (table, index_name),
    )
    return int((row or {}).get("n") or 0) > 0


def _ensure_tables() -> None:
    global _TABLES_READY
    if _TABLES_READY:
        return
    execute(CREATE_TABLE_SQL["france"])
    execute(CREATE_TABLE_SQL["dalian"])
    execute(CREATE_LOG_TABLE_SQL)
    for region, table in REGION_TABLES.items():
        for column, stmt in (*AUDIT_COLUMN_ALTERS, *VERSION_COLUMN_ALTERS, *VAT_COLUMN_ALTERS):
            if not _column_exists(table, column):
                execute(stmt.format(table=table))
        if _column_exists(table, "valid_slot"):
            execute(f"UPDATE {table} SET valid_slot=1 WHERE is_valid=1 AND valid_slot IS NULL")
        if _index_exists(table, "uk_group_material"):
            execute(f"ALTER TABLE {table} DROP INDEX uk_group_material")
        if not _index_exists(table, "uk_valid_identity"):
            execute(UNIQUE_INDEX_SQL[region])
        if not _index_exists(table, "idx_is_valid"):
            execute(f"ALTER TABLE {table} ADD INDEX idx_is_valid (is_valid)")
    _TABLES_READY = True


def _table(region: str) -> str:
    key = (region or "").strip().lower()
    if key not in REGION_TABLES:
        raise HTTPException(status_code=400, detail="区域仅支持 france 或 dalian")
    return REGION_TABLES[key]


def _normalize_group(value: Any) -> str:
    text = _normalize_str(value).upper()
    if text in ("304", "316", "304&316", "SS", "不锈钢"):
        return "SS"
    return "A105"


def _row_from_db(row: dict) -> dict:
    out = {"id": int(row["id"])}
    for key in ALL_DATA_COLUMNS:
        out[key] = _jsonable(row.get(key))
        if key in STRING_FIELDS:
            out[key] = _normalize_str(out[key])
        elif key in INT_FIELDS:
            raw = out[key]
            out[key] = None if raw in (None, "") else int(raw)
        elif key in ("excel_net_price", "vat_rate"):
            raw = out[key]
            out[key] = None if raw in (None, "") else _safe_float(raw, 0.0)
        elif key in NUMERIC_FIELDS:
            out[key] = _safe_float(out[key], 0.0)
    out["created_by"] = _normalize_str(row.get("created_by"))
    out["created_by_user_id"] = _safe_int(row.get("created_by_user_id"), None)
    out["created_at"] = _jsonable(row.get("created_at")) or ""
    out["updated_by"] = _normalize_str(row.get("updated_by"))
    out["updated_by_user_id"] = _safe_int(row.get("updated_by_user_id"), None)
    out["updated_at"] = _jsonable(row.get("updated_at")) or ""
    out["is_valid"] = int(row.get("is_valid") or 0)
    out["valid_slot"] = _safe_int(row.get("valid_slot"), None)
    out["superseded_by_id"] = _safe_int(row.get("superseded_by_id"), None)
    return out


def _payload_from_raw(raw: dict, region: str, *, recalc: bool = True) -> dict[str, Any]:
    item: dict[str, Any] = {}
    for key in ALL_DATA_COLUMNS:
        if key in STRING_FIELDS:
            item[key] = _normalize_str(raw.get(key))
        elif key in INT_FIELDS:
            item[key] = _safe_int(raw.get(key), None if key in ("source_row", "replaces_id", "superseded_by_id") else 0)
        elif key in ("excel_net_price", "vat_rate"):
            raw_val = raw.get(key)
            item[key] = None if raw_val in (None, "") else _safe_float(raw_val, 0.0)
        else:
            item[key] = _safe_float(raw.get(key), 0.0)
    item["material_group"] = _normalize_group(item.get("material_group") or raw.get("material_group"))
    if item["material_group"] not in MATERIAL_GROUPS:
        item["material_group"] = "A105"
    if not item.get("profit_rate"):
        item["profit_rate"] = 1.07
    sanitize_inactive_fields(item, region)
    if recalc:
        apply_derived(item, region)
    return item


def _insert_values(item: dict, user: CurrentUser) -> tuple[Any, ...]:
    values = [item.get(key) for key in ALL_DATA_COLUMNS]
    values.extend((user.name, user.id, user.name, user.id, 1, 1))
    return tuple(values)


def _update_row(table: str, row_id: int, item: dict, user: CurrentUser) -> None:
    assignments = ", ".join(f"{col}=%s" for col in ALL_DATA_COLUMNS)
    execute(
        f"""
        UPDATE {table}
        SET {assignments}, updated_by=%s, updated_by_user_id=%s, updated_at=NOW()
        WHERE id=%s
        """,
        (*tuple(item.get(col) for col in ALL_DATA_COLUMNS), user.name, user.id, row_id),
    )


def _insert_valid_row(table: str, item: dict, user: CurrentUser) -> int:
    placeholders = ", ".join(["%s"] * (len(ALL_DATA_COLUMNS) + 6))
    columns = ", ".join(
        (*ALL_DATA_COLUMNS, "created_by", "created_by_user_id", "updated_by", "updated_by_user_id", "is_valid", "valid_slot")
    )
    return execute(
        f"INSERT INTO {table} ({columns}) VALUES ({placeholders})",
        _insert_values(item, user),
    )


def _price_changed(before: dict, after: dict) -> bool:
    for key in PRICE_COMPARE_FIELDS:
        if abs(_safe_float(before.get(key)) - _safe_float(after.get(key))) > 1e-6:
            return True
    return False


def _find_valid_row(table: str, region: str, item: dict) -> dict | None:
    return fetch_one(
        f"SELECT {SELECT_SQL} FROM {table} WHERE {identity_clause(region)} LIMIT 1",
        identity_values(region, item),
    )


def _invalidate_row(table: str, old_id: int, new_id: int, user: CurrentUser) -> None:
    execute(
        f"""
        UPDATE {table}
        SET is_valid=0, valid_slot=NULL, superseded_by_id=%s,
            updated_by=%s, updated_by_user_id=%s, updated_at=NOW()
        WHERE id=%s
        """,
        (new_id, user.name, user.id, old_id),
    )


def _save_versioned_row(
    table: str,
    region: str,
    item: dict,
    user: CurrentUser,
    *,
    remark: str,
    source_id: int | None = None,
) -> tuple[int, str, dict | None]:
    if not _normalize_str(item.get("drawing_no")):
        raise HTTPException(status_code=400, detail="图纸号不能为空")
    if not _normalize_str(item.get("material_no")):
        raise HTTPException(status_code=400, detail="物料号不能为空")
    if region == "france" and item.get("vat_rate") in (None, ""):
        raise HTTPException(status_code=400, detail="法国增值税为必填项")
    version = _normalize_str(item.get("price_version"))
    if not version:
        raise HTTPException(status_code=400, detail="价格版本不能为空")
    item["price_version"] = version
    remark = _normalize_str(remark)

    existing = None
    if source_id:
        existing = fetch_one(f"SELECT {SELECT_SQL} FROM {table} WHERE id=%s AND is_valid=1", (int(source_id),))
    if not existing:
        existing = _find_valid_row(table, region, item)

    if not existing:
        item["price_change_remark"] = remark or "首次建档"
        item["replaces_id"] = None
        new_id = _insert_valid_row(table, item, user)
        return int(new_id), "inserted", None

    before = _row_from_db(existing)
    old_id = int(existing["id"])
    if not _price_changed(before, item):
        item["price_change_remark"] = remark or _normalize_str(before.get("price_change_remark"))
        item["replaces_id"] = before.get("replaces_id")
        _update_row(table, old_id, item, user)
        return old_id, "updated", before

    if version == _normalize_str(before.get("price_version")):
        raise HTTPException(
            status_code=400,
            detail=f"价格已变动，请先修改价格版本后再保存（物料号 {item['material_no']}，当前版本 {version}）",
        )
    if not remark:
        raise HTTPException(
            status_code=400,
            detail=f"价格已变动，请填写价格变化记录说明（物料号 {item['material_no']}）",
        )
    item["price_change_remark"] = remark
    item["replaces_id"] = old_id
    new_id = _insert_valid_row(table, item, user)
    _invalidate_row(table, old_id, int(new_id), user)
    return int(new_id), "adjusted", before


def _log_row_change(
    region: str,
    action: str,
    user: CurrentUser,
    row_id: int | None,
    item: dict,
    *,
    before: dict | None = None,
    extra: dict | None = None,
) -> None:
    summary: dict[str, Any] = {"op": action, **snapshot_keys(item)}
    if action in ("update", "adjust") and before:
        summary["fields"] = diff_fields(before, item)
        summary["history_net_price"] = before.get("net_price")
        summary["history_price_version"] = before.get("price_version")
        summary["history_id"] = before.get("id")
    if extra:
        summary.update(extra)
    write_op_log(
        region=region,
        action=action,
        user=user,
        row_id=row_id,
        material_group=_normalize_str(item.get("material_group")),
        material_no=_normalize_str(item.get("material_no")),
        change_summary=summary,
    )


def _sync_materials(region: str, item: dict) -> dict[str, Any]:
    field = "france_agreement_price" if region == "france" else "dalian_agreement_price"
    material_no = _normalize_str(item.get("material_no"))
    drawing_no = _normalize_str(item.get("drawing_no"))
    drawing_code = _normalize_str(item.get("drawing_code"))
    net_price = round(_safe_float(item.get("net_price")), 2)
    if not material_no:
        return {"ok": False, "reason": "物料号为空"}

    matched = fetch_one(
        """
        SELECT id, drawing_no FROM materials
        WHERE material_no=%s AND COALESCE(drawing_no, '')=%s
        ORDER BY id LIMIT 1
        """,
        (material_no, drawing_no),
    )
    if not matched and drawing_code:
        matched = fetch_one(
            """
            SELECT id, drawing_no FROM materials
            WHERE material_no=%s AND COALESCE(drawing_no, '')=%s
            ORDER BY id LIMIT 1
            """,
            (material_no, drawing_code),
        )
    if not matched:
        rows = fetch_all("SELECT id, drawing_no FROM materials WHERE material_no=%s ORDER BY id", (material_no,))
        if len(rows) == 1:
            matched = rows[0]
    if not matched:
        return {"ok": False, "reason": "未找到对应物料建档", "material_no": material_no, "drawing_no": drawing_no}

    execute(f"UPDATE materials SET {field}=%s, updated_at=NOW() WHERE id=%s", (net_price, int(matched["id"])))
    return {"ok": True, "material_id": int(matched["id"]), "field": field, "net_price": net_price}


class PriceSplitSaveRequest(BaseModel):
    rows: list[dict[str, Any]] = Field(default_factory=list)
    sync_materials: bool = False
    price_version: str = ""
    price_change_remark: str = ""


class PriceSplitRecalcRequest(BaseModel):
    rows: list[dict[str, Any]] = Field(default_factory=list)


class PriceSplitDeleteBatchRequest(BaseModel):
    ids: list[int] = Field(default_factory=list)


@router.get("/{region}/logs")
def list_price_split_logs(
    region: str,
    _user: UserDep,
    material_group: str = "",
    material_no: str = "",
    row_id: int | None = None,
    action: str = "",
    limit: int = Query(200, ge=1, le=500),
):
    _ensure_tables()
    _table(region)
    clauses = ["region=%s"]
    params: list[Any] = [region]
    if _normalize_str(material_group):
        clauses.append("material_group=%s")
        params.append(_normalize_group(material_group))
    if _normalize_str(material_no):
        clauses.append("material_no LIKE %s")
        params.append(f"%{_normalize_str(material_no)}%")
    if row_id:
        clauses.append("row_id=%s")
        params.append(int(row_id))
    if _normalize_str(action):
        clauses.append("action=%s")
        params.append(_normalize_str(action))
    rows = fetch_all(
        f"""
        SELECT id, region, row_id, material_group, material_no, action,
               operator_user_id, operator_name, operated_at, change_summary
        FROM price_split_operation_logs
        WHERE {' AND '.join(clauses)}
        ORDER BY id DESC
        LIMIT %s
        """,
        (*params, limit),
    )
    data = []
    for row in rows:
        item = {k: _jsonable(v) for k, v in row.items()}
        item["operator_name"] = _normalize_str(item.get("operator_name"))
        item["change_summary"] = item.get("change_summary") or ""
        data.append(item)
    return {"status": "success", "region": region, "total": len(data), "rows": data}


@router.get("/{region}")
def list_price_split(
    region: str,
    _user: UserDep,
    material_group: str = "",
    drawing_no: str = "",
    material_no: str = "",
    part_no: str = "",
):
    _ensure_tables()
    table = _table(region)
    clauses = ["is_valid=1"]
    params: list[Any] = []
    group = _normalize_str(material_group)
    if group:
        clauses.append("material_group=%s")
        params.append(_normalize_group(group))
    if _normalize_str(drawing_no):
        clauses.append("drawing_no LIKE %s")
        params.append(f"%{_normalize_str(drawing_no)}%")
    if _normalize_str(material_no):
        clauses.append("material_no LIKE %s")
        params.append(f"%{_normalize_str(material_no)}%")
    if _normalize_str(part_no):
        clauses.append("part_no LIKE %s")
        params.append(f"%{_normalize_str(part_no)}%")
    rows = fetch_all(
        f"SELECT {SELECT_SQL} FROM {table} WHERE {' AND '.join(clauses)} ORDER BY material_group, id",
        tuple(params),
    )
    data = [_row_from_db(r) for r in rows]
    return {"status": "success", "region": region, "total": len(data), "rows": data}


@router.get("/{region}/lookup")
def lookup_price_split(
    region: str,
    _user: UserDep,
    material_no: str = Query(...),
    drawing_no: str = "",
    material_group: str = "",
    part_no: str = "",
):
    _ensure_tables()
    table = _table(region)
    clauses = ["material_no=%s", "is_valid=1"]
    params: list[Any] = [_normalize_str(material_no)]
    if _normalize_str(material_group):
        clauses.append("material_group=%s")
        params.append(_normalize_group(material_group))
    if _normalize_str(drawing_no):
        clauses.append("(drawing_no=%s OR drawing_code=%s)")
        params.extend([_normalize_str(drawing_no), _normalize_str(drawing_no)])
    if region == "dalian" and _normalize_str(part_no):
        clauses.append("COALESCE(part_no,'')=%s")
        params.append(_normalize_str(part_no))
    row = fetch_one(
        f"SELECT {SELECT_SQL} FROM {table} WHERE {' AND '.join(clauses)} ORDER BY id LIMIT 1",
        tuple(params),
    )
    if not row:
        return {"status": "success", "found": False, "row": None}
    data = _row_from_db(row)
    return {"status": "success", "found": True, "row": data, "net_price": data.get("net_price")}


@router.get("/{region}/export")
def export_price_split(region: str, user: UserDep, material_group: str = Query(...)):
    _ensure_tables()
    table = _table(region)
    group = _normalize_group(material_group)
    rows = fetch_all(
        f"SELECT {SELECT_SQL} FROM {table} WHERE material_group=%s AND is_valid=1 ORDER BY id",
        (group,),
    )
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise HTTPException(status_code=500, detail="缺少 openpyxl 依赖，无法导出 Excel") from exc
    headers = [
        "ID",
        "图纸号",
        "物料号",
        "AP1",
        "材质",
        "类型",
        "DN",
        "描述",
        "钢材价格",
        "下料毛重",
        "净重",
        "材料成本",
        "总加工成本",
        "最终报价",
        "增值税",
        "含税单价",
        "价格版本",
        "调价说明",
        "上传者",
        "上传时间",
        "更改者",
        "更改时间",
    ]
    keys = [
        "id",
        "drawing_no",
        "material_no",
        "part_no",
        "material",
        "flange_type",
        "dn",
        "description",
        "steel_price",
        "blanking_weight",
        "net_weight",
        "material_cost",
        "total_process_cost",
        "net_price",
        "vat_rate",
        "tax_inclusive_price",
        "price_version",
        "price_change_remark",
        "created_by",
        "created_at",
        "updated_by",
        "updated_at",
    ]
    wb = Workbook()
    ws = wb.active
    ws.title = f"{region}-{group}"
    ws.append(headers)
    for raw in rows:
        item = _row_from_db(raw)
        ws.append([item.get(key) for key in keys])
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    region_label = "法国" if region == "france" else "大连"
    group_label = "不锈钢" if group == "SS" else "A105"
    file_name = f"价格拆分_{region_label}_{group_label}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    write_op_log(
        region=region,
        action="export",
        user=user,
        material_group=group,
        change_summary={"op": "export", "material_group": group, "file_name": file_name, "count": len(rows)},
    )
    return StreamingResponse(
        bio,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(file_name)}"},
    )


def _save_uploaded_xlsx(file: UploadFile) -> Path:
    filename = (file.filename or "").lower()
    if not filename.endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="仅支持 xlsx 文件")
    temp_dir = get_storage_dir() / "temp"
    temp_dir.mkdir(parents=True, exist_ok=True)
    temp_file = temp_dir / f"price_split_{_now_compact()}.xlsx"
    return temp_file


@router.post("/{region}/preview")
async def preview_price_split(region: str, _user: UserDep, file: UploadFile = File(...)):
    _ensure_tables()
    _table(region)
    temp_file = _save_uploaded_xlsx(file)
    try:
        content = await file.read()
        temp_file.write_bytes(content)
        parsed = parse_price_split_workbook(temp_file, region)
    except HTTPException:
        raise
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"文件读取失败: {exc}") from exc
    finally:
        if temp_file.exists():
            temp_file.unlink()
    return {"status": "success", "region": region, **parsed}


@router.post("/{region}/import")
async def import_price_split(
    region: str,
    user: UserDep,
    file: UploadFile = File(...),
    price_version: str = Form(""),
    price_change_remark: str = Form(""),
    vat_rate: str = Form(""),
    sync_materials: str = Form("false"),
):
    _ensure_tables()
    table = _table(region)
    do_sync = str(sync_materials).strip().lower() in ("1", "true", "yes", "on")
    temp_file = _save_uploaded_xlsx(file)
    source_name = file.filename or temp_file.name
    try:
        content = await file.read()
        temp_file.write_bytes(content)
        parsed = parse_price_split_workbook(temp_file, region)
    except HTTPException:
        raise
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"文件读取失败: {exc}") from exc
    finally:
        if temp_file.exists():
            temp_file.unlink()

    version = _normalize_str(price_version)
    if not version:
        raise HTTPException(status_code=400, detail="价格版本不能为空")
    remark = _normalize_str(price_change_remark)
    vat_value = None if not _normalize_str(vat_rate) else _safe_float(vat_rate)
    if region == "france" and vat_value is None:
        raise HTTPException(status_code=400, detail="法国增值税为必填项，请填写如 13（代表 13%）")
    inserted = 0
    updated = 0
    adjusted = 0
    synced = 0
    sync_missing: list[dict] = []
    saved_ids: list[int] = []
    for raw in (*parsed["rows_a105"], *parsed["rows_ss"]):
        raw["price_version"] = version
        item = _payload_from_raw(raw, region, recalc=False)
        if not item["material_no"]:
            continue
        if region == "france" and item.get("vat_rate") in (None, "") and vat_value is not None:
            item["vat_rate"] = vat_value
            apply_derived(item, region)
        row_id, action, before = _save_versioned_row(table, region, item, user, remark=remark)
        saved_ids.append(row_id)
        if action == "inserted":
            inserted += 1
        elif action == "adjusted":
            adjusted += 1
        else:
            updated += 1
        log_action = "create" if action == "inserted" else ("adjust" if action == "adjusted" else "update")
        _log_row_change(region, log_action, user, row_id, item, before=before, extra={"price_change_remark": remark or item.get("price_change_remark")})
        if do_sync:
            result = _sync_materials(region, item)
            if result.get("ok"):
                synced += 1
            else:
                sync_missing.append(result)

    write_op_log(
        region=region,
        action="upload",
        user=user,
        change_summary={
            "op": "upload",
            "file_name": source_name,
            "count_a105": parsed.get("count_a105", 0),
            "count_ss": parsed.get("count_ss", 0),
            "inserted_count": inserted,
            "updated_count": updated,
            "adjusted_count": adjusted,
            "price_version": version,
            "price_change_remark": remark,
            "sync_materials": do_sync,
        },
    )

    return {
        "status": "success",
        "region": region,
        "inserted_count": inserted,
        "updated_count": updated,
        "adjusted_count": adjusted,
        "saved_count": inserted + updated + adjusted,
        "synced_count": synced,
        "sync_missing": sync_missing[:50],
        "warning_count": parsed.get("warning_count", 0),
        "warnings": parsed.get("warnings", [])[:50],
        "count_a105": parsed.get("count_a105", 0),
        "count_ss": parsed.get("count_ss", 0),
        "sheets": parsed.get("sheets", []),
    }


@router.post("/{region}/rows")
def save_price_split_rows(region: str, req: PriceSplitSaveRequest, user: UserDep):
    _ensure_tables()
    table = _table(region)
    inserted = 0
    updated = 0
    adjusted = 0
    synced = 0
    sync_missing: list[dict] = []
    saved_ids: list[int] = []
    remark = _normalize_str(req.price_change_remark)
    for raw in req.rows:
        material_no = _normalize_str(raw.get("material_no"))
        if not material_no:
            raise HTTPException(status_code=400, detail="物料号不能为空")
        if req.price_version:
            raw["price_version"] = req.price_version
        item = _payload_from_raw(raw, region, recalc=True)
        row_id, action, before = _save_versioned_row(
            table,
            region,
            item,
            user,
            remark=remark or _normalize_str(raw.get("price_change_remark")),
            source_id=int(raw["id"]) if raw.get("id") else None,
        )
        saved_ids.append(row_id)
        if action == "inserted":
            inserted += 1
        elif action == "adjusted":
            adjusted += 1
        else:
            updated += 1
        log_action = "create" if action == "inserted" else ("adjust" if action == "adjusted" else "update")
        _log_row_change(
            region,
            log_action,
            user,
            row_id,
            item,
            before=before,
            extra={"price_change_remark": item.get("price_change_remark")},
        )
        if req.sync_materials:
            result = _sync_materials(region, item)
            if result.get("ok"):
                synced += 1
            else:
                sync_missing.append(result)

    rows = []
    if saved_ids:
        placeholders = ", ".join(["%s"] * len(saved_ids))
        fetched = fetch_all(
            f"SELECT {SELECT_SQL} FROM {table} WHERE id IN ({placeholders}) ORDER BY id",
            tuple(saved_ids),
        )
        rows = [_row_from_db(r) for r in fetched]
    return {
        "status": "success",
        "saved_count": len(saved_ids),
        "inserted_count": inserted,
        "updated_count": updated,
        "adjusted_count": adjusted,
        "synced_count": synced,
        "sync_missing": sync_missing[:50],
        "rows": rows,
    }


@router.post("/{region}/recalc")
def recalc_price_split(region: str, req: PriceSplitRecalcRequest, _user: UserDep):
    _table(region)
    rows = []
    for raw in req.rows:
        item = _payload_from_raw(raw, region, recalc=True)
        if raw.get("id"):
            item["id"] = int(raw["id"])
        rows.append(item)
    return {"status": "success", "rows": rows}


@router.post("/{region}/delete-batch")
def delete_price_split_batch(region: str, req: PriceSplitDeleteBatchRequest, user: UserDep):
    _ensure_tables()
    table = _table(region)
    ids = sorted({int(item) for item in req.ids if int(item) > 0})
    if not ids:
        raise HTTPException(status_code=400, detail="请选择要删除的记录")
    placeholders = ", ".join(["%s"] * len(ids))
    existing = fetch_all(f"SELECT {SELECT_SQL} FROM {table} WHERE id IN ({placeholders})", tuple(ids))
    found_ids = [int(row["id"]) for row in existing]
    for row in existing:
        item = _row_from_db(row)
        _log_row_change(region, "delete", user, int(row["id"]), item, extra={"snapshot": snapshot_keys(item)})
    if found_ids:
        found_placeholders = ", ".join(["%s"] * len(found_ids))
        execute(f"DELETE FROM {table} WHERE id IN ({found_placeholders})", tuple(found_ids))
    return {
        "status": "success",
        "deleted_count": len(found_ids),
        "missing_count": len(ids) - len(found_ids),
    }


@router.post("/{region}/delete-all")
def delete_price_split_all(region: str, user: UserDep, material_group: str = Query(...)):
    _ensure_tables()
    table = _table(region)
    group = _normalize_group(material_group)
    existing = fetch_all(
        f"SELECT id, material_group, material_no, drawing_no, net_price FROM {table} WHERE material_group=%s AND is_valid=1",
        (group,),
    )
    total = len(existing)
    if total:
        execute(f"DELETE FROM {table} WHERE material_group=%s AND is_valid=1", (group,))
    write_op_log(
        region=region,
        action="delete",
        user=user,
        material_group=group,
        change_summary={
            "op": "delete-all",
            "material_group": group,
            "deleted_count": total,
            "material_nos": [row.get("material_no") for row in existing[:80]],
        },
    )
    return {"status": "success", "deleted_count": total, "material_group": group}


@router.delete("/{region}/{row_id}")
def delete_price_split_row(region: str, row_id: int, user: UserDep):
    _ensure_tables()
    table = _table(region)
    existing = fetch_one(f"SELECT {SELECT_SQL} FROM {table} WHERE id=%s", (row_id,))
    if not existing:
        raise HTTPException(status_code=404, detail="未找到要删除的价格拆分记录")
    item = _row_from_db(existing)
    _log_row_change(region, "delete", user, row_id, item, extra={"snapshot": snapshot_keys(item)})
    execute(f"DELETE FROM {table} WHERE id=%s", (row_id,))
    return {"status": "success", "deleted_id": row_id, "material_no": existing.get("material_no") or ""}
