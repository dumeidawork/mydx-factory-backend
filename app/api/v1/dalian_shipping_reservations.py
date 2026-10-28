"""大连西门子客户发货预定。部分预定会把订单明细拆成本次数量和余量。"""
from __future__ import annotations

import io
import json
from datetime import date, datetime
from typing import Any
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from typing_extensions import Annotated

from app.api.deps import CurrentUser, get_current_user
from app.core.database import execute, fetch_all, fetch_one
from app.services.dalian_order_split import (
    align_order_line_to_quantity,
    assert_planned_qty_fits,
    reservation_is_locked,
)

router = APIRouter(prefix="/partners/dalian-siemens", tags=["合作伙伴-大连西门子发货预定"])
UserDep = Annotated[CurrentUser, Depends(get_current_user)]

_TABLES_READY = False

STATUS_RESERVED = "已预定"
STATUS_ASSIGNED = "已指定对接"
STATUS_CHANGED = "存在变更内容"
STATUS_PACKING = "开始装箱"
STATUS_PACKED = "已装箱"
STATUS_SHIPPED = "已发货"

PROGRESS_STATUSES = (
    STATUS_RESERVED,
    STATUS_ASSIGNED,
    STATUS_CHANGED,
    STATUS_PACKING,
    STATUS_PACKED,
    STATUS_SHIPPED,
)

ORDER_DETAIL_STATUSES = (
    "开始",
    "分割",
    "加工",
    "入库",
    "质检",
    "再加工",
    "包装",
    "装箱",
    "打包",
    "发货",
    "完成",
)

CREATE_RESERVATION_SQL = """
CREATE TABLE IF NOT EXISTS dalian_shipping_reservations (
    id BIGINT AUTO_INCREMENT PRIMARY KEY,
    order_detail_id INT NOT NULL COMMENT '订单明细ID',
    order_no VARCHAR(64) NOT NULL DEFAULT '' COMMENT '订单号快照',
    factory_order_no VARCHAR(128) NOT NULL DEFAULT '' COMMENT '本厂单号快照',
    part_no VARCHAR(128) NOT NULL DEFAULT '' COMMENT '料号快照',
    material_no VARCHAR(128) NOT NULL DEFAULT '' COMMENT '物料号快照',
    spec_model VARCHAR(256) NOT NULL DEFAULT '' COMMENT '规格型号快照',
    material VARCHAR(64) NOT NULL DEFAULT '' COMMENT '材质快照',
    item_no INT NULL COMMENT '条目，来自订单明细，箱单显示',
    drawing_no VARCHAR(255) NOT NULL DEFAULT '' COMMENT '图纸号，来自订单明细，箱单显示',
    unit_weight DECIMAL(12,2) NOT NULL DEFAULT 0.00 COMMENT '单重，来自订单明细，箱单显示',
    total_weight DECIMAL(12,2) NOT NULL DEFAULT 0.00 COMMENT '总重=单重×预定数量，箱单显示',
    order_qty INT NOT NULL DEFAULT 0 COMMENT '保存时订单数量快照',
    planned_arrival_date DATE NOT NULL COMMENT '预定到货日期',
    planned_qty INT NOT NULL DEFAULT 0 COMMENT '预定到货数量，0=该日取消留痕',
    remark VARCHAR(255) NOT NULL DEFAULT '' COMMENT '备注',
    status VARCHAR(32) NOT NULL DEFAULT '已预定' COMMENT '已预定/已指定对接/存在变更内容/开始装箱/已装箱/已发货',
    booker_user_id INT NULL COMMENT '预定者用户ID',
    booker_name VARCHAR(64) NOT NULL DEFAULT '' COMMENT '预定者',
    booked_at DATETIME NULL COMMENT '预定时间',
    assignee_user_id INT NULL COMMENT '对接人员用户ID',
    assignee_name VARCHAR(64) NOT NULL DEFAULT '' COMMENT '对接人员',
    assigned_at DATETIME NULL COMMENT '指定对接时间',
    assigned_by_user_id INT NULL COMMENT '指定人用户ID',
    assigned_by_name VARCHAR(64) NOT NULL DEFAULT '' COMMENT '指定人',
    created_by VARCHAR(64) NULL COMMENT '上传者',
    created_by_user_id INT NULL COMMENT '上传者用户ID',
    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP COMMENT '上传时间',
    updated_by VARCHAR(64) NULL COMMENT '更改者',
    updated_by_user_id INT NULL COMMENT '更改者用户ID',
    updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP COMMENT '更改时间',
    UNIQUE KEY uk_dsr_detail_date (order_detail_id, planned_arrival_date),
    INDEX idx_dsr_date (planned_arrival_date),
    INDEX idx_dsr_order (order_no),
    INDEX idx_dsr_material (material_no),
    INDEX idx_dsr_booker (booker_user_id),
    INDEX idx_dsr_assignee (assignee_user_id),
    INDEX idx_dsr_status (status)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COMMENT='大连西门子客户发货预定'
"""

CREATE_DATE_SQL = """
CREATE TABLE IF NOT EXISTS dalian_shipping_date_assignments (
    planned_arrival_date DATE PRIMARY KEY COMMENT '预定到货日期',
    assignee_user_id INT NULL COMMENT '对接人员用户ID',
    assignee_name VARCHAR(64) NOT NULL DEFAULT '' COMMENT '对接人员',
    assigned_at DATETIME NULL COMMENT '指定时间',
    assigned_by_user_id INT NULL COMMENT '指定人用户ID',
    assigned_by_name VARCHAR(64) NOT NULL DEFAULT '' COMMENT '指定人',
    packing_generated_by VARCHAR(64) NULL COMMENT '箱单最后生成者',
    packing_generated_by_user_id INT NULL COMMENT '箱单最后生成者用户ID',
    packing_generated_at DATETIME NULL COMMENT '箱单最后生成时间',
    packing_content_stale TINYINT NOT NULL DEFAULT 0 COMMENT '生成后预定内容已变，须重新生成才能开始装箱',
    created_by VARCHAR(64) NULL COMMENT '上传者',
    created_by_user_id INT NULL COMMENT '上传者用户ID',
    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP COMMENT '上传时间',
    updated_by VARCHAR(64) NULL COMMENT '更改者',
    updated_by_user_id INT NULL COMMENT '更改者用户ID',
    updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP COMMENT '更改时间'
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COMMENT='大连发货预定按日对接人（删光预定后仍保留）'
"""

CREATE_UNREAD_SQL = """
CREATE TABLE IF NOT EXISTS dalian_shipping_progress_unreads (
    planned_arrival_date DATE NOT NULL,
    user_id INT NOT NULL,
    unread TINYINT NOT NULL DEFAULT 1,
    updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
    PRIMARY KEY (planned_arrival_date, user_id),
    INDEX idx_dsr_unread_user (user_id, unread)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COMMENT='大连发货进度未读（预定者/对接人）'
"""

CREATE_LOG_SQL = """
CREATE TABLE IF NOT EXISTS dalian_shipping_reservation_logs (
    id BIGINT AUTO_INCREMENT PRIMARY KEY,
    row_id BIGINT NULL COMMENT '预定行ID，删除后可空',
    order_detail_id INT NULL,
    order_no VARCHAR(64) NULL,
    material_no VARCHAR(128) NULL,
    planned_arrival_date DATE NULL,
    action VARCHAR(32) NOT NULL COMMENT 'create/update/delete/assign/reassign/generate/regenerate/export/start_packing',
    operator_user_id INT NULL,
    operator_name VARCHAR(64) NOT NULL DEFAULT '',
    operated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
    change_summary TEXT NULL COMMENT '变更内容 JSON',
    INDEX idx_dsr_log_row (row_id),
    INDEX idx_dsr_log_order (order_no),
    INDEX idx_dsr_log_material (material_no),
    INDEX idx_dsr_log_date (planned_arrival_date),
    INDEX idx_dsr_log_action (action),
    INDEX idx_dsr_log_time (operated_at)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COMMENT='大连发货预定操作日志'
"""

RESERVATION_SELECT = """
    id, order_detail_id, order_no, factory_order_no, part_no, material_no, spec_model, material,
    item_no, drawing_no, unit_weight, total_weight,
    order_qty, planned_arrival_date, planned_qty, remark, status,
    booker_user_id, booker_name, booked_at,
    assignee_user_id, assignee_name, assigned_at, assigned_by_user_id, assigned_by_name,
    created_by, created_by_user_id, created_at, updated_by, updated_by_user_id, updated_at
"""


class ReservationUpsertIn(BaseModel):
    order_detail_id: int
    planned_arrival_date: str
    planned_qty: int
    remark: str = ""


class ReservationBatchRequest(BaseModel):
    items: list[ReservationUpsertIn] = Field(default_factory=list)


class ReservationPatchRequest(BaseModel):
    planned_arrival_date: str | None = None
    planned_qty: int | None = None
    remark: str | None = None


class AssignRequest(BaseModel):
    planned_arrival_date: str
    assignee_user_id: int


class PackingDateRequest(BaseModel):
    planned_arrival_date: str


class PackingExportRequest(BaseModel):
    planned_arrival_date: str
    columns: list[str] = Field(default_factory=list)


PACKING_EXPORT_COLUMNS: tuple[tuple[str, str], ...] = (
    ("planned_arrival_date", "预定到货日期"),
    ("order_no", "订单号"),
    ("factory_order_no", "本厂单号"),
    ("item_no", "条目"),
    ("part_no", "料号"),
    ("material_no", "物料号"),
    ("drawing_no", "图纸号"),
    ("spec_model", "规格型号"),
    ("material", "材质"),
    ("heat_no", "炉号"),
    ("heat_treatment_batch_no", "热处理号"),
    ("unit_weight", "单重"),
    ("planned_qty", "预定到货数量"),
    ("total_weight", "总重"),
    ("remark", "备注"),
    ("booker_name", "预定者"),
    ("booked_at", "预定时间"),
    ("assignee_name", "对接人员"),
    ("status", "进度"),
)

_LATER_STATUSES = frozenset({STATUS_PACKING, STATUS_PACKED, STATUS_SHIPPED})


def _column_exists(table: str, column: str) -> bool:
    row = fetch_one(
        """
        SELECT 1 AS n FROM information_schema.columns
        WHERE table_schema = DATABASE() AND table_name=%s AND column_name=%s
        LIMIT 1
        """,
        (table, column),
    )
    return bool(row)


def _ensure_column(table: str, column: str, ddl: str) -> None:
    if not _column_exists(table, column):
        execute(ddl)


def _ensure_tables() -> None:
    global _TABLES_READY
    if _TABLES_READY:
        return
    execute(CREATE_RESERVATION_SQL)
    execute(CREATE_DATE_SQL)
    execute(CREATE_UNREAD_SQL)
    execute(CREATE_LOG_SQL)
    for column, ddl in (
        ("item_no", "ALTER TABLE dalian_shipping_reservations ADD COLUMN item_no INT NULL COMMENT '条目，来自订单明细，箱单显示' AFTER material"),
        ("drawing_no", "ALTER TABLE dalian_shipping_reservations ADD COLUMN drawing_no VARCHAR(255) NOT NULL DEFAULT '' COMMENT '图纸号，来自订单明细，箱单显示' AFTER item_no"),
        ("unit_weight", "ALTER TABLE dalian_shipping_reservations ADD COLUMN unit_weight DECIMAL(12,2) NOT NULL DEFAULT 0.00 COMMENT '单重，来自订单明细，箱单显示' AFTER drawing_no"),
        ("total_weight", "ALTER TABLE dalian_shipping_reservations ADD COLUMN total_weight DECIMAL(12,2) NOT NULL DEFAULT 0.00 COMMENT '总重=单重×预定数量，箱单显示' AFTER unit_weight"),
    ):
        _ensure_column("dalian_shipping_reservations", column, ddl)
    for column, ddl in (
        ("packing_generated_by", "ALTER TABLE dalian_shipping_date_assignments ADD COLUMN packing_generated_by VARCHAR(64) NULL COMMENT '箱单最后生成者' AFTER assigned_by_name"),
        ("packing_generated_by_user_id", "ALTER TABLE dalian_shipping_date_assignments ADD COLUMN packing_generated_by_user_id INT NULL COMMENT '箱单最后生成者用户ID' AFTER packing_generated_by"),
        ("packing_generated_at", "ALTER TABLE dalian_shipping_date_assignments ADD COLUMN packing_generated_at DATETIME NULL COMMENT '箱单最后生成时间' AFTER packing_generated_by_user_id"),
        ("packing_content_stale", "ALTER TABLE dalian_shipping_date_assignments ADD COLUMN packing_content_stale TINYINT NOT NULL DEFAULT 0 COMMENT '生成后预定内容已变，须重新生成才能开始装箱' AFTER packing_generated_at"),
    ):
        _ensure_column("dalian_shipping_date_assignments", column, ddl)
    _TABLES_READY = True


def _jsonable(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    return value


def _norm(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _operator_name(user: CurrentUser) -> str:
    return _norm(user.name) or _norm(user.username) or str(user.id)


def _as_date(value: Any, label: str = "预定到货日期") -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return _parse_date(str(value), label)


def _parse_date(raw: str, label: str = "预定到货日期") -> date:
    text = _norm(raw).replace("/", "-").replace(".", "-")
    if not text:
        raise HTTPException(status_code=400, detail=f"{label}不能为空")
    digits = "".join(ch for ch in text if ch.isdigit())
    try:
        if len(digits) >= 8 and "-" not in text[:10]:
            return date(int(digits[0:4]), int(digits[4:6]), int(digits[6:8]))
        return date.fromisoformat(text[:10])
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=f"{label}格式无效，请使用 YYYY-MM-DD") from exc


def _as_int(value: Any, label: str, *, allow_zero: bool = True) -> int:
    if value is None or value == "":
        raise HTTPException(status_code=400, detail=f"{label}不能为空")
    try:
        number = int(value)
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=400, detail=f"{label}必须是整数") from exc
    if number < 0:
        raise HTTPException(status_code=400, detail=f"{label}不能为负数")
    if number == 0 and not allow_zero:
        raise HTTPException(status_code=400, detail=f"{label}不能为 0")
    return number


def _write_log(
    *,
    action: str,
    user: CurrentUser,
    row_id: int | None,
    order_detail_id: int | None,
    order_no: str,
    material_no: str,
    planned_arrival_date: date | None,
    change_summary: dict[str, Any],
) -> None:
    execute(
        """
        INSERT INTO dalian_shipping_reservation_logs
            (row_id, order_detail_id, order_no, material_no, planned_arrival_date,
             action, operator_user_id, operator_name, change_summary)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
        """,
        (
            row_id,
            order_detail_id,
            order_no,
            material_no,
            planned_arrival_date,
            action,
            user.id,
            _operator_name(user),
            json.dumps(change_summary, ensure_ascii=False, default=str),
        ),
    )


def _reservation_from_db(row: dict[str, Any] | None) -> dict[str, Any]:
    if not row:
        return {}
    return {
        "id": int(row["id"]),
        "order_detail_id": int(row["order_detail_id"]),
        "order_no": _norm(row.get("order_no")),
        "factory_order_no": _norm(row.get("factory_order_no")),
        "part_no": _norm(row.get("part_no")),
        "material_no": _norm(row.get("material_no")),
        "spec_model": _norm(row.get("spec_model")),
        "material": _norm(row.get("material")),
        "item_no": int(row["item_no"]) if row.get("item_no") is not None else None,
        "drawing_no": _norm(row.get("drawing_no")),
        "unit_weight": _weight(row.get("unit_weight")),
        "total_weight": _weight(row.get("total_weight")),
        "order_qty": int(row.get("order_qty") or 0),
        "planned_arrival_date": _jsonable(row.get("planned_arrival_date")) or "",
        "planned_qty": int(row.get("planned_qty") or 0),
        "remark": _norm(row.get("remark")),
        "status": _norm(row.get("status")) or STATUS_RESERVED,
        "booker_user_id": row.get("booker_user_id"),
        "booker_name": _norm(row.get("booker_name")),
        "booked_at": _jsonable(row.get("booked_at")) or "",
        "assignee_user_id": row.get("assignee_user_id"),
        "assignee_name": _norm(row.get("assignee_name")),
        "assigned_at": _jsonable(row.get("assigned_at")) or "",
        "assigned_by_user_id": row.get("assigned_by_user_id"),
        "assigned_by_name": _norm(row.get("assigned_by_name")),
        "created_by": _norm(row.get("created_by")),
        "created_by_user_id": row.get("created_by_user_id"),
        "created_at": _jsonable(row.get("created_at")) or "",
        "updated_by": _norm(row.get("updated_by")),
        "updated_by_user_id": row.get("updated_by_user_id"),
        "updated_at": _jsonable(row.get("updated_at")) or "",
    }


def _load_reservation(row_id: int) -> dict[str, Any]:
    row = fetch_one(f"SELECT {RESERVATION_SELECT} FROM dalian_shipping_reservations WHERE id=%s", (row_id,))
    if not row:
        raise HTTPException(status_code=404, detail="未找到该发货预定")
    return _reservation_from_db(row)


def _load_order_detail(order_detail_id: int) -> dict[str, Any]:
    row = fetch_one(
        """
        SELECT id, order_no, factory_order_no, name, material_no, spec_model, material, quantity,
               item_no, drawing_no, unit_weight
        FROM order_details
        WHERE id=%s
        """,
        (order_detail_id,),
    )
    if not row:
        raise HTTPException(status_code=404, detail="未找到订单明细")
    order_no = _norm(row.get("order_no"))
    if not order_no.startswith("4"):
        raise HTTPException(status_code=400, detail="仅支持大连订单（订单号 4 开头）")
    return {
        "id": int(row["id"]),
        "order_no": order_no,
        "factory_order_no": _norm(row.get("factory_order_no")),
        "part_no": _norm(row.get("name")),
        "material_no": _norm(row.get("material_no")),
        "spec_model": _norm(row.get("spec_model")),
        "material": _norm(row.get("material")),
        "quantity": int(row.get("quantity") or 0),
        "item_no": int(row["item_no"]) if row.get("item_no") is not None else None,
        "drawing_no": _norm(row.get("drawing_no")),
        "unit_weight": _weight(row.get("unit_weight")),
    }


def _reserved_qty(order_detail_id: int, exclude_id: int | None = None) -> int:
    sql = """
        SELECT COALESCE(SUM(planned_qty), 0) AS qty
        FROM dalian_shipping_reservations
        WHERE order_detail_id=%s AND planned_qty > 0
    """
    params: list[Any] = [order_detail_id]
    if exclude_id is not None:
        sql += " AND id<>%s"
        params.append(int(exclude_id))
    row = fetch_one(sql, tuple(params))
    return int((row or {}).get("qty") or 0)


def _assert_remaining(order_qty: int, order_detail_id: int, planned_qty: int, exclude_id: int | None) -> None:
    if planned_qty <= 0:
        return
    reserved = _reserved_qty(order_detail_id, exclude_id=exclude_id)
    remaining = int(order_qty) - reserved
    if planned_qty > remaining:
        raise HTTPException(
            status_code=400,
            detail=f"预定数量 {planned_qty} 超过可预定余量 {max(remaining, 0)}（订单数量 {order_qty}，已预定 {reserved}）",
        )


def _get_date_assignment(day: date) -> dict[str, Any] | None:
    row = fetch_one(
        """
        SELECT planned_arrival_date, assignee_user_id, assignee_name, assigned_at,
               assigned_by_user_id, assigned_by_name,
               packing_generated_by, packing_generated_by_user_id, packing_generated_at,
               packing_content_stale
        FROM dalian_shipping_date_assignments
        WHERE planned_arrival_date=%s
        """,
        (day,),
    )
    return row


def _load_user(user_id: int) -> dict[str, Any]:
    row = fetch_one("SELECT id, name, username FROM users WHERE id=%s", (user_id,))
    if not row:
        raise HTTPException(status_code=404, detail="未找到对接人员")
    return {
        "id": int(row["id"]),
        "name": _norm(row.get("name")) or _norm(row.get("username")),
    }


def _mark_unread(day: date, user_ids: set[int]) -> None:
    for uid in user_ids:
        if not uid:
            continue
        execute(
            """
            INSERT INTO dalian_shipping_progress_unreads (planned_arrival_date, user_id, unread)
            VALUES (%s, %s, 1)
            ON DUPLICATE KEY UPDATE unread=1
            """,
            (day, int(uid)),
        )


def mark_reservations_packed(reservation_ids: list[int], user: CurrentUser) -> int:
    """把唛头用到的预定行改为已装箱。已发货的不改。涉及的到货日标未读。"""
    _ensure_tables()
    ids = sorted({int(item) for item in reservation_ids if item})
    if not ids:
        return 0
    placeholders = ", ".join(["%s"] * len(ids))
    rows = fetch_all(
        f"""
        SELECT id, order_detail_id, order_no, material_no, planned_arrival_date, status
        FROM dalian_shipping_reservations
        WHERE id IN ({placeholders})
        """,
        tuple(ids),
    )
    operator = _operator_name(user)
    updated = 0
    notify_days: set[date] = set()
    for row in rows:
        status = _norm(row.get("status"))
        if status == STATUS_SHIPPED:
            continue
        day = _as_date(row.get("planned_arrival_date")) if row.get("planned_arrival_date") else None
        if day:
            notify_days.add(day)
        if status == STATUS_PACKED:
            continue
        execute(
            """
            UPDATE dalian_shipping_reservations
            SET status=%s, updated_by=%s, updated_by_user_id=%s, updated_at=NOW()
            WHERE id=%s AND status<>%s
            """,
            (STATUS_PACKED, operator, user.id, int(row["id"]), STATUS_SHIPPED),
        )
        _write_log(
            action="update",
            user=user,
            row_id=int(row["id"]),
            order_detail_id=int(row["order_detail_id"]) if row.get("order_detail_id") else None,
            order_no=_norm(row.get("order_no")),
            material_no=_norm(row.get("material_no")),
            planned_arrival_date=day,
            change_summary={"op": "packed", "fields": {"status": {"before": status, "after": STATUS_PACKED}}},
        )
        updated += 1
    for day in notify_days:
        _notify_date(day, extra_user_ids={int(user.id)})
    return updated


def _notify_date(day: date, extra_user_ids: set[int] | None = None) -> None:
    ids: set[int] = set(extra_user_ids or set())
    assignment = _get_date_assignment(day)
    if assignment and assignment.get("assignee_user_id"):
        ids.add(int(assignment["assignee_user_id"]))
    rows = fetch_all(
        """
        SELECT DISTINCT booker_user_id
        FROM dalian_shipping_reservations
        WHERE planned_arrival_date=%s AND booker_user_id IS NOT NULL
        """,
        (day,),
    )
    for row in rows:
        if row.get("booker_user_id"):
            ids.add(int(row["booker_user_id"]))
    _mark_unread(day, ids)


def _inherit_assignee(day: date) -> dict[str, Any]:
    assignment = _get_date_assignment(day)
    if not assignment or not assignment.get("assignee_user_id"):
        return {
            "assignee_user_id": None,
            "assignee_name": "",
            "assigned_at": None,
            "assigned_by_user_id": None,
            "assigned_by_name": "",
            "has_assignee": False,
        }
    return {
        "assignee_user_id": int(assignment["assignee_user_id"]),
        "assignee_name": _norm(assignment.get("assignee_name")),
        "assigned_at": assignment.get("assigned_at"),
        "assigned_by_user_id": assignment.get("assigned_by_user_id"),
        "assigned_by_name": _norm(assignment.get("assigned_by_name")),
        "has_assignee": True,
    }


def _weight(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    return round(float(value), 2)


def _snapshot_weights(detail: dict[str, Any], qty: int) -> tuple[Any, str, float, float]:
    unit = _weight(detail.get("unit_weight"))
    return detail.get("item_no"), _norm(detail.get("drawing_no")), unit, round(unit * int(qty), 2)


def _content_changed(existing: dict[str, Any] | None, qty: int, day: date) -> bool:
    if not existing:
        return True
    old_qty = int(existing.get("planned_qty") or 0)
    old_day = _as_date(existing.get("planned_arrival_date"))
    return old_qty != int(qty) or old_day != day


def _status_for_save(*, has_assignee: bool, is_create: bool, had_assignee_before: bool) -> str:
    if not has_assignee:
        return STATUS_RESERVED
    if is_create or had_assignee_before:
        return STATUS_CHANGED
    return STATUS_ASSIGNED


def _status_after_edit(existing: dict[str, Any] | None, qty: int, day: date, has_assignee: bool) -> str:
    if not existing:
        return _status_for_save(has_assignee=has_assignee, is_create=True, had_assignee_before=False)
    if not _content_changed(existing, qty, day):
        return _norm(existing.get("status")) or STATUS_RESERVED
    return _status_for_save(
        has_assignee=has_assignee,
        is_create=False,
        had_assignee_before=bool(existing.get("assignee_user_id")) or has_assignee,
    )


def _mark_packing_stale(day: date) -> None:
    execute(
        """
        UPDATE dalian_shipping_date_assignments
        SET packing_content_stale=1
        WHERE planned_arrival_date=%s AND packing_generated_at IS NOT NULL
        """,
        (day,),
    )


def _revert_started_packing(day: date, user: CurrentUser) -> None:
    execute(
        """
        UPDATE dalian_shipping_reservations
        SET status=%s, updated_by=%s, updated_by_user_id=%s, updated_at=NOW()
        WHERE planned_arrival_date=%s AND status=%s
        """,
        (STATUS_CHANGED, _operator_name(user), user.id, day, STATUS_PACKING),
    )


def _note_packing_content_changed(days: set[date], user: CurrentUser) -> None:
    for day in days:
        _mark_packing_stale(day)
        _revert_started_packing(day, user)


def _upsert_reservation(item: ReservationUpsertIn, user: CurrentUser) -> tuple[dict[str, Any], dict[str, Any]]:
    detail = _load_order_detail(item.order_detail_id)
    day = _parse_date(item.planned_arrival_date)
    qty = _as_int(item.planned_qty, "预定到货数量")
    remark = _norm(item.remark)
    existing = fetch_one(
        f"SELECT {RESERVATION_SELECT} FROM dalian_shipping_reservations WHERE order_detail_id=%s AND planned_arrival_date=%s",
        (detail["id"], day),
    )
    exclude_id = int(existing["id"]) if existing else None
    split_meta = {"changed": False, "remainder_created": False, "remainder_qty": None}
    old_planned = int(existing.get("planned_qty") or 0) if existing else None
    qty_changed = old_planned is None or old_planned != qty
    if existing and reservation_is_locked(_norm(existing.get("status")), int(existing["id"])):
        existing_day = _as_date(existing.get("planned_arrival_date"))
        if qty_changed or existing_day != day:
            raise HTTPException(status_code=400, detail="已装箱或已有发货明细，只能修改备注")
    if qty_changed and (qty > 0 or existing):
        if qty > 0:
            assert_planned_qty_fits(int(detail["id"]), qty, exclude_id)
        split_meta = align_order_line_to_quantity(int(detail["id"]), qty if qty > 0 else 0, user)
        detail = _load_order_detail(int(detail["id"]))
    inherited = _inherit_assignee(day)
    operator = _operator_name(user)
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    item_no, drawing_no, unit_weight, total_weight = _snapshot_weights(detail, qty)
    changed_days = {day} if _content_changed(existing, qty, day) else set()

    if existing:
        before = _reservation_from_db(existing)
        new_status = _status_after_edit(existing, qty, day, inherited["has_assignee"])
        execute(
            """
            UPDATE dalian_shipping_reservations
            SET planned_qty=%s, remark=%s, status=%s, order_qty=%s,
                order_no=%s, factory_order_no=%s, part_no=%s, material_no=%s,
                spec_model=%s, material=%s, item_no=%s, drawing_no=%s,
                unit_weight=%s, total_weight=%s,
                assignee_user_id=%s, assignee_name=%s, assigned_at=%s,
                assigned_by_user_id=%s, assigned_by_name=%s,
                updated_by=%s, updated_by_user_id=%s, updated_at=NOW()
            WHERE id=%s
            """,
            (
                qty,
                remark,
                new_status,
                detail["quantity"],
                detail["order_no"],
                detail["factory_order_no"],
                detail["part_no"],
                detail["material_no"],
                detail["spec_model"],
                detail["material"],
                item_no,
                drawing_no,
                unit_weight,
                total_weight,
                inherited["assignee_user_id"],
                inherited["assignee_name"],
                inherited["assigned_at"],
                inherited["assigned_by_user_id"],
                inherited["assigned_by_name"],
                operator,
                user.id,
                int(existing["id"]),
            ),
        )
        after = _load_reservation(int(existing["id"]))
        _write_log(
            action="update",
            user=user,
            row_id=after["id"],
            order_detail_id=detail["id"],
            order_no=detail["order_no"],
            material_no=detail["material_no"],
            planned_arrival_date=day,
            change_summary={
                "op": "update",
                "fields": {
                    "planned_qty": {"before": before["planned_qty"], "after": after["planned_qty"]},
                    "remark": {"before": before["remark"], "after": after["remark"]},
                    "status": {"before": before["status"], "after": after["status"]},
                },
            },
        )
        if changed_days:
            _note_packing_content_changed(changed_days, user)
        if inherited["has_assignee"]:
            _notify_date(day, extra_user_ids={user.id})
        return after, split_meta

    status = _status_for_save(has_assignee=inherited["has_assignee"], is_create=True, had_assignee_before=False)
    row_id = execute(
        """
        INSERT INTO dalian_shipping_reservations (
            order_detail_id, order_no, factory_order_no, part_no, material_no, spec_model, material,
            item_no, drawing_no, unit_weight, total_weight, order_qty,
            planned_arrival_date, planned_qty, remark, status,
            booker_user_id, booker_name, booked_at,
            assignee_user_id, assignee_name, assigned_at, assigned_by_user_id, assigned_by_name,
            created_by, created_by_user_id, updated_by, updated_by_user_id
        ) VALUES (
            %s, %s, %s, %s, %s, %s, %s,
            %s, %s, %s, %s, %s,
            %s, %s, %s, %s,
            %s, %s, %s,
            %s, %s, %s, %s, %s,
            %s, %s, %s, %s
        )
        """,
        (
            detail["id"],
            detail["order_no"],
            detail["factory_order_no"],
            detail["part_no"],
            detail["material_no"],
            detail["spec_model"],
            detail["material"],
            item_no,
            drawing_no,
            unit_weight,
            total_weight,
            detail["quantity"],
            day,
            qty,
            remark,
            status,
            user.id,
            operator,
            now,
            inherited["assignee_user_id"],
            inherited["assignee_name"],
            inherited["assigned_at"],
            inherited["assigned_by_user_id"],
            inherited["assigned_by_name"],
            operator,
            user.id,
            operator,
            user.id,
        ),
    )
    after = _load_reservation(row_id)
    _write_log(
        action="create",
        user=user,
        row_id=after["id"],
        order_detail_id=detail["id"],
        order_no=detail["order_no"],
        material_no=detail["material_no"],
        planned_arrival_date=day,
        change_summary={"op": "create", "planned_qty": qty, "remark": remark, "status": status},
    )
    if qty > 0:
        _note_packing_content_changed({day}, user)
    if inherited["has_assignee"]:
        _notify_date(day, extra_user_ids={user.id})
    return after, split_meta


def _unshipped_filters(
    order_no: str,
    factory_order_no: str,
    material_no: str,
    part_no: str,
    statuses: list[str] | None = None,
) -> tuple[str, list[Any]]:
    clauses = ["o.order_no LIKE '4%'"]
    params: list[Any] = []
    if _norm(order_no):
        clauses.append("o.order_no LIKE %s")
        params.append(f"%{_norm(order_no)}%")
    if _norm(factory_order_no):
        clauses.append("COALESCE(o.factory_order_no,'') LIKE %s")
        params.append(f"%{_norm(factory_order_no)}%")
    if _norm(material_no):
        clauses.append("o.material_no LIKE %s")
        params.append(f"%{_norm(material_no)}%")
    if _norm(part_no):
        clauses.append("COALESCE(o.name,'') LIKE %s")
        params.append(f"%{_norm(part_no)}%")
    allowed = list(dict.fromkeys(s for s in (_norm(item) for item in (statuses or [])) if s in ORDER_DETAIL_STATUSES))
    if allowed:
        placeholders = ", ".join(["%s"] * len(allowed))
        clauses.append(f"o.order_status IN ({placeholders})")
        params.extend(allowed)
    return " AND ".join(clauses), params


@router.get("/unshipped-details")
def list_unshipped_details(
    _user: UserDep,
    order_no: str = "",
    factory_order_no: str = "",
    material_no: str = "",
    part_no: str = "",
    status: list[str] = Query(default=[]),
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=200),
):
    _ensure_tables()
    where_sql, params = _unshipped_filters(order_no, factory_order_no, material_no, part_no, status)
    count_row = fetch_one(
        f"""
        SELECT COUNT(*) AS cnt
        FROM order_details o
        LEFT JOIN (
            SELECT order_detail_id, COALESCE(SUM(planned_qty), 0) AS reserved_qty
            FROM dalian_shipping_reservations
            WHERE planned_qty > 0
            GROUP BY order_detail_id
        ) r ON r.order_detail_id = o.id
        WHERE {where_sql}
          AND (o.quantity - COALESCE(r.reserved_qty, 0)) > 0
        """,
        tuple(params),
    )
    total = int((count_row or {}).get("cnt") or 0)
    offset = (page - 1) * page_size
    rows = fetch_all(
        f"""
        SELECT o.id, o.order_no, o.factory_order_no, o.name AS part_no, o.material_no,
               o.spec_model, o.material, o.order_status AS status, o.quantity AS order_qty,
               COALESCE(r.reserved_qty, 0) AS reserved_qty,
               (o.quantity - COALESCE(r.reserved_qty, 0)) AS remaining_qty
        FROM order_details o
        LEFT JOIN (
            SELECT order_detail_id, COALESCE(SUM(planned_qty), 0) AS reserved_qty
            FROM dalian_shipping_reservations
            WHERE planned_qty > 0
            GROUP BY order_detail_id
        ) r ON r.order_detail_id = o.id
        WHERE {where_sql}
          AND (o.quantity - COALESCE(r.reserved_qty, 0)) > 0
        ORDER BY o.order_no, o.id
        LIMIT %s OFFSET %s
        """,
        (*params, page_size, offset),
    )
    detail_ids = [int(r["id"]) for r in rows]
    existing_map: dict[int, list[dict[str, Any]]] = {did: [] for did in detail_ids}
    if detail_ids:
        placeholders = ", ".join(["%s"] * len(detail_ids))
        existing_rows = fetch_all(
            f"""
            SELECT id, order_detail_id, planned_arrival_date, planned_qty, remark, status
            FROM dalian_shipping_reservations
            WHERE order_detail_id IN ({placeholders})
            ORDER BY planned_arrival_date
            """,
            tuple(detail_ids),
        )
        for item in existing_rows:
            existing_map.setdefault(int(item["order_detail_id"]), []).append(
                {
                    "id": int(item["id"]),
                    "planned_arrival_date": _jsonable(item.get("planned_arrival_date")) or "",
                    "planned_qty": int(item.get("planned_qty") or 0),
                    "remark": _norm(item.get("remark")),
                    "status": _norm(item.get("status")),
                }
            )
    data = []
    for row in rows:
        oid = int(row["id"])
        data.append(
            {
                "order_detail_id": oid,
                "order_no": _norm(row.get("order_no")),
                "factory_order_no": _norm(row.get("factory_order_no")),
                "part_no": _norm(row.get("part_no")),
                "material_no": _norm(row.get("material_no")),
                "spec_model": _norm(row.get("spec_model")),
                "material": _norm(row.get("material")),
                "status": _norm(row.get("status")),
                "order_qty": int(row.get("order_qty") or 0),
                "reserved_qty": int(row.get("reserved_qty") or 0),
                "remaining_qty": int(row.get("remaining_qty") or 0),
                "existing_reservations": existing_map.get(oid, []),
            }
        )
    return {"status": "success", "total": total, "page": page, "page_size": page_size, "rows": data}


@router.get("/reservations")
def list_reservations(
    _user: UserDep,
    planned_arrival_date: str = Query(..., min_length=8),
):
    _ensure_tables()
    day = _parse_date(planned_arrival_date)
    rows = fetch_all(
        f"""
        SELECT r.*, o.quantity AS live_order_qty,
               COALESCE(s.reserved_qty, 0) AS reserved_qty
        FROM dalian_shipping_reservations r
        LEFT JOIN order_details o ON o.id = r.order_detail_id
        LEFT JOIN (
            SELECT order_detail_id, COALESCE(SUM(planned_qty), 0) AS reserved_qty
            FROM dalian_shipping_reservations
            WHERE planned_qty > 0
            GROUP BY order_detail_id
        ) s ON s.order_detail_id = r.order_detail_id
        WHERE r.planned_arrival_date=%s
        ORDER BY r.order_no, r.id
        """,
        (day,),
    )
    assignment = _get_date_assignment(day)
    data = []
    for row in rows:
        item = _reservation_from_db(row)
        live_qty = int(row.get("live_order_qty") if row.get("live_order_qty") is not None else item["order_qty"])
        reserved = int(row.get("reserved_qty") or 0)
        item["order_qty"] = live_qty
        item["reserved_qty"] = reserved
        item["remaining_qty"] = live_qty - reserved
        data.append(item)
    return {
        "status": "success",
        "planned_arrival_date": day.strftime("%Y-%m-%d"),
        "assignee_user_id": int(assignment["assignee_user_id"]) if assignment and assignment.get("assignee_user_id") else None,
        "assignee_name": _norm(assignment.get("assignee_name")) if assignment else "",
        "assigned_at": _jsonable(assignment.get("assigned_at")) if assignment else "",
        "assigned_by_name": _norm(assignment.get("assigned_by_name")) if assignment else "",
        "rows": data,
    }


@router.post("/reservations")
def upsert_reservation(item: ReservationUpsertIn, user: UserDep):
    _ensure_tables()
    row, split_meta = _upsert_reservation(item, user)
    return {"status": "success", "row": row, "order_split": split_meta}


@router.post("/reservations/batch")
def batch_upsert_reservations(req: ReservationBatchRequest, user: UserDep):
    _ensure_tables()
    saved: list[dict[str, Any]] = []
    errors: list[str] = []
    skipped = 0
    for idx, item in enumerate(req.items, start=1):
        if not _norm(item.planned_arrival_date):
            skipped += 1
            continue
        try:
            row, _split_meta = _upsert_reservation(item, user)
            saved.append(row)
        except HTTPException as exc:
            errors.append(f"第{idx}行：{exc.detail}")
    if errors and not saved:
        raise HTTPException(status_code=400, detail="；".join(errors))
    return {
        "status": "success",
        "saved_count": len(saved),
        "skipped_count": skipped,
        "error_count": len(errors),
        "errors": errors,
        "rows": saved,
    }


@router.patch("/reservations/{row_id}")
def patch_reservation(row_id: int, req: ReservationPatchRequest, user: UserDep):
    _ensure_tables()
    current = fetch_one(f"SELECT {RESERVATION_SELECT} FROM dalian_shipping_reservations WHERE id=%s", (row_id,))
    if not current:
        raise HTTPException(status_code=404, detail="未找到该发货预定")
    before = _reservation_from_db(current)
    detail = _load_order_detail(int(current["order_detail_id"]))
    old_day = current["planned_arrival_date"]
    if isinstance(old_day, datetime):
        old_day = old_day.date()
    elif not isinstance(old_day, date):
        old_day = _parse_date(str(old_day))
    new_day = _parse_date(req.planned_arrival_date) if req.planned_arrival_date is not None else old_day
    qty = _as_int(req.planned_qty, "预定到货数量") if req.planned_qty is not None else int(current.get("planned_qty") or 0)
    remark = _norm(req.remark) if req.remark is not None else _norm(current.get("remark"))
    if new_day != old_day:
        clash = fetch_one(
            "SELECT id FROM dalian_shipping_reservations WHERE order_detail_id=%s AND planned_arrival_date=%s AND id<>%s",
            (detail["id"], new_day, row_id),
        )
        if clash:
            raise HTTPException(status_code=400, detail="该明细在目标到货日已有预定，请直接修改那一条")
    qty_changed = qty != int(current.get("planned_qty") or 0)
    if reservation_is_locked(_norm(current.get("status")), row_id) and (qty_changed or new_day != old_day):
        raise HTTPException(status_code=400, detail="已装箱或已有发货明细，只能修改备注")
    if qty_changed:
        if qty > 0:
            assert_planned_qty_fits(int(detail["id"]), qty, row_id)
        align_order_line_to_quantity(int(detail["id"]), qty if qty > 0 else 0, user)
        detail = _load_order_detail(int(detail["id"]))
    inherited = _inherit_assignee(new_day)
    has_assignee = inherited["has_assignee"] or bool(current.get("assignee_user_id"))
    new_status = _status_after_edit(current, qty, new_day, has_assignee)
    item_no, drawing_no, unit_weight, total_weight = _snapshot_weights(detail, qty)
    assignee = inherited if inherited["has_assignee"] else {
        "assignee_user_id": current.get("assignee_user_id"),
        "assignee_name": _norm(current.get("assignee_name")),
        "assigned_at": current.get("assigned_at"),
        "assigned_by_user_id": current.get("assigned_by_user_id"),
        "assigned_by_name": _norm(current.get("assigned_by_name")),
    }
    execute(
        """
        UPDATE dalian_shipping_reservations
        SET planned_arrival_date=%s, planned_qty=%s, remark=%s, status=%s, order_qty=%s,
            item_no=%s, drawing_no=%s, unit_weight=%s, total_weight=%s,
            assignee_user_id=%s, assignee_name=%s, assigned_at=%s,
            assigned_by_user_id=%s, assigned_by_name=%s,
            updated_by=%s, updated_by_user_id=%s, updated_at=NOW()
        WHERE id=%s
        """,
        (
            new_day,
            qty,
            remark,
            new_status,
            detail["quantity"],
            item_no,
            drawing_no,
            unit_weight,
            total_weight,
            assignee.get("assignee_user_id"),
            assignee.get("assignee_name") or "",
            assignee.get("assigned_at"),
            assignee.get("assigned_by_user_id"),
            assignee.get("assigned_by_name") or "",
            _operator_name(user),
            user.id,
            row_id,
        ),
    )
    after = _load_reservation(row_id)
    _write_log(
        action="update",
        user=user,
        row_id=row_id,
        order_detail_id=detail["id"],
        order_no=detail["order_no"],
        material_no=detail["material_no"],
        planned_arrival_date=new_day,
        change_summary={
            "op": "patch",
            "fields": {
                "planned_arrival_date": {"before": before["planned_arrival_date"], "after": after["planned_arrival_date"]},
                "planned_qty": {"before": before["planned_qty"], "after": after["planned_qty"]},
                "remark": {"before": before["remark"], "after": after["remark"]},
                "status": {"before": before["status"], "after": after["status"]},
            },
        },
    )
    if _content_changed(current, qty, new_day):
        _note_packing_content_changed({old_day, new_day}, user)
    if has_assignee:
        _notify_date(new_day, extra_user_ids={user.id})
        if new_day != old_day:
            _notify_date(old_day, extra_user_ids={user.id})
    return {"status": "success", "row": after}


@router.delete("/reservations/{row_id}")
def delete_reservation(row_id: int, user: UserDep):
    _ensure_tables()
    current = fetch_one(f"SELECT {RESERVATION_SELECT} FROM dalian_shipping_reservations WHERE id=%s", (row_id,))
    if not current:
        raise HTTPException(status_code=404, detail="未找到该发货预定")
    before = _reservation_from_db(current)
    day_raw = current["planned_arrival_date"]
    day = day_raw.date() if isinstance(day_raw, datetime) else (day_raw if isinstance(day_raw, date) else _parse_date(str(day_raw)))
    if reservation_is_locked(_norm(current.get("status")), row_id):
        raise HTTPException(status_code=400, detail="已装箱或已有发货明细，不能删除预定")
    if int(before.get("planned_qty") or 0) > 0:
        align_order_line_to_quantity(int(current["order_detail_id"]), 0, user, remove_if_zero=True)
    execute("DELETE FROM dalian_shipping_reservations WHERE id=%s", (row_id,))
    if int(before.get("planned_qty") or 0) > 0:
        _note_packing_content_changed({day}, user)
    _write_log(
        action="delete",
        user=user,
        row_id=None,
        order_detail_id=before["order_detail_id"],
        order_no=before["order_no"],
        material_no=before["material_no"],
        planned_arrival_date=day,
        change_summary={"op": "delete", "snapshot": before},
    )
    assignment = _get_date_assignment(day)
    if assignment and assignment.get("assignee_user_id"):
        remaining = fetch_all(
            "SELECT id FROM dalian_shipping_reservations WHERE planned_arrival_date=%s AND assignee_user_id IS NOT NULL",
            (day,),
        )
        if remaining:
            execute(
                """
                UPDATE dalian_shipping_reservations
                SET status=%s, updated_by=%s, updated_by_user_id=%s, updated_at=NOW()
                WHERE planned_arrival_date=%s AND status IN (%s, %s, %s)
                """,
                (STATUS_CHANGED, _operator_name(user), user.id, day, STATUS_ASSIGNED, STATUS_CHANGED, STATUS_PACKING),
            )
        _notify_date(day, extra_user_ids={user.id})
    return {"status": "success", "deleted_id": row_id}


@router.post("/assign")
def assign_date(req: AssignRequest, user: UserDep):
    _ensure_tables()
    day = _parse_date(req.planned_arrival_date)
    target = _load_user(req.assignee_user_id)
    existing = _get_date_assignment(day)
    operator = _operator_name(user)
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    action = "reassign" if existing and existing.get("assignee_user_id") and int(existing["assignee_user_id"]) != target["id"] else "assign"
    if existing:
        execute(
            """
            UPDATE dalian_shipping_date_assignments
            SET assignee_user_id=%s, assignee_name=%s, assigned_at=%s,
                assigned_by_user_id=%s, assigned_by_name=%s,
                updated_by=%s, updated_by_user_id=%s, updated_at=NOW()
            WHERE planned_arrival_date=%s
            """,
            (target["id"], target["name"], now, user.id, operator, operator, user.id, day),
        )
    else:
        execute(
            """
            INSERT INTO dalian_shipping_date_assignments (
                planned_arrival_date, assignee_user_id, assignee_name, assigned_at,
                assigned_by_user_id, assigned_by_name,
                created_by, created_by_user_id, updated_by, updated_by_user_id
            ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
            """,
            (day, target["id"], target["name"], now, user.id, operator, operator, user.id, operator, user.id),
        )
    new_status = STATUS_CHANGED if action == "reassign" else STATUS_ASSIGNED
    execute(
        """
        UPDATE dalian_shipping_reservations
        SET assignee_user_id=%s, assignee_name=%s, assigned_at=%s,
            assigned_by_user_id=%s, assigned_by_name=%s, status=%s,
            updated_by=%s, updated_by_user_id=%s, updated_at=NOW()
        WHERE planned_arrival_date=%s
        """,
        (target["id"], target["name"], now, user.id, operator, new_status, operator, user.id, day),
    )
    _write_log(
        action=action,
        user=user,
        row_id=None,
        order_detail_id=None,
        order_no="",
        material_no="",
        planned_arrival_date=day,
        change_summary={
            "op": action,
            "assignee": {"before": _norm(existing.get("assignee_name")) if existing else "", "after": target["name"]},
            "assignee_user_id": {"before": existing.get("assignee_user_id") if existing else None, "after": target["id"]},
        },
    )
    extra = {user.id, target["id"]}
    if existing and existing.get("assignee_user_id"):
        extra.add(int(existing["assignee_user_id"]))
    _notify_date(day, extra_user_ids=extra)
    return {"status": "success", "planned_arrival_date": day.strftime("%Y-%m-%d"), "assignee_name": target["name"], "action": action}


@router.get("/logs")
def list_logs(
    _user: UserDep,
    order_no: str = "",
    material_no: str = "",
    planned_arrival_date: str = "",
    action: str = "",
    limit: int = Query(200, ge=1, le=500),
):
    _ensure_tables()
    clauses = ["1=1"]
    params: list[Any] = []
    if _norm(order_no):
        clauses.append("order_no LIKE %s")
        params.append(f"%{_norm(order_no)}%")
    if _norm(material_no):
        clauses.append("material_no LIKE %s")
        params.append(f"%{_norm(material_no)}%")
    if _norm(planned_arrival_date):
        clauses.append("planned_arrival_date=%s")
        params.append(_parse_date(planned_arrival_date))
    if _norm(action):
        clauses.append("action=%s")
        params.append(_norm(action))
    rows = fetch_all(
        f"""
        SELECT id, row_id, order_detail_id, order_no, material_no, planned_arrival_date,
               action, operator_user_id, operator_name, operated_at, change_summary
        FROM dalian_shipping_reservation_logs
        WHERE {' AND '.join(clauses)}
        ORDER BY id DESC
        LIMIT %s
        """,
        (*params, limit),
    )
    data = []
    for row in rows:
        item = {k: _jsonable(v) for k, v in row.items()}
        item["operator_name"] = _norm(item.get("operator_name"))
        item["change_summary"] = item.get("change_summary") or ""
        data.append(item)
    return {"status": "success", "total": len(data), "rows": data}


def _attach_order_heat_fields(items: list[dict[str, Any]]) -> None:
    """箱单显示用：炉号、热处理号取订单明细当前值，不写入预定表。"""
    ids = sorted({int(item["order_detail_id"]) for item in items if item.get("order_detail_id")})
    by_id: dict[int, dict[str, Any]] = {}
    if ids:
        placeholders = ",".join(["%s"] * len(ids))
        rows = fetch_all(
            f"SELECT id, heat_no, heat_treatment_batch_no FROM order_details WHERE id IN ({placeholders})",
            tuple(ids),
        )
        by_id = {int(row["id"]): row for row in rows}
    for item in items:
        source = by_id.get(int(item.get("order_detail_id") or 0), {})
        item["heat_no"] = _norm(source.get("heat_no"))
        item["heat_treatment_batch_no"] = _norm(source.get("heat_treatment_batch_no"))


def _packing_view(day: date) -> dict[str, Any]:
    assignment = _get_date_assignment(day) or {}
    lines = fetch_all(
        f"""
        SELECT {RESERVATION_SELECT}
        FROM dalian_shipping_reservations
        WHERE planned_arrival_date=%s
        ORDER BY order_no, id
        """,
        (day,),
    )
    items = [_reservation_from_db(row) for row in lines]
    active = [item for item in items if int(item.get("planned_qty") or 0) > 0]
    _attach_order_heat_fields(active)
    has_changes = any(item.get("status") == STATUS_CHANGED for item in items)
    generated_at = _jsonable(assignment.get("packing_generated_at")) or ""
    stale = int(assignment.get("packing_content_stale") or 0) == 1
    has_assignee = bool(assignment.get("assignee_user_id"))
    return {
        "planned_arrival_date": day.strftime("%Y-%m-%d"),
        "assignee_user_id": int(assignment["assignee_user_id"]) if has_assignee else None,
        "assignee_name": _norm(assignment.get("assignee_name")),
        "packing_generated_by": _norm(assignment.get("packing_generated_by")),
        "packing_generated_by_user_id": assignment.get("packing_generated_by_user_id"),
        "packing_generated_at": generated_at,
        "packing_content_stale": stale,
        "has_changes": has_changes,
        "can_generate": has_assignee and bool(active),
        "can_start_packing": bool(has_assignee and generated_at and active and not stale and not has_changes),
        "rows": active,
    }


def _require_packing_date(day: date) -> dict[str, Any]:
    view = _packing_view(day)
    if not view["assignee_user_id"]:
        raise HTTPException(status_code=400, detail="该日尚未指定对接人员，不能生成箱单")
    if not view["rows"]:
        raise HTTPException(status_code=400, detail="该日没有预定数量大于 0 的明细")
    return view


@router.get("/packing-list")
def get_packing_list(planned_arrival_date: str, _user: UserDep):
    _ensure_tables()
    day = _parse_date(planned_arrival_date)
    return {"status": "success", **_packing_view(day)}


@router.post("/packing-list/generate")
def generate_packing_list(req: PackingDateRequest, user: UserDep):
    _ensure_tables()
    day = _parse_date(req.planned_arrival_date)
    before = _require_packing_date(day)
    assignment = _get_date_assignment(day) or {}
    action = "generate" if not assignment.get("packing_generated_at") else "regenerate"
    rows = fetch_all(
        f"SELECT {RESERVATION_SELECT} FROM dalian_shipping_reservations WHERE planned_arrival_date=%s",
        (day,),
    )
    operator = _operator_name(user)
    refreshed = 0
    for raw in rows:
        detail = _load_order_detail(int(raw["order_detail_id"]))
        qty = int(raw.get("planned_qty") or 0)
        item_no, drawing_no, unit_weight, total_weight = _snapshot_weights(detail, qty)
        status = _norm(raw.get("status")) or STATUS_RESERVED
        if status == STATUS_CHANGED:
            status = STATUS_ASSIGNED
        execute(
            """
            UPDATE dalian_shipping_reservations
            SET item_no=%s, drawing_no=%s, unit_weight=%s, total_weight=%s,
                order_no=%s, factory_order_no=%s, part_no=%s, material_no=%s,
                spec_model=%s, material=%s, order_qty=%s, status=%s,
                updated_by=%s, updated_by_user_id=%s, updated_at=NOW()
            WHERE id=%s
            """,
            (
                item_no,
                drawing_no,
                unit_weight,
                total_weight,
                detail["order_no"],
                detail["factory_order_no"],
                detail["part_no"],
                detail["material_no"],
                detail["spec_model"],
                detail["material"],
                detail["quantity"],
                status,
                operator,
                user.id,
                int(raw["id"]),
            ),
        )
        refreshed += 1
    execute(
        """
        UPDATE dalian_shipping_date_assignments
        SET packing_generated_by=%s, packing_generated_by_user_id=%s, packing_generated_at=NOW(),
            packing_content_stale=0, updated_by=%s, updated_by_user_id=%s, updated_at=NOW()
        WHERE planned_arrival_date=%s
        """,
        (operator, user.id, operator, user.id, day),
    )
    _write_log(
        action=action,
        user=user,
        row_id=None,
        order_detail_id=None,
        order_no="",
        material_no="",
        planned_arrival_date=day,
        change_summary={
            "op": action,
            "line_count": len(before["rows"]),
            "refreshed_count": refreshed,
            "row_ids": [item["id"] for item in before["rows"]],
        },
    )
    return {"status": "success", "action": action, **_packing_view(day)}


@router.post("/packing-list/export")
def export_packing_list(req: PackingExportRequest, user: UserDep):
    _ensure_tables()
    day = _parse_date(req.planned_arrival_date)
    view = _packing_view(day)
    if not view["rows"]:
        raise HTTPException(status_code=400, detail="该日没有可导出的预定明细")
    allowed = {key: label for key, label in PACKING_EXPORT_COLUMNS}
    columns = [_norm(key) for key in req.columns if _norm(key)]
    if not columns:
        raise HTTPException(status_code=400, detail="请选择要导出的列")
    unknown = [key for key in columns if key not in allowed]
    if unknown:
        raise HTTPException(status_code=400, detail="包含不支持的导出列")
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise HTTPException(status_code=500, detail="缺少 openpyxl 依赖，无法导出 Excel") from exc
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "大连客户箱单"
    sheet.append([allowed[key] for key in columns])
    for item in view["rows"]:
        sheet.append(["" if item.get(key) is None else item.get(key) for key in columns])
    bio = io.BytesIO()
    workbook.save(bio)
    bio.seek(0)
    file_name = f"大连客户箱单_{day.strftime('%Y-%m-%d')}.xlsx"
    _write_log(
        action="export",
        user=user,
        row_id=None,
        order_detail_id=None,
        order_no="",
        material_no="",
        planned_arrival_date=day,
        change_summary={"op": "export", "file_name": file_name, "columns": columns, "line_count": len(view["rows"])},
    )
    return StreamingResponse(
        bio,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(file_name)}"},
    )


@router.post("/packing-list/start")
def start_packing_list(req: PackingDateRequest, user: UserDep):
    _ensure_tables()
    day = _parse_date(req.planned_arrival_date)
    view = _require_packing_date(day)
    if not view["packing_generated_at"]:
        raise HTTPException(status_code=400, detail="请先生成箱单")
    if view["packing_content_stale"] or view["has_changes"]:
        raise HTTPException(status_code=400, detail="预定内容已变化，请先重新生成箱单后再开始自动装箱")
    operator = _operator_name(user)
    execute(
        """
        UPDATE dalian_shipping_reservations
        SET status=%s, updated_by=%s, updated_by_user_id=%s, updated_at=NOW()
        WHERE planned_arrival_date=%s AND planned_qty>0 AND status NOT IN (%s, %s)
        """,
        (STATUS_PACKING, operator, user.id, day, STATUS_PACKED, STATUS_SHIPPED),
    )
    _write_log(
        action="start_packing",
        user=user,
        row_id=None,
        order_detail_id=None,
        order_no="",
        material_no="",
        planned_arrival_date=day,
        change_summary={"op": "start_packing", "line_count": len(view["rows"]), "row_ids": [item["id"] for item in view["rows"]]},
    )
    _notify_date(day, extra_user_ids={user.id})
    return {"status": "success", "action": "start_packing", **_packing_view(day)}


def get_dalian_shipping_progress(user_id: int) -> dict[str, Any]:
    """工作台发货进度：不计入待办 count。"""
    _ensure_tables()
    date_rows = fetch_all(
        """
        SELECT DISTINCT planned_arrival_date AS d FROM (
            SELECT planned_arrival_date FROM dalian_shipping_reservations WHERE booker_user_id=%s
            UNION
            SELECT planned_arrival_date FROM dalian_shipping_reservations WHERE assignee_user_id=%s
            UNION
            SELECT planned_arrival_date FROM dalian_shipping_date_assignments WHERE assignee_user_id=%s
        ) t
        ORDER BY d
        """,
        (user_id, user_id, user_id),
    )
    unread_rows = fetch_all(
        """
        SELECT planned_arrival_date, unread
        FROM dalian_shipping_progress_unreads
        WHERE user_id=%s
        """,
        (user_id,),
    )
    unread_map = {
        _as_date(r["planned_arrival_date"]).strftime("%Y-%m-%d"): int(r.get("unread") or 0) == 1
        for r in unread_rows
    }
    groups: list[dict[str, Any]] = []
    unread_count = 0
    for date_row in date_rows:
        day = _as_date(date_row["d"])
        day_key = day.strftime("%Y-%m-%d")
        assignment = _get_date_assignment(day)
        lines = fetch_all(
            f"""
            SELECT {RESERVATION_SELECT}
            FROM dalian_shipping_reservations
            WHERE planned_arrival_date=%s
            ORDER BY order_no, id
            """,
            (day,),
        )
        items = [_reservation_from_db(r) for r in lines]
        total_qty = sum(int(x["planned_qty"] or 0) for x in items if int(x["planned_qty"] or 0) > 0)
        has_changes = any(x["status"] == STATUS_CHANGED for x in items)
        unread = bool(unread_map.get(day_key))
        if unread:
            unread_count += 1
        groups.append(
            {
                "planned_arrival_date": day_key,
                "assignee_user_id": int(assignment["assignee_user_id"]) if assignment and assignment.get("assignee_user_id") else None,
                "assignee_name": _norm(assignment.get("assignee_name")) if assignment else "",
                "assigned_at": _jsonable(assignment.get("assigned_at")) if assignment else "",
                "unread": unread,
                "has_changes": has_changes,
                "line_count": len(items),
                "total_qty": total_qty,
                "rows": items,
            }
        )
    return {"dates": groups, "unread_count": unread_count}


@router.post("/progress/{planned_arrival_date}/ack")
def ack_progress(planned_arrival_date: str, user: UserDep):
    _ensure_tables()
    day = _parse_date(planned_arrival_date)
    execute(
        """
        INSERT INTO dalian_shipping_progress_unreads (planned_arrival_date, user_id, unread)
        VALUES (%s, %s, 0)
        ON DUPLICATE KEY UPDATE unread=0
        """,
        (day, user.id),
    )
    assignment = _get_date_assignment(day)
    flipped = 0
    if assignment and assignment.get("assignee_user_id") and int(assignment["assignee_user_id"]) == int(user.id):
        execute(
            """
            UPDATE dalian_shipping_reservations
            SET status=%s, updated_at=NOW()
            WHERE planned_arrival_date=%s AND status=%s
            """,
            (STATUS_ASSIGNED, day, STATUS_CHANGED),
        )
        flipped_row = fetch_one(
            "SELECT COUNT(*) AS cnt FROM dalian_shipping_reservations WHERE planned_arrival_date=%s AND status=%s",
            (day, STATUS_ASSIGNED),
        )
        flipped = int((flipped_row or {}).get("cnt") or 0)
    return {"status": "success", "planned_arrival_date": day.strftime("%Y-%m-%d"), "cleared": True, "restored_assigned": flipped}
