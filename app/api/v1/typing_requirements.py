"""打字要求：按订单号（可选物料号）从订单明细/物料/图纸三表联合生成，可人工覆盖，独立表与操作日志。"""
from __future__ import annotations

import json
import re
from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Any

from fastapi import APIRouter, Depends, HTTPException, Query
from pydantic import BaseModel, Field
from typing_extensions import Annotated

from app.api.deps import CurrentUser, get_current_user
from app.api.v1.dn_inch_maps import load_inch_map
from app.api.v1.material_typing_maps import load_typing_material_map
from app.core.database import execute, fetch_all, fetch_one
from app.core.paths import get_templates_dir
from app.utils.drawing_no import (
    drawing_no_in_sql,
    drawing_no_match_variants,
    normalize_drawing_input,
    normalize_drawing_rev_for_db,
)
from app.utils.typing_content import build_typing_content, format_qr_spec

router = APIRouter(prefix="/typing-requirements", tags=["打字要求"])

UserDep = Annotated[CurrentUser, Depends(get_current_user)]

_TABLES_READY = False

TEMPLATE_NAME = "打字要求模板.xlsx"
MEDIA_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MISSING_FILL_RGB = "FFFFC7CE"

REMARK_GREEN_PAINT = "绿漆"
REMARK_NO_PAINT = "/"

# Excel 模板：A1 标题、第 2 行表头、第 3 行起数据、预铺样式到第 21 行
TEMPLATE_HEADER_ROW = 2
TEMPLATE_FIRST_DATA_ROW = 3
TEMPLATE_LAST_STYLED_ROW = 21
TEMPLATE_COL_COUNT = 9
EXPORT_COL_WIDTHS = (25.25, 21.3, 18, 106.6, 18.6, 22.3, 8.1, 10, 11.25)

EXPORT_COLUMNS = (
    ("drawing_rev_no", "图纸号"),
    ("material_no", "R3P物料号"),
    ("part_name", "料号"),
    ("typing_content", "打字内容如下："),
    ("heat_no", "炉号"),
    ("heat_treatment_batch_no", "热处理批号"),
    ("quantity", "数量"),
    ("unit_weight", "单重"),
    ("remark", "备注"),
)

CREATE_TABLE_SQL = """
        CREATE TABLE IF NOT EXISTS typing_requirements (
            id BIGINT PRIMARY KEY AUTO_INCREMENT,
            order_no VARCHAR(64) NOT NULL COMMENT '订单号',
            factory_order_no VARCHAR(128) NULL COMMENT '本厂单号',
            order_detail_id INT NULL COMMENT '订单明细ID，手工追加行为空',
            item_no INT NULL COMMENT '条目号',
            seq INT NULL COMMENT '订单明细序号',
            drawing_rev_no VARCHAR(255) NULL COMMENT '图纸号-版本号（Rev 形）',
            material_no VARCHAR(128) NULL COMMENT 'R3P物料号',
            part_name VARCHAR(128) NULL COMMENT '料号：法国默认-，大连取订单明细名称/料号',
            typing_content TEXT NULL COMMENT '打字内容',
            spec VARCHAR(255) NULL COMMENT '规格（展示：spec_display 或 spec）',
            model_standard VARCHAR(255) NULL COMMENT '型号（标准）',
            qr_spec VARCHAR(128) NULL COMMENT '二维码规格信息',
            heat_no VARCHAR(64) NULL COMMENT '炉号',
            heat_treatment_batch_no VARCHAR(64) NULL COMMENT '热处理批号',
            quantity INT NULL COMMENT '数量',
            unit_weight DECIMAL(12,2) NULL COMMENT '单重，取自订单明细',
            remark VARCHAR(255) NULL COMMENT '备注：绿漆 / /',
            manual_fields TEXT NULL COMMENT '人工改过的字段 JSON 数组',
            source_snapshot TEXT NULL COMMENT '生成时三表取值 JSON，便于溯源',
            created_by VARCHAR(64) NULL COMMENT '上传者',
            created_by_user_id INT NULL COMMENT '上传者用户ID',
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP COMMENT '上传时间',
            updated_by VARCHAR(64) NULL COMMENT '更改者',
            updated_by_user_id INT NULL COMMENT '更改者用户ID',
            updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP COMMENT '更改时间',
            UNIQUE KEY uk_typing_order_detail (order_detail_id),
            INDEX idx_typing_order_no (order_no),
            INDEX idx_typing_drawing (drawing_rev_no),
            INDEX idx_typing_updated (updated_at)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COMMENT='打字要求'
    """

CREATE_LOG_TABLE_SQL = """
        CREATE TABLE IF NOT EXISTS typing_requirement_logs (
            id BIGINT PRIMARY KEY AUTO_INCREMENT,
            row_id BIGINT NULL COMMENT '打字要求ID，删除后可空',
            order_no VARCHAR(64) NULL,
            drawing_rev_no VARCHAR(255) NULL,
            action VARCHAR(32) NOT NULL COMMENT 'generate/create/update/delete/export',
            operator_user_id INT NULL,
            operator_name VARCHAR(64) NOT NULL DEFAULT '',
            operated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            change_summary TEXT NULL COMMENT '变更内容 JSON',
            INDEX idx_typing_log_row (row_id),
            INDEX idx_typing_log_order (order_no),
            INDEX idx_typing_log_action (action),
            INDEX idx_typing_log_time (operated_at)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COMMENT='打字要求操作日志'
    """

SELECT_SQL = """
        id, order_no, factory_order_no, order_detail_id, item_no, seq,
        drawing_rev_no, material_no, part_name, typing_content,
        spec, model_standard, qr_spec, heat_no,
        heat_treatment_batch_no, quantity, unit_weight, remark, manual_fields, source_snapshot,
        created_by, created_by_user_id, created_at,
        updated_by, updated_by_user_id, updated_at
    """

EDITABLE_KEYS = (
    "drawing_rev_no",
    "material_no",
    "part_name",
    "typing_content",
    "heat_no",
    "heat_treatment_batch_no",
    "quantity",
    "unit_weight",
    "remark",
)

FIELD_LABELS = {
    "drawing_rev_no": "图纸号",
    "material_no": "R3P物料号",
    "part_name": "料号",
    "typing_content": "打字内容",
    "spec": "规格",
    "model_standard": "型号（标准）",
    "qr_spec": "二维码规格信息",
    "heat_no": "炉号",
    "heat_treatment_batch_no": "热处理批号",
    "quantity": "数量",
    "unit_weight": "单重",
    "remark": "备注",
}


class GenerateRequest(BaseModel):
    order_no: str = Field(..., min_length=1)
    material_no: str = ""
    order_detail_ids: list[int] = Field(default_factory=list)
    force_overwrite: bool = False


class BatchDeleteRequest(BaseModel):
    ids: list[int] = Field(default_factory=list)


class TypingRowCreate(BaseModel):
    order_no: str = Field(..., min_length=1)
    drawing_rev_no: str = ""
    material_no: str = ""
    part_name: str = ""
    typing_content: str = ""
    heat_no: str = ""
    heat_treatment_batch_no: str = ""
    quantity: int | None = None
    unit_weight: float | None = None
    remark: str = ""


class TypingRowUpdate(BaseModel):
    drawing_rev_no: str | None = None
    material_no: str | None = None
    part_name: str | None = None
    typing_content: str | None = None
    heat_no: str | None = None
    heat_treatment_batch_no: str | None = None
    quantity: int | None = None
    unit_weight: float | None = None
    remark: str | None = None


def _ensure_tables() -> None:
    global _TABLES_READY
    if _TABLES_READY:
        return
    execute(CREATE_TABLE_SQL)
    execute(CREATE_LOG_TABLE_SQL)
    _ensure_part_name_column()
    _ensure_unit_weight_column()
    _ensure_derived_spec_columns()
    _TABLES_READY = True


def _has_typing_column(column_name: str) -> bool:
    found = fetch_one(
        """
        SELECT COUNT(*) AS cnt
        FROM information_schema.COLUMNS
        WHERE TABLE_SCHEMA = DATABASE()
          AND TABLE_NAME = 'typing_requirements'
          AND COLUMN_NAME = %s
        """,
        (column_name,),
    )
    return bool(found and int(found.get("cnt") or 0) > 0)


def _ensure_part_name_column() -> None:
    if _has_typing_column("part_name"):
        return
    execute(
        """
        ALTER TABLE typing_requirements
        ADD COLUMN part_name VARCHAR(128) NULL
            COMMENT '料号：法国默认-，大连取订单明细名称/料号'
            AFTER material_no
        """
    )


def _ensure_unit_weight_column() -> None:
    if _has_typing_column("unit_weight"):
        return
    execute(
        """
        ALTER TABLE typing_requirements
        ADD COLUMN unit_weight DECIMAL(12,2) NULL
            COMMENT '单重，取自订单明细'
            AFTER quantity
        """
    )


def _ensure_derived_spec_columns() -> None:
    alters = (
        (
            "spec",
            """
            ALTER TABLE typing_requirements
            ADD COLUMN spec VARCHAR(255) NULL
                COMMENT '规格（展示：spec_display 或 spec）'
                AFTER typing_content
            """,
        ),
        (
            "model_standard",
            """
            ALTER TABLE typing_requirements
            ADD COLUMN model_standard VARCHAR(255) NULL
                COMMENT '型号（标准）'
                AFTER spec
            """,
        ),
        (
            "qr_spec",
            """
            ALTER TABLE typing_requirements
            ADD COLUMN qr_spec VARCHAR(128) NULL
                COMMENT '二维码规格信息'
                AFTER model_standard
            """,
        ),
    )
    for column_name, stmt in alters:
        if not _has_typing_column(column_name):
            execute(stmt)
    _backfill_derived_spec_columns()


def _derive_spec_fields(snapshot: dict[str, Any] | None) -> tuple[str, str, str]:
    data = snapshot or {}
    spec = _normalize_str(data.get("spec_display") or data.get("spec"))
    model = _normalize_str(data.get("model"))
    standard_no = _normalize_str(data.get("standard_no"))
    if model and standard_no:
        model_standard = f"{model}（{standard_no}）"
    else:
        model_standard = model or standard_no
    qr_spec = format_qr_spec(data.get("spec"), data.get("model"))
    return spec, model_standard, qr_spec


def _backfill_derived_spec_columns() -> None:
    rows = fetch_all(
        f"""
        SELECT {SELECT_SQL}
        FROM typing_requirements
        WHERE (IFNULL(spec, '') = '' OR IFNULL(model_standard, '') = '' OR IFNULL(qr_spec, '') = '')
          AND IFNULL(source_snapshot, '') NOT IN ('', '{{}}')
        """
    )
    for raw in rows or []:
        row = _row_from_db(raw)
        spec, model_standard, qr_spec = _derive_spec_fields(row.get("source_snapshot"))
        if (
            spec == _normalize_str(row.get("spec"))
            and model_standard == _normalize_str(row.get("model_standard"))
            and qr_spec == _normalize_str(row.get("qr_spec"))
        ):
            continue
        execute(
            """
            UPDATE typing_requirements
            SET spec=%s, model_standard=%s, qr_spec=%s
            WHERE id=%s
            """,
            (spec or None, model_standard or None, qr_spec or None, row["id"]),
        )


def _jsonable(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    return value


def _normalize_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _as_int(value: Any) -> int | None:
    if value is None or value == "":
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def _as_unit_weight(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        return round(float(value), 2)
    except (TypeError, ValueError):
        return None


def _load_json_list(raw: Any) -> list[str]:
    text = _normalize_str(raw)
    if not text:
        return []
    try:
        parsed = json.loads(text)
    except (TypeError, ValueError):
        return []
    if isinstance(parsed, list):
        return [str(item) for item in parsed]
    return []


def _load_json_dict(raw: Any) -> dict[str, Any]:
    text = _normalize_str(raw)
    if not text:
        return {}
    try:
        parsed = json.loads(text)
    except (TypeError, ValueError):
        return {}
    return parsed if isinstance(parsed, dict) else {}


def _row_from_db(row: dict[str, Any] | None) -> dict[str, Any]:
    if not row:
        return {}
    return {
        "id": int(row["id"]),
        "order_no": _normalize_str(row.get("order_no")),
        "factory_order_no": _normalize_str(row.get("factory_order_no")),
        "order_detail_id": int(row["order_detail_id"]) if row.get("order_detail_id") else None,
        "item_no": row.get("item_no"),
        "seq": row.get("seq"),
        "drawing_rev_no": _normalize_str(row.get("drawing_rev_no")),
        "material_no": _normalize_str(row.get("material_no")),
        "part_name": _normalize_str(row.get("part_name")),
        "typing_content": _normalize_str(row.get("typing_content")),
        "spec": _normalize_str(row.get("spec")),
        "model_standard": _normalize_str(row.get("model_standard")),
        "qr_spec": _normalize_str(row.get("qr_spec")),
        "heat_no": _normalize_str(row.get("heat_no")),
        "heat_treatment_batch_no": _normalize_str(row.get("heat_treatment_batch_no")),
        "quantity": row.get("quantity"),
        "unit_weight": _as_unit_weight(row.get("unit_weight")),
        "remark": _normalize_str(row.get("remark")),
        "manual_fields": _load_json_list(row.get("manual_fields")),
        "source_snapshot": _load_json_dict(row.get("source_snapshot")),
        "created_by": _normalize_str(row.get("created_by")),
        "created_by_user_id": row.get("created_by_user_id"),
        "created_at": _jsonable(row.get("created_at")) or "",
        "updated_by": _normalize_str(row.get("updated_by")),
        "updated_by_user_id": row.get("updated_by_user_id"),
        "updated_at": _jsonable(row.get("updated_at")) or "",
    }


def _snapshot(row: dict[str, Any]) -> dict[str, Any]:
    out: dict[str, Any] = {}
    for key in EDITABLE_KEYS:
        if key == "quantity":
            out[key] = row.get(key)
        elif key == "unit_weight":
            out[key] = _as_unit_weight(row.get(key))
        else:
            out[key] = _normalize_str(row.get(key))
    return out


def _diff_fields(before: dict[str, Any], after: dict[str, Any]) -> dict[str, dict[str, Any]]:
    out: dict[str, dict[str, Any]] = {}
    for key in EDITABLE_KEYS:
        left = before.get(key)
        right = after.get(key)
        if key == "quantity":
            if _as_int(left) != _as_int(right):
                out[key] = {"before": _as_int(left), "after": _as_int(right)}
            continue
        if key == "unit_weight":
            if _as_unit_weight(left) != _as_unit_weight(right):
                out[key] = {"before": _as_unit_weight(left), "after": _as_unit_weight(right)}
            continue
        if _normalize_str(left) != _normalize_str(right):
            out[key] = {"before": _normalize_str(left), "after": _normalize_str(right)}
    return out


def _write_log(
    *,
    action: str,
    user: CurrentUser,
    row_id: int | None,
    order_no: str,
    drawing_rev_no: str,
    change_summary: dict[str, Any],
) -> None:
    execute(
        """
        INSERT INTO typing_requirement_logs
            (row_id, order_no, drawing_rev_no, action, operator_user_id, operator_name, change_summary)
        VALUES (%s, %s, %s, %s, %s, %s, %s)
        """,
        (
            row_id,
            order_no,
            drawing_rev_no,
            action,
            user.id,
            user.name,
            json.dumps(change_summary, ensure_ascii=False, default=str),
        ),
    )


CONTENT_MISSING_LABELS = {
    "spec_model": "规格型号",
    "spec": "规格（DN值）",
    "inch_value": "英寸（去「DN值-英寸」补录）",
    "typing_material": "打字材质（去「材质-打字材质」补录）",
    "heat_no": "炉号",
    "heat_treatment_batch_no": "热处理批号",
}


def _missing_info(row: dict[str, Any]) -> tuple[list[str], list[str]]:
    """返回（要标粉色的列, 缺失原因）。列名与表格/Excel 列一一对应。"""
    missing: list[str] = []
    for key in ("drawing_rev_no", "material_no", "heat_no", "heat_treatment_batch_no"):
        if not _normalize_str(row.get(key)):
            missing.append(key)
    order_no = _normalize_str(row.get("order_no"))
    if not _normalize_str(row.get("part_name")) and (
        order_no.startswith("2") or order_no.startswith("4")
    ):
        missing.append("part_name")
    if _as_int(row.get("quantity")) is None:
        missing.append("quantity")
    unit_weight = _as_unit_weight(row.get("unit_weight"))
    if unit_weight is None or unit_weight == 0:
        missing.append("unit_weight")

    snapshot = row.get("source_snapshot") or {}
    content_missing = [str(k) for k in (snapshot.get("content_missing") or [])]
    if not _normalize_str(row.get("typing_content")) or content_missing:
        missing.append("typing_content")
    if not _normalize_str(row.get("spec")) or "spec" in content_missing or "spec_model" in content_missing:
        missing.append("spec")
    if not _normalize_str(row.get("model_standard")):
        missing.append("model_standard")
    if not _normalize_str(row.get("qr_spec")):
        missing.append("qr_spec")

    reasons = [CONTENT_MISSING_LABELS.get(key, key) for key in content_missing]
    for key in missing:
        if key in ("drawing_rev_no", "material_no", "part_name", "quantity", "unit_weight", "spec", "model_standard", "qr_spec"):
            label = FIELD_LABELS.get(key, key)
            if label not in reasons:
                reasons.append(label)
    return missing, reasons


def _spec_display_fields(row: dict[str, Any]) -> tuple[str, str, str]:
    snapshot = row.get("source_snapshot") or {}
    spec, model_standard, qr_spec = _derive_spec_fields(snapshot)
    if spec or model_standard or qr_spec:
        return spec, model_standard, qr_spec
    return (
        _normalize_str(row.get("spec")),
        _normalize_str(row.get("model_standard")),
        _normalize_str(row.get("qr_spec")),
    )


def _with_missing(row: dict[str, Any]) -> dict[str, Any]:
    spec, model_standard, qr_spec = _spec_display_fields(row)
    decorated = {**row, "spec": spec, "model_standard": model_standard, "qr_spec": qr_spec}
    missing, reasons = _missing_info(decorated)
    return {**decorated, "missing": missing, "missing_reasons": reasons}


def _lookup_material(material_no: str, drawing_no: str) -> dict[str, Any]:
    if not material_no:
        return {}
    clause, params = drawing_no_in_sql("drawing_no", drawing_no)
    if drawing_no:
        row = fetch_one(
            f"""
            SELECT material_no, part_no, drawing_no, standard_no, spec_model, spec, model, material
            FROM materials
            WHERE material_no=%s AND {clause}
            ORDER BY id LIMIT 1
            """,
            (material_no, *params),
        )
        if row:
            return row
    return (
        fetch_one(
            """
            SELECT material_no, part_no, drawing_no, standard_no, spec_model, spec, model, material
            FROM materials
            WHERE material_no=%s
            ORDER BY id LIMIT 1
            """,
            (material_no,),
        )
        or {}
    )


def _lookup_drawing(drawing_no: str) -> dict[str, Any]:
    if not drawing_no:
        return {}
    clause, params = drawing_no_in_sql("drawing_no", drawing_no)
    rev_clause, rev_params = drawing_no_in_sql("drawing_rev_no", drawing_no)
    return (
        fetch_one(
            f"""
            SELECT drawing_no, drawing_rev_no, standard_no, spec_model, material, material_no, drawing_type
            FROM drawing_archives
            WHERE {clause} OR {rev_clause}
            ORDER BY id DESC LIMIT 1
            """,
            (*params, *rev_params),
        )
        or {}
    )


def _load_paint_remarks(drawing_nos: list[str]) -> dict[str, str]:
    """不划绿漆 → /，其余（含未录入）→ 绿漆。按图纸号变体匹配划漆要求。"""
    if not drawing_nos:
        return {}
    rows = fetch_all("SELECT drawing_rev_no, skip_green_paint FROM paint_requirements")
    skip_keys: set[str] = set()
    for row in rows:
        if not int(row.get("skip_green_paint") or 0):
            continue
        for variant in drawing_no_match_variants(_normalize_str(row.get("drawing_rev_no"))):
            skip_keys.add(variant.upper())
    out: dict[str, str] = {}
    for drawing_no in drawing_nos:
        variants = {v.upper() for v in drawing_no_match_variants(drawing_no)}
        out[drawing_no] = REMARK_NO_PAINT if variants & skip_keys else REMARK_GREEN_PAINT
    return out


def _pick_first(*values: Any) -> str:
    for value in values:
        text = _normalize_str(value)
        if text:
            return text
    return ""


def default_part_name(order_no: str, detail_name: Any = "") -> str:
    """2 开头法国默认「-」；4 开头大连取订单明细「名称/料号」。"""
    text = _normalize_str(order_no)
    if text.startswith("2"):
        return "-"
    if text.startswith("4"):
        return _normalize_str(detail_name)
    return ""


def _build_generated_row(
    detail: dict[str, Any],
    *,
    inch_by_dn: dict[str, str],
    typing_material_by_key: dict[str, str],
    paint_remarks: dict[str, str],
) -> dict[str, Any]:
    drawing_no = _normalize_str(detail.get("drawing_no"))
    material_no = _normalize_str(detail.get("material_no"))
    material_row = _lookup_material(material_no, drawing_no)
    drawing_row = _lookup_drawing(drawing_no)

    # 三表取值优先级：订单明细 → 物料 → 图纸
    spec_model = _pick_first(
        detail.get("spec_model"), material_row.get("spec_model"), drawing_row.get("spec_model")
    )
    material = _pick_first(detail.get("material"), material_row.get("material"), drawing_row.get("material"))
    heat_no = _normalize_str(detail.get("heat_no"))
    batch_no = _normalize_str(detail.get("heat_treatment_batch_no"))

    result = build_typing_content(
        spec_model=spec_model,
        material=material,
        material_no=material_no,
        part_no=material_row.get("part_no") or "",
        heat_no=heat_no,
        heat_treatment_batch_no=batch_no,
        inch_by_dn=inch_by_dn,
        typing_material_by_key=typing_material_by_key,
    )

    drawing_rev_no = normalize_drawing_rev_for_db(drawing_no) if drawing_no else ""
    snapshot = {
        "spec_model": spec_model,
        "material": material,
        "standard_no": result["standard_no"],
        "spec": result["spec"],
        "model": result["model"],
        "spec_display": result["spec_display"],
        "typing_material": result["typing_material"],
        "content_missing": result["missing"],
        "detail_name": _normalize_str(detail.get("name")),
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }
    spec, model_standard, qr_spec = _derive_spec_fields(snapshot)
    return {
        "order_no": _normalize_str(detail.get("order_no")),
        "factory_order_no": _normalize_str(detail.get("factory_order_no")),
        "order_detail_id": _as_int(detail.get("id")),
        "item_no": _as_int(detail.get("item_no")),
        "seq": _as_int(detail.get("seq")),
        "drawing_rev_no": drawing_rev_no,
        "material_no": material_no,
        "part_name": default_part_name(detail.get("order_no"), detail.get("name")),
        "typing_content": result["content"],
        "spec": spec,
        "model_standard": model_standard,
        "qr_spec": qr_spec,
        "heat_no": heat_no,
        "heat_treatment_batch_no": batch_no,
        "quantity": _as_int(detail.get("quantity")),
        "unit_weight": _as_unit_weight(detail.get("unit_weight")),
        "remark": paint_remarks.get(drawing_no, REMARK_GREEN_PAINT),
        "source_snapshot": snapshot,
    }


def _fetch_row(row_id: int) -> dict[str, Any] | None:
    return fetch_one(f"SELECT {SELECT_SQL} FROM typing_requirements WHERE id=%s", (row_id,))


def _list_rows(order_no: str, limit: int = 2000) -> list[dict[str, Any]]:
    clauses = ["1=1"]
    params: list[Any] = []
    if order_no:
        clauses.append("order_no=%s")
        params.append(order_no)
    rows = fetch_all(
        f"""
        SELECT {SELECT_SQL}
        FROM typing_requirements
        WHERE {' AND '.join(clauses)}
        ORDER BY COALESCE(seq, 999999), COALESCE(item_no, 999999), id
        LIMIT %s
        """,
        (*params, limit),
    )
    return [_with_missing(_row_from_db(r)) for r in rows]


@router.post("/generate")
def generate_typing_requirements(req: GenerateRequest, user: UserDep):
    _ensure_tables()
    order_no = _normalize_str(req.order_no)
    material_no = _normalize_str(req.material_no)
    detail_ids = sorted({int(item) for item in req.order_detail_ids if int(item) > 0})
    order_exists = fetch_one(
        "SELECT id FROM order_details WHERE order_no=%s LIMIT 1",
        (order_no,),
    )
    if not order_exists:
        raise HTTPException(status_code=404, detail=f"未找到订单号 {order_no} 的订单明细")

    clauses = ["order_no=%s"]
    params: list[Any] = [order_no]
    if material_no:
        clauses.append("material_no=%s")
        params.append(material_no)
    if detail_ids:
        clauses.append("id IN (" + ", ".join(["%s"] * len(detail_ids)) + ")")
        params.extend(detail_ids)
    details = fetch_all(
        f"""
        SELECT id, order_no, factory_order_no, seq, item_no, drawing_no, material_no, name,
               spec_model, material, quantity, unit_weight, heat_no, heat_treatment_batch_no
        FROM order_details
        WHERE {' AND '.join(clauses)}
        ORDER BY COALESCE(seq, 999999), COALESCE(item_no, 999999), id
        """,
        params,
    )
    if not details:
        if detail_ids:
            raise HTTPException(status_code=404, detail=f"订单号 {order_no} 下未找到所选订单明细")
        raise HTTPException(
            status_code=404,
            detail=f"订单号 {order_no} 下未找到物料号 {material_no} 的订单明细",
        )

    inch_by_dn = load_inch_map()
    typing_material_by_key = load_typing_material_map()
    paint_remarks = _load_paint_remarks(
        list({_normalize_str(d.get("drawing_no")) for d in details if _normalize_str(d.get("drawing_no"))})
    )

    inserted = 0
    updated = 0
    kept_manual: dict[str, list[str]] = {}
    for detail in details:
        generated = _build_generated_row(
            detail,
            inch_by_dn=inch_by_dn,
            typing_material_by_key=typing_material_by_key,
            paint_remarks=paint_remarks,
        )
        current = fetch_one(
            f"SELECT {SELECT_SQL} FROM typing_requirements WHERE order_detail_id=%s",
            (generated["order_detail_id"],),
        )
        snapshot_json = json.dumps(generated["source_snapshot"], ensure_ascii=False, default=str)
        if not current:
            new_id = execute(
                """
                INSERT INTO typing_requirements (
                    order_no, factory_order_no, order_detail_id, item_no, seq,
                    drawing_rev_no, material_no, part_name, typing_content,
                    spec, model_standard, qr_spec, heat_no,
                    heat_treatment_batch_no, quantity, unit_weight, remark, manual_fields, source_snapshot,
                    created_by, created_by_user_id, updated_by, updated_by_user_id
                ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                """,
                (
                    generated["order_no"],
                    generated["factory_order_no"] or None,
                    generated["order_detail_id"],
                    generated["item_no"],
                    generated["seq"],
                    generated["drawing_rev_no"],
                    generated["material_no"],
                    generated["part_name"],
                    generated["typing_content"],
                    generated["spec"] or None,
                    generated["model_standard"] or None,
                    generated["qr_spec"] or None,
                    generated["heat_no"],
                    generated["heat_treatment_batch_no"],
                    generated["quantity"],
                    generated["unit_weight"],
                    generated["remark"],
                    "[]",
                    snapshot_json,
                    user.name,
                    user.id,
                    user.name,
                    user.id,
                ),
            )
            inserted += 1
            _write_log(
                action="create",
                user=user,
                row_id=new_id,
                order_no=generated["order_no"],
                drawing_rev_no=generated["drawing_rev_no"],
                change_summary={"op": "generate_create", "fields": _snapshot(generated)},
            )
            continue

        before = _row_from_db(current)
        manual = [] if req.force_overwrite else before["manual_fields"]
        target = dict(generated)
        for key in manual:
            if key in EDITABLE_KEYS:
                target[key] = before.get(key)
        fields = _diff_fields(_snapshot(before), _snapshot(target))
        next_manual = [] if req.force_overwrite else manual
        if not fields and next_manual == before["manual_fields"]:
            execute(
                """
                UPDATE typing_requirements
                SET source_snapshot=%s, spec=%s, model_standard=%s, qr_spec=%s
                WHERE id=%s
                """,
                (
                    snapshot_json,
                    generated["spec"] or None,
                    generated["model_standard"] or None,
                    generated["qr_spec"] or None,
                    before["id"],
                ),
            )
            continue
        execute(
            """
            UPDATE typing_requirements
            SET order_no=%s, factory_order_no=%s, item_no=%s, seq=%s,
                drawing_rev_no=%s, material_no=%s, part_name=%s, typing_content=%s,
                spec=%s, model_standard=%s, qr_spec=%s, heat_no=%s,
                heat_treatment_batch_no=%s, quantity=%s, unit_weight=%s, remark=%s,
                manual_fields=%s, source_snapshot=%s,
                updated_by=%s, updated_by_user_id=%s, updated_at=NOW()
            WHERE id=%s
            """,
            (
                target["order_no"],
                target["factory_order_no"] or None,
                target["item_no"],
                target["seq"],
                target["drawing_rev_no"],
                target["material_no"],
                target["part_name"],
                target["typing_content"],
                target.get("spec") or None,
                target.get("model_standard") or None,
                target.get("qr_spec") or None,
                target["heat_no"],
                target["heat_treatment_batch_no"],
                target["quantity"],
                target["unit_weight"],
                target["remark"],
                json.dumps(next_manual, ensure_ascii=False),
                snapshot_json,
                user.name,
                user.id,
                before["id"],
            ),
        )
        updated += 1
        if manual:
            kept_manual[str(before["id"])] = manual
        _write_log(
            action="update",
            user=user,
            row_id=before["id"],
            order_no=target["order_no"],
            drawing_rev_no=target["drawing_rev_no"],
            change_summary={
                "op": "generate_overwrite" if req.force_overwrite else "generate_update",
                "fields": fields,
                "kept_manual_fields": manual,
            },
        )

    _write_log(
        action="generate",
        user=user,
        row_id=None,
        order_no=order_no,
        drawing_rev_no="",
        change_summary={
            "op": "generate",
            "order_no": order_no,
            "material_no": material_no,
            "order_detail_ids": detail_ids,
            "scope": "details" if detail_ids else ("material" if material_no else "order"),
            "force_overwrite": bool(req.force_overwrite),
            "detail_count": len(details),
            "inserted_count": inserted,
            "updated_count": updated,
            "kept_manual_fields": kept_manual,
        },
    )
    rows = _list_rows(order_no)
    return {
        "status": "success",
        "order_no": order_no,
        "material_no": material_no,
        "order_detail_ids": detail_ids,
        "scope": "details" if detail_ids else ("material" if material_no else "order"),
        "factory_order_no": _normalize_str(details[0].get("factory_order_no")),
        "inserted_count": inserted,
        "updated_count": updated,
        "detail_count": len(details),
        "total": len(rows),
        "rows": rows,
    }


@router.get("")
def list_typing_requirements(_user: UserDep, order_no: str = "", limit: int = Query(2000, ge=1, le=5000)):
    _ensure_tables()
    order_no = _normalize_str(order_no)
    rows = _list_rows(order_no, limit)
    factory_order_no = ""
    if rows:
        factory_order_no = _pick_first(*[r.get("factory_order_no") for r in rows])
    elif order_no:
        found = fetch_one(
            "SELECT factory_order_no FROM order_details WHERE order_no=%s AND factory_order_no IS NOT NULL LIMIT 1",
            (order_no,),
        )
        factory_order_no = _normalize_str((found or {}).get("factory_order_no"))
    return {
        "status": "success",
        "order_no": order_no,
        "factory_order_no": factory_order_no,
        "total": len(rows),
        "rows": rows,
    }


@router.get("/logs")
def list_typing_requirement_logs(
    _user: UserDep,
    order_no: str = "",
    action: str = "",
    limit: int = Query(200, ge=1, le=500),
):
    _ensure_tables()
    clauses = ["1=1"]
    params: list[Any] = []
    if _normalize_str(order_no):
        clauses.append("order_no LIKE %s")
        params.append(f"%{_normalize_str(order_no)}%")
    if _normalize_str(action):
        clauses.append("action=%s")
        params.append(_normalize_str(action))
    rows = fetch_all(
        f"""
        SELECT id, row_id, order_no, drawing_rev_no, action, operator_user_id,
               operator_name, operated_at, change_summary
        FROM typing_requirement_logs
        WHERE {' AND '.join(clauses)}
        ORDER BY id DESC
        LIMIT %s
        """,
        (*params, limit),
    )
    data = []
    for row in rows:
        item = {k: _jsonable(v) for k, v in row.items()}
        item["operator_name"] = _normalize_str(item.get("operator_name"))
        item["change_summary"] = item.get("change_summary") or ""
        data.append(item)
    return {"status": "success", "total": len(data), "rows": data}


@router.post("")
def create_typing_requirement(req: TypingRowCreate, user: UserDep):
    _ensure_tables()
    order_no = _normalize_str(req.order_no)
    drawing_rev_no = normalize_drawing_input(req.drawing_rev_no)
    found = fetch_one(
        "SELECT factory_order_no FROM order_details WHERE order_no=%s AND factory_order_no IS NOT NULL LIMIT 1",
        (order_no,),
    )
    payload = {
        "drawing_rev_no": drawing_rev_no,
        "material_no": _normalize_str(req.material_no),
        "part_name": _normalize_str(req.part_name) or default_part_name(order_no),
        "typing_content": _normalize_str(req.typing_content),
        "heat_no": _normalize_str(req.heat_no),
        "heat_treatment_batch_no": _normalize_str(req.heat_treatment_batch_no),
        "quantity": _as_int(req.quantity),
        "unit_weight": _as_unit_weight(req.unit_weight),
        "remark": _normalize_str(req.remark),
    }
    new_id = execute(
        """
        INSERT INTO typing_requirements (
            order_no, factory_order_no, order_detail_id, drawing_rev_no, material_no, part_name,
            typing_content, spec, model_standard, qr_spec,
            heat_no, heat_treatment_batch_no, quantity, unit_weight, remark,
            manual_fields, source_snapshot,
            created_by, created_by_user_id, updated_by, updated_by_user_id
        ) VALUES (%s, %s, NULL, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
        """,
        (
            order_no,
            _normalize_str((found or {}).get("factory_order_no")) or None,
            payload["drawing_rev_no"],
            payload["material_no"],
            payload["part_name"],
            payload["typing_content"],
            None,
            None,
            None,
            payload["heat_no"],
            payload["heat_treatment_batch_no"],
            payload["quantity"],
            payload["unit_weight"],
            payload["remark"],
            # 手工追加行没有 order_detail_id，生成时不会被匹配到，无需逐字段标人工
            "[]",
            json.dumps({"op": "manual_create"}, ensure_ascii=False),
            user.name,
            user.id,
            user.name,
            user.id,
        ),
    )
    row = _with_missing(_row_from_db(_fetch_row(new_id)))
    _write_log(
        action="create",
        user=user,
        row_id=new_id,
        order_no=order_no,
        drawing_rev_no=payload["drawing_rev_no"],
        change_summary={"op": "manual_create", "fields": payload},
    )
    return {"status": "success", "row": row}


@router.patch("/{row_id}")
def update_typing_requirement(row_id: int, req: TypingRowUpdate, user: UserDep):
    _ensure_tables()
    current = _fetch_row(row_id)
    if not current:
        raise HTTPException(status_code=404, detail="未找到打字要求记录")
    before = _row_from_db(current)

    target = dict(before)
    for key, value in req.model_dump(exclude_unset=True).items():
        if key not in EDITABLE_KEYS or value is None:
            continue
        if key == "quantity":
            target[key] = _as_int(value)
        elif key == "unit_weight":
            target[key] = _as_unit_weight(value)
        elif key == "drawing_rev_no":
            target[key] = normalize_drawing_input(value)
        else:
            target[key] = _normalize_str(value)

    fields = _diff_fields(_snapshot(before), _snapshot(target))
    if not fields:
        return {"status": "success", "row": _with_missing(before), "changed": False}

    manual = list(before["manual_fields"])
    for key in fields:
        if key not in manual:
            manual.append(key)

    execute(
        """
        UPDATE typing_requirements
        SET drawing_rev_no=%s, material_no=%s, part_name=%s, typing_content=%s, heat_no=%s,
            heat_treatment_batch_no=%s, quantity=%s, unit_weight=%s, remark=%s, manual_fields=%s,
            updated_by=%s, updated_by_user_id=%s, updated_at=NOW()
        WHERE id=%s
        """,
        (
            target["drawing_rev_no"],
            target["material_no"],
            target["part_name"],
            target["typing_content"],
            target["heat_no"],
            target["heat_treatment_batch_no"],
            target["quantity"],
            target["unit_weight"],
            target["remark"],
            json.dumps(manual, ensure_ascii=False),
            user.name,
            user.id,
            row_id,
        ),
    )
    row = _with_missing(_row_from_db(_fetch_row(row_id)))
    _write_log(
        action="update",
        user=user,
        row_id=row_id,
        order_no=before["order_no"],
        drawing_rev_no=target["drawing_rev_no"],
        change_summary={"op": "manual_update", "fields": fields, "manual_fields": manual},
    )
    return {"status": "success", "row": row, "changed": True}


@router.delete("/{row_id}")
def delete_typing_requirement(row_id: int, user: UserDep):
    _ensure_tables()
    current = _fetch_row(row_id)
    if not current:
        raise HTTPException(status_code=404, detail="未找到打字要求记录")
    before = _row_from_db(current)
    execute("DELETE FROM typing_requirements WHERE id=%s", (row_id,))
    _write_log(
        action="delete",
        user=user,
        row_id=row_id,
        order_no=before["order_no"],
        drawing_rev_no=before["drawing_rev_no"],
        change_summary={"op": "delete", "fields": _snapshot(before)},
    )
    return {"status": "success", "id": row_id}


@router.post("/delete-batch")
def delete_typing_requirements_batch(req: BatchDeleteRequest, user: UserDep):
    _ensure_tables()
    ids = sorted({int(item) for item in req.ids if int(item) > 0})
    if not ids:
        raise HTTPException(status_code=400, detail="请选择要删除的记录")
    placeholders = ", ".join(["%s"] * len(ids))
    existing = fetch_all(
        f"SELECT {SELECT_SQL} FROM typing_requirements WHERE id IN ({placeholders})",
        tuple(ids),
    )
    if not existing:
        raise HTTPException(status_code=404, detail="未找到要删除的打字要求记录")
    found_ids = [int(row["id"]) for row in existing]
    order_no = _normalize_str((existing[0] or {}).get("order_no"))
    for row in existing:
        before = _row_from_db(row)
        _write_log(
            action="delete",
            user=user,
            row_id=before["id"],
            order_no=before["order_no"],
            drawing_rev_no=before["drawing_rev_no"],
            change_summary={"op": "delete", "fields": _snapshot(before)},
        )
    found_placeholders = ", ".join(["%s"] * len(found_ids))
    execute(f"DELETE FROM typing_requirements WHERE id IN ({found_placeholders})", tuple(found_ids))
    _write_log(
        action="delete",
        user=user,
        row_id=None,
        order_no=order_no,
        drawing_rev_no="",
        change_summary={
            "op": "delete_batch",
            "order_no": order_no,
            "deleted_count": len(found_ids),
            "ids": found_ids,
        },
    )
    return {
        "status": "success",
        "deleted_count": len(found_ids),
        "deleted_ids": found_ids,
        "missing_count": len(ids) - len(found_ids),
    }


def _resolve_template() -> Path | None:
    path = get_templates_dir() / TEMPLATE_NAME
    return path if path.is_file() else None


def _copy_cell_style(src, dst) -> None:
    if not src.has_style:
        return
    dst.font = copy(src.font)
    dst.border = copy(src.border)
    dst.fill = copy(src.fill)
    dst.number_format = copy(src.number_format)
    dst.protection = copy(src.protection)
    dst.alignment = copy(src.alignment)


def _copy_row_style(ws, src_row: int, dst_row: int) -> None:
    for col in range(1, TEMPLATE_COL_COUNT + 1):
        _copy_cell_style(ws.cell(src_row, col), ws.cell(dst_row, col))
    height = ws.row_dimensions[src_row].height
    if height:
        ws.row_dimensions[dst_row].height = height


def _apply_export_column_widths(ws) -> None:
    for col, width in zip("ABCDEFGHI", EXPORT_COL_WIDTHS):
        ws.column_dimensions[col].width = width


def _insert_export_column(ws, col_idx: int, header: str) -> None:
    for merged in list(ws.merged_cells.ranges):
        if merged.min_row == 1:
            ws.unmerge_cells(str(merged))
    ws.insert_cols(col_idx)
    header_src = ws.cell(TEMPLATE_HEADER_ROW, max(col_idx - 1, 1))
    header_dst = ws.cell(TEMPLATE_HEADER_ROW, col_idx)
    header_dst.value = header
    _copy_cell_style(header_src, header_dst)
    for row_idx in range(TEMPLATE_FIRST_DATA_ROW, TEMPLATE_LAST_STYLED_ROW + 1):
        _copy_cell_style(ws.cell(row_idx, max(col_idx - 1, 1)), ws.cell(row_idx, col_idx))


def _header_col(ws, label: str) -> int | None:
    for col in range(1, TEMPLATE_COL_COUNT + 3):
        if _normalize_str(ws.cell(TEMPLATE_HEADER_ROW, col).value) == label:
            return col
    return None


def _ensure_export_template_layout(ws) -> None:
    """兼容旧模板：补「料号」「单重」，扩到 9 列。"""
    if _header_col(ws, "料号") is None:
        _insert_export_column(ws, 3, "料号")
    qty_col = _header_col(ws, "数量") or 7
    after_qty = _normalize_str(ws.cell(TEMPLATE_HEADER_ROW, qty_col + 1).value)
    if after_qty != "单重":
        _insert_export_column(ws, qty_col + 1, "单重")
    title_merged = False
    for merged in ws.merged_cells.ranges:
        if merged.min_row == 1 and merged.min_col == 1 and merged.max_col >= TEMPLATE_COL_COUNT:
            title_merged = True
            break
    if not title_merged:
        for merged in list(ws.merged_cells.ranges):
            if merged.min_row == 1:
                ws.unmerge_cells(str(merged))
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=TEMPLATE_COL_COUNT)
    for idx, (_key, label) in enumerate(EXPORT_COLUMNS, start=1):
        ws.cell(TEMPLATE_HEADER_ROW, idx).value = label
    _apply_export_column_widths(ws)


def _build_export_workbook(order_no: str, factory_order_no: str, rows: list[dict[str, Any]]):
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    template = _resolve_template()
    if template:
        wb = load_workbook(template)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "打字规则"
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=TEMPLATE_COL_COUNT)
        thin = Side(style="thin")
        header_font = Font(name="宋体", size=14, bold=True)
        header_fill = PatternFill(fill_type="solid", fgColor="FFFFFF00")
        ws.cell(1, 1).font = Font(name="宋体", size=22)
        for idx, (_key, label) in enumerate(EXPORT_COLUMNS, start=1):
            cell = ws.cell(TEMPLATE_HEADER_ROW, idx, label)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        _apply_export_column_widths(ws)

    _ensure_export_template_layout(ws)
    title = " ".join(part for part in ("订单号：" + factory_order_no if factory_order_no else "订单号：", order_no, "打字要求") if part)
    ws.cell(1, 1).value = title

    missing_fill = PatternFill(fill_type="solid", fgColor=MISSING_FILL_RGB)
    missing_cells = 0
    for offset, row in enumerate(rows):
        excel_row = TEMPLATE_FIRST_DATA_ROW + offset
        if excel_row > TEMPLATE_LAST_STYLED_ROW:
            _copy_row_style(ws, TEMPLATE_FIRST_DATA_ROW, excel_row)
        missing = set(row.get("missing") or [])
        for col, (key, _label) in enumerate(EXPORT_COLUMNS, start=1):
            cell = ws.cell(excel_row, col)
            value = row.get(key)
            if key == "quantity":
                cell.value = _as_int(value)
            elif key == "unit_weight":
                cell.value = _as_unit_weight(value)
            else:
                cell.value = _normalize_str(value)
            if key in missing:
                cell.fill = missing_fill
                missing_cells += 1

    # 清掉模板预铺样式行里多出来的旧内容
    last_used = TEMPLATE_FIRST_DATA_ROW + len(rows) - 1
    for excel_row in range(max(last_used + 1, TEMPLATE_FIRST_DATA_ROW), TEMPLATE_LAST_STYLED_ROW + 1):
        for col in range(1, TEMPLATE_COL_COUNT + 1):
            ws.cell(excel_row, col).value = None

    return wb, missing_cells


@router.get("/export")
def export_typing_requirements(
    user: UserDep,
    order_no: str = Query(..., min_length=1),
    ids: list[int] = Query(default=[]),
):
    import io

    from fastapi.responses import StreamingResponse
    from urllib.parse import quote

    _ensure_tables()
    order_no = _normalize_str(order_no)
    rows = _list_rows(order_no)
    if not rows:
        raise HTTPException(status_code=404, detail=f"订单号 {order_no} 还没有打字要求，请先点「生成」")
    id_set = {int(item) for item in ids if int(item) > 0}
    if id_set:
        rows = [row for row in rows if int(row.get("id") or 0) in id_set]
        if not rows:
            raise HTTPException(status_code=404, detail="所选范围内没有可导出的打字要求")
    factory_order_no = _pick_first(*[r.get("factory_order_no") for r in rows])

    wb, missing_cells = _build_export_workbook(order_no, factory_order_no, rows)
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)

    safe_order = re.sub(r"[\\/:*?\"<>|]", "_", order_no)
    file_name = f"打字要求_{safe_order}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    _write_log(
        action="export",
        user=user,
        row_id=None,
        order_no=order_no,
        drawing_rev_no="",
        change_summary={
            "op": "export",
            "order_no": order_no,
            "factory_order_no": factory_order_no,
            "file_name": file_name,
            "row_count": len(rows),
            "missing_cell_count": missing_cells,
            "ids": sorted(id_set),
        },
    )
    return StreamingResponse(
        bio,
        media_type=MEDIA_XLSX,
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(file_name)}"},
    )
