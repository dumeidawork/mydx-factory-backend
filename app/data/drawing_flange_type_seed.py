"""图纸-法兰类型预置对照：优先读已抽出常量，否则解析附件 Excel。"""
from __future__ import annotations

import re
from pathlib import Path

from app.utils.drawing_no import drawing_no_base

FLANGE_TYPES = ("滑动", "平板", "盲板", "高径")
_FLANGE_ALIASES = {
    "滑动": "滑动",
    "so": "滑动",
    "slip": "滑动",
    "slip-on": "滑动",
    "slip on": "滑动",
    "slipon": "滑动",
    "sop": "滑动",
    "soh": "滑动",
    "平板": "平板",
    "pl": "平板",
    "plate": "平板",
    "flat": "平板",
    "盲板": "盲板",
    "bl": "盲板",
    "blind": "盲板",
    "高径": "高径",
    "高颈": "高径",
    "wn": "高径",
    "lwn": "高径",
    "weld neck": "高径",
    "weldneck": "高径",
    "welding neck": "高径",
}


def _normalize_flange_type(raw: str) -> str:
    text = (raw or "").strip()
    if not text:
        return ""
    if text in FLANGE_TYPES:
        return text
    compact = re.sub(r"[\s_\-]+", " ", text).strip().lower().replace("/", " ")
    if compact in _FLANGE_ALIASES:
        return _FLANGE_ALIASES[compact]
    first = compact.split()[0] if compact else ""
    return _FLANGE_ALIASES.get(first, "")

EXCEL_CANDIDATES = (
    Path(r"c:\Users\dumei\OneDrive\デスクトップ\铭远数字化\00-前期要件\大连+法国1-3号更新-1xlsx(3)(1).xlsx"),
)
_FALLBACK_CANDIDATES = (
    Path(__file__).resolve().parents[3] / "00-temp" / "dalian_price_split.xlsx",
    Path(__file__).resolve().parents[3] / "00-temp" / "france_price_split.xlsx",
)

DRAWING_HEADERS = (
    "图纸号",
    "drawing no and version no",
    "drawing no and version",
    "drawing no,",
    "drawing no",
    "drawing",
)
TYPE_HEADERS = ("法兰类型", "flange type", "类型", "type", "备注")


def _norm_header(value: object) -> str:
    text = str(value or "").replace("\n", " ").replace("\r", " ").lower()
    return re.sub(r"\s+", " ", text).strip().strip(":")


def _pick_col(headers: dict[int, str], aliases: tuple[str, ...]) -> int | None:
    for alias in aliases:
        for col, name in headers.items():
            if name == alias or name.startswith(alias):
                return col
    return None


def _row_index(row) -> int | None:
    for cell in row:
        idx = getattr(cell, "row", None)
        if idx:
            return int(idx)
    return None


def _extract_file(path: Path) -> list[tuple[str, str, str]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    rows_out: list[tuple[str, str, str]] = []
    try:
        for sheet in wb.worksheets:
            header_row = None
            headers: dict[int, str] = {}
            for row in sheet.iter_rows(min_row=1, max_row=20, values_only=False):
                current = {
                    cell.column: _norm_header(cell.value)
                    for cell in row
                    if cell.value not in (None, "")
                }
                if _pick_col(current, DRAWING_HEADERS) and _pick_col(current, TYPE_HEADERS):
                    header_row = _row_index(row)
                    headers = current
                    break
            if not header_row:
                continue
            drawing_col = _pick_col(headers, DRAWING_HEADERS)
            type_col = _pick_col(headers, TYPE_HEADERS)
            if not drawing_col or not type_col:
                continue
            for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
                if not row:
                    continue
                drawing_raw = row[drawing_col - 1] if drawing_col - 1 < len(row) else None
                type_raw = row[type_col - 1] if type_col - 1 < len(row) else None
                drawing_no = drawing_no_base(str(drawing_raw or "").strip())
                flange_type = _normalize_flange_type(str(type_raw or "").strip())
                if not drawing_no or not flange_type:
                    continue
                if drawing_no.lower() in {"drawing", "drawing no", "图纸号"}:
                    continue
                rows_out.append((drawing_no, flange_type, ""))
    finally:
        wb.close()
    return rows_out


def _dedupe(rows: list[tuple[str, str, str]]) -> tuple[tuple[str, str, str], ...]:
    by_key: dict[str, tuple[str, str, str]] = {}
    conflicts: dict[str, set[str]] = {}
    for drawing_no, flange_type, remark in rows:
        key = drawing_no.upper()
        if key not in by_key:
            by_key[key] = (drawing_no, flange_type, remark)
            conflicts[key] = {flange_type}
            continue
        conflicts[key].add(flange_type)
    out: list[tuple[str, str, str]] = []
    for key, (drawing_no, flange_type, remark) in by_key.items():
        types = conflicts[key]
        note = remark
        if len(types) > 1:
            extra = "、".join(sorted(types))
            note = f"附件中存在多种类型：{extra}"
        out.append((drawing_no, flange_type, note))
    out.sort(key=lambda item: item[0].upper())
    return tuple(out)


def _collect_existing(paths: tuple[Path, ...]) -> list[Path]:
    found: list[Path] = []
    seen: set[str] = set()
    for path in paths:
        try:
            key = str(path.resolve()).lower()
        except OSError:
            key = str(path).lower()
        if key in seen:
            continue
        try:
            if path.is_file():
                seen.add(key)
                found.append(path)
        except OSError:
            continue
    return found


def _iter_excel_files() -> list[Path]:
    primary = _collect_existing(EXCEL_CANDIDATES)
    if primary:
        return primary
    return _collect_existing(_FALLBACK_CANDIDATES)


def write_precomputed(rows: tuple[tuple[str, str, str], ...]) -> Path:
    dest = Path(__file__).resolve().parent / "drawing_flange_type_seed_data.py"
    lines = [
        '"""已核对的图纸-法兰类型预置对照（由当前库导出）。"""',
        "",
        "PRECOMPUTED_MAPS = (",
    ]
    for drawing_no, flange_type, remark in rows:
        lines.append(f"    ({drawing_no!r}, {flange_type!r}, {remark!r}),")
    lines.append(")")
    lines.append("")
    dest.write_text("\n".join(lines), encoding="utf-8")
    return dest


def _write_precomputed(rows: tuple[tuple[str, str, str], ...]) -> None:
    write_precomputed(rows)


def load_default_maps() -> tuple[tuple[str, str, str], ...]:
    try:
        from app.data.drawing_flange_type_seed_data import PRECOMPUTED_MAPS

        if PRECOMPUTED_MAPS:
            return PRECOMPUTED_MAPS
    except ImportError:
        pass
    collected: list[tuple[str, str, str]] = []
    for path in _iter_excel_files():
        collected.extend(_extract_file(path))
    rows = _dedupe(collected)
    if rows:
        try:
            _write_precomputed(rows)
        except OSError:
            pass
    return rows
