"""解析图纸尺寸表 xlsx：活跃 sheet、双行表头、稀疏字段。"""
from __future__ import annotations

import re
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.utils.drawing_no import looks_like_drawing_rev, normalize_drawing_rev_for_db
from app.utils.drawing_source_fields import (
    BOM_HEADER_MARKERS,
    DIM_FIELDS,
    DIM_LABELS,
    match_header,
    normalize_header,
)
from app.utils.spec_parser import extract_spec_model_fields

_MAX_COLS = 40
_MAX_ROWS = 80
_STOP_PREFIXES = ("注意", "交货期", "合计", "制图", "审核", "技术要求")
_DUAL_HINTS = frozenset(
    {
        normalize_header("外径"),
        normalize_header("内径"),
        normalize_header("盘厚"),
        normalize_header("总厚"),
        normalize_header("上口"),
        normalize_header("孔径"),
        normalize_header("孔数"),
    }
)
_DRAWING_NO_RE = re.compile(r"图纸号\s*[：:]\s*(.+)$")
_MATERIAL_RE = re.compile(r"材质\s*[：:]\s*(.+)$")


class DrawingSourceParseError(ValueError):
    """尺寸表无法解析。"""


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _filled_grid(ws: Worksheet) -> list[list[str]]:
    max_row = min(ws.max_row or 1, _MAX_ROWS)
    max_col = min(ws.max_column or 1, _MAX_COLS)
    grid = [[_cell_text(ws.cell(r, c).value) for c in range(1, max_col + 1)] for r in range(1, max_row + 1)]
    for merged in ws.merged_cells.ranges:
        if merged.min_row > max_row or merged.min_col > max_col:
            continue
        origin = grid[merged.min_row - 1][merged.min_col - 1]
        r1 = max(merged.min_row, 1)
        r2 = min(merged.max_row, max_row)
        c1 = max(merged.min_col, 1)
        c2 = min(merged.max_col, max_col)
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                if not grid[r - 1][c - 1]:
                    grid[r - 1][c - 1] = origin
    return grid


def _is_dual_row(cells: list[str]) -> bool:
    norms = [normalize_header(c) for c in cells if c]
    if not norms:
        return False
    hits = sum(1 for n in norms if n in _DUAL_HINTS)
    return hits >= 2


def _is_header_row(cells: list[str]) -> bool:
    blob = "".join(cells)
    return "规格" in blob and sum(1 for c in cells if c) >= 4


def _looks_like_bom(header_cells: list[str]) -> bool:
    norms = {normalize_header(c) for c in header_cells if c}
    return BOM_HEADER_MARKERS.issubset(norms)


def _stop_row(cells: list[str]) -> bool:
    first = next((c for c in cells if c), "")
    compact = re.sub(r"\s+", "", first)
    return any(compact.startswith(p) for p in _STOP_PREFIXES)


def _scan_meta(grid: list[list[str]]) -> tuple[str, str]:
    drawing_no = ""
    material = ""
    for row in grid:
        for cell in row:
            if not drawing_no:
                m = _DRAWING_NO_RE.search(cell.replace("\n", " "))
                if m:
                    drawing_no = m.group(1).strip()
            if not material:
                m = _MATERIAL_RE.search(cell.replace("\n", " "))
                if m:
                    material = m.group(1).strip()
            if drawing_no and material:
                return drawing_no, material
    return drawing_no, material


def _map_columns(header_cells: list[str]) -> list[str | None]:
    return [match_header(cell) if cell else None for cell in header_cells]


def parse_drawing_source_xlsx(
    content: bytes,
    *,
    sheet_name: str | None = None,
) -> dict[str, Any]:
    try:
        wb = load_workbook(BytesIO(content), data_only=True, read_only=False)
    except Exception as exc:
        raise DrawingSourceParseError("无法读取 Excel 文件") from exc

    visible = [name for name in wb.sheetnames if wb[name].sheet_state == "visible"]
    if not visible:
        visible = list(wb.sheetnames)
    if not visible:
        raise DrawingSourceParseError("工作簿没有工作表")

    chosen = (sheet_name or "").strip()
    if chosen:
        if chosen not in wb.sheetnames:
            raise DrawingSourceParseError(f"找不到工作表：{chosen}")
        ws = wb[chosen]
    else:
        active = wb.active
        active_name = getattr(active, "title", None)
        if active_name in wb.sheetnames:
            ws = wb[active_name]
            chosen = active_name
        else:
            chosen = visible[0]
            ws = wb[chosen]

    grid = _filled_grid(ws)
    drawing_raw, material = _scan_meta(grid)

    header_idx = None
    for i, row in enumerate(grid):
        if _is_header_row(row):
            header_idx = i
            break
    if header_idx is None:
        raise DrawingSourceParseError("未找到尺寸表表头（需含「规格」）")

    header_cells = list(grid[header_idx])
    if _looks_like_bom(header_cells):
        raise DrawingSourceParseError("这是物料清单，不是尺寸表")

    data_start = header_idx + 1
    if data_start < len(grid) and _is_dual_row(grid[data_start]):
        sub = grid[data_start]
        merged: list[str] = []
        for i, top in enumerate(header_cells):
            bottom = sub[i] if i < len(sub) else ""
            merged.append(" ".join(p for p in (top, bottom) if p))
        header_cells = merged
        data_start += 1

    col_keys = _map_columns(header_cells)
    if "spec_model" not in col_keys:
        raise DrawingSourceParseError("表头未识别到规格型号列")

    used_headers: list[dict[str, str]] = []
    seen_keys: set[str] = set()
    for i, key in enumerate(col_keys):
        if not key or key in seen_keys:
            continue
        seen_keys.add(key)
        used_headers.append(
            {
                "key": key,
                "label": DIM_LABELS.get(key, key),
                "source": header_cells[i],
            }
        )

    rows: list[dict[str, str]] = []
    for row in grid[data_start:]:
        if not any(row):
            if rows:
                break
            continue
        if _stop_row(row):
            break
        rec: dict[str, str] = {}
        for i, key in enumerate(col_keys):
            if not key or i >= len(row):
                continue
            val = row[i].replace("\n", " ").strip()
            if not val:
                continue
            rec.setdefault(key, val)
        spec = rec.get("spec_model", "").strip()
        if not spec:
            continue
        if spec.startswith("注意"):
            break
        parsed = extract_spec_model_fields(spec)
        item = {k: rec[k] for k in rec if rec[k]}
        if parsed["material_no"] and "material_no" not in item:
            item["material_no"] = parsed["material_no"]
        rows.append(item)

    if not rows:
        raise DrawingSourceParseError("表头下没有规格数据行")

    drawing_rev_no = ""
    standard_no = ""
    if drawing_raw:
        if looks_like_drawing_rev(drawing_raw):
            drawing_rev_no = normalize_drawing_rev_for_db(drawing_raw)
        else:
            standard_no = drawing_raw

    return {
        "sheet": chosen,
        "sheet_names": visible,
        "drawing_rev_no": drawing_rev_no,
        "standard_no": standard_no,
        "material": material,
        "headers": used_headers,
        "rows": rows,
        "dim_fields": list(DIM_FIELDS),
    }
