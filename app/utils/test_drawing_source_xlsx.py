"""drawing_source_xlsx 解析单测。运行：python -m unittest app.utils.test_drawing_source_xlsx"""
from __future__ import annotations

import unittest
from io import BytesIO

from openpyxl import Workbook

from app.utils.drawing_source_fields import match_header
from app.utils.drawing_source_xlsx import DrawingSourceParseError, parse_drawing_source_xlsx


def _xlsx_bytes(build) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    build(ws)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


class MatchHeaderTests(unittest.TestCase):
    def test_side_hole_variants(self):
        self.assertEqual(match_header("侧孔A深度最小"), "side_a_depth")
        self.assertEqual(match_header("侧孔B"), "side_b")
        self.assertEqual(match_header("侧孔"), "side_hole")
        self.assertEqual(match_header("侧孔深度 最小"), "side_hole_depth")

    def test_od_and_seat(self):
        self.assertEqual(match_header("D外径（+-0.5）"), "od_d")
        self.assertEqual(match_header("A外径（+0.3）"), "od_d")
        self.assertEqual(match_header("d +-0.5 RF台径"), "seat_a")
        self.assertIsNone(match_header("件数"))


class ParseSourceXlsxTests(unittest.TestCase):
    def test_simple_sorf(self):
        def build(ws):
            ws["A9"] = "订单号：法国侧孔"
            ws["A10"] = "图纸号：A5E02970610A Rev015"
            ws["A11"] = "材质：A105"
            headers = [
                "规格型号",
                "D外径（+0.5）",
                "C内径（+1.5-0）",
                "T厚度（+0.5）",
                "K中心距（+-1）",
                "A台径（+0.8）",
                "F台高（+-0.2）",
                "H孔径（+1.5-0）",
                "N孔数",
                "侧孔B",
                "侧孔B深度",
                "单重",
                "数量(件)",
            ]
            for i, h in enumerate(headers, 1):
                ws.cell(12, i, h)
            values = [
                "FDK-115-518 AS4087 PN16 DN450 SO RF",
                640,
                461,
                30,
                584,
                552,
                4,
                26,
                12,
                "M6",
                "14-20",
                32.35,
                10,
            ]
            for i, v in enumerate(values, 1):
                ws.cell(13, i, v)
            ws["A14"] = "注意：侧孔M6"

        result = parse_drawing_source_xlsx(_xlsx_bytes(build))
        self.assertEqual(result["drawing_rev_no"], "A5E02970610A Rev015")
        self.assertEqual(result["material"], "A105")
        self.assertEqual(len(result["rows"]), 1)
        row = result["rows"][0]
        self.assertTrue(row["spec_model"].startswith("FDK-115-518"))
        self.assertEqual(row["material_no"], "FDK-115-518")
        self.assertEqual(row["od_d"], "640")
        self.assertEqual(row["side_b"], "M6")
        self.assertNotIn("qty", row)
        keys = {h["key"] for h in result["headers"]}
        self.assertNotIn("unit_weight" if False else "hole_n", set())  # keep linter calm
        self.assertIn("unit_weight", keys)
        self.assertNotIn(None, keys)

    def test_soff_through_hole(self):
        def build(ws):
            ws["A9"] = "图纸号：A5E03496033A Rev011"
            ws["A11"] = "材质：A105"
            headers = [
                "规格型号",
                "A外径（+0.3）",
                "D小内径（+1-0）",
                "E大内径（+0.8-0）",
                "B厚度（+0.3-0）",
                "S11中心距（+-1）",
                "G（+-0.5）",
                "F深度（-1）",
                "S9孔径（+-0.2）",
                "S10孔数",
                "通孔C（+0.3-0）",
                "V1°",
                "V2°",
                "侧孔",
                "侧孔深度最小",
                "单重",
                "数量(件)",
            ]
            for i, h in enumerate(headers, 1):
                ws.cell(12, i, h)
            values = [
                "FDK-083N5035 ASME B16.5 CL150 DN300 SO FF",
                483,
                306.5,
                326.6,
                42,
                431.8,
                12.5,
                13,
                25.9,
                12,
                10,
                "15°",
                "30°",
                "M6",
                16,
                34.1,
                20,
            ]
            for i, v in enumerate(values, 1):
                ws.cell(13, i, v)

        row = parse_drawing_source_xlsx(_xlsx_bytes(build))["rows"][0]
        self.assertEqual(row["od_d"], "483")
        self.assertEqual(row["id_small_d"], "306.5")
        self.assertEqual(row["id_large_e"], "326.6")
        self.assertEqual(row["thru_c"], "10")
        self.assertEqual(row["side_hole"], "M6")
        self.assertEqual(row["g_val"], "12.5")
        self.assertEqual(row["v1_deg"], "15°")

    def test_wn_dual_row(self):
        def build(ws):
            ws["A10"] = "图纸号：A5E45038171A Rev005"
            ws["A11"] = "材质：A105"
            l1 = [
                "规格",
                "D +-0.5",
                "B+-0.5",
                "C +0.5",
                "H +-0.5",
                "A+0.5-0",
                "N+-0.5",
                "K +-0.8",
                "d +-0.5",
                "f1 -0.5",
                "R2 +-1",
                "E ≧",
                "L +1",
                "n",
                "侧孔",
                "侧孔深度",
                "焊口  θ",
                "R",
                "单重",
                "件数",
            ]
            l2 = [
                "",
                "外径",
                "内径",
                "盘厚",
                "总厚",
                "上口",
                "下口",
                "中心距",
                "RF台径",
                "台高",
                "",
                "",
                "孔径",
                "孔数",
                "",
                "",
                "+-2.5°",
                "",
                "",
                "",
            ]
            for i, h in enumerate(l1, 1):
                ws.cell(12, i, h)
            for i, h in enumerate(l2, 1):
                ws.cell(13, i, h)
            values = [
                "A5E45033247 ASME B16.5 CL600 DN40 WN RF",
                155,
                40.94,
                29.3,
                77,
                48.3,
                70,
                114.3,
                73,
                7,
                5,
                6,
                22.2,
                4,
                "M6",
                "16-20",
                "37.5°",
                3,
                3.4,
                4,
            ]
            for i, v in enumerate(values, 1):
                ws.cell(14, i, v)

        parsed = parse_drawing_source_xlsx(_xlsx_bytes(build))
        row = parsed["rows"][0]
        self.assertEqual(row["od_d"], "155")
        self.assertEqual(row["id_c"], "40.94")
        self.assertEqual(row["thk_t"], "29.3")
        self.assertEqual(row["thk_f"], "77")
        self.assertEqual(row["mouth_m"], "48.3")
        self.assertEqual(row["mouth_n"], "70")
        self.assertEqual(row["seat_a"], "73")
        self.assertEqual(row["hole_h"], "22.2")
        self.assertEqual(row["material_no"], "A5E45033247")
        self.assertEqual(row["side_hole"], "M6")
        self.assertEqual(row["weld_theta"], "37.5°")
        self.assertEqual(row["r2_val"], "5")
        self.assertEqual(row["e_val"], "6")

    def test_pcd2(self):
        def build(ws):
            ws["A10"] = "图纸号：A5E35929718A Rev010"
            ws["A11"] = "材质：A105"
            headers = [
                "规格型号",
                "D外径（+-0.5）",
                "C内径（+1-0）",
                "T厚度（+-0.5）",
                "K中心距（+-1）",
                "A台径（+-0.8）",
                "F台高（-0.5）",
                "H1孔径（+2.5）",
                "N1孔数",
                "PCD2中心距（+-1）",
                "H2孔径（+0.3）",
                "N2孔数",
                "侧孔B",
                "侧孔B深度",
                "单重",
                "数量(件)",
            ]
            for i, h in enumerate(headers, 1):
                ws.cell(12, i, h)
            values = [
                "FDK-083N4864 EN1092-1 DN250 PN25 SO RF",
                425,
                275.5,
                34.5,
                370,
                335,
                3,
                30,
                12,
                365,
                10,
                4,
                "M6",
                "14-20",
                18.23,
                6,
            ]
            for i, v in enumerate(values, 1):
                ws.cell(13, i, v)

        row = parse_drawing_source_xlsx(_xlsx_bytes(build))["rows"][0]
        self.assertEqual(row["pcd2"], "365")
        self.assertEqual(row["hole_h2"], "10")
        self.assertEqual(row["hole_n2"], "4")
        self.assertEqual(row["hole_h"], "30")

    def test_ff_bottom_hole(self):
        def build(ws):
            ws["A10"] = "图纸号：A5E52565483A Rev001"
            ws["A11"] = "材质：A105"
            headers = [
                "规格型号",
                "D外径（+-0.5）",
                "C内径（+1-0）",
                "T厚度（+-0.5）",
                "K中心距（+-1）",
                "H孔径（+1-0）",
                "孔数",
                "侧孔A",
                "侧孔A深度最小",
                "底孔（+-1）",
                "单重",
                "数量(件)",
            ]
            for i, h in enumerate(headers, 1):
                ws.cell(12, i, h)
            values = [
                "A5E02531260 EN1092-1 DN125 PN16 SO FF",
                250,
                142.5,
                24,
                210,
                18,
                8,
                "M6",
                "12",
                "16",
                5.4,
                10,
            ]
            for i, v in enumerate(values, 1):
                ws.cell(13, i, v)

        row = parse_drawing_source_xlsx(_xlsx_bytes(build))["rows"][0]
        self.assertEqual(row["side_a"], "M6")
        self.assertEqual(row["bottom_hole"], "16")
        self.assertEqual(row["material_no"], "A5E02531260")

    def test_bom_rejected(self):
        def build(ws):
            for i, h in enumerate(["序号", "名称", "图纸号", "物料号", "规格型号", "材质", "数量"], 1):
                ws.cell(3, i, h)
            ws["A4"] = 1
            ws["B4"] = "法兰"
            ws["E4"] = "EN1092"

        with self.assertRaises(DrawingSourceParseError) as ctx:
            parse_drawing_source_xlsx(_xlsx_bytes(build))
        self.assertIn("物料清单", str(ctx.exception))


if __name__ == "__main__":
    unittest.main()
