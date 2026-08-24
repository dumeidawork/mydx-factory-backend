"""文档模板占位符渲染服务。"""
from __future__ import annotations

import re
import tempfile
import zipfile
from copy import copy
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from docx import Document
from docx.oxml.ns import qn

_LEGACY_OOXML_PREFIX = "http://purl.oclc.org/ooxml/"
_OOXML_URI_REPLACEMENTS = (
    ("http://purl.oclc.org/ooxml/wordprocessingml/main", "http://schemas.openxmlformats.org/wordprocessingml/2006/main"),
    ("http://purl.oclc.org/ooxml/drawingml/main", "http://schemas.openxmlformats.org/drawingml/2006/main"),
    ("http://purl.oclc.org/ooxml/drawingml/picture", "http://schemas.openxmlformats.org/drawingml/2006/picture"),
    (
        "http://purl.oclc.org/ooxml/drawingml/wordprocessingDrawing",
        "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    ),
    ("http://purl.oclc.org/ooxml/officeDocument/math", "http://schemas.openxmlformats.org/officeDocument/2006/math"),
    (
        "http://purl.oclc.org/ooxml/officeDocument/relationships",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    ),
)


def _render_text(text: str, mapping: dict[str, Any]) -> str:
    out = text
    for k, v in mapping.items():
        value = str(v if v is not None else "")
        out = out.replace(f"{{{{{k}}}}}", value)
        out = re.sub(r"\{\{\s*" + re.escape(k) + r"\s*\}\}", value, out)
    return out


def render_placeholder_text(text: str, mapping: dict[str, Any]) -> str:
    """占位符替换（Word / PDF 模板共用）。"""
    return _render_text(text, mapping)


def _replace_paragraph_text(paragraph, mapping: dict[str, Any]) -> None:
    original = paragraph.text
    if "{{" not in original:
        return
    rendered = _render_text(original, mapping)
    if rendered == original:
        return
    if paragraph.runs:
        paragraph.runs[0].text = rendered
        for i in range(1, len(paragraph.runs)):
            paragraph.runs[i].text = ""
    else:
        paragraph.text = rendered


def _replace_docx_in_paragraphs(doc: Document, mapping: dict[str, Any]) -> None:
    for p in doc.paragraphs:
        _replace_paragraph_text(p, mapping)


def _normalize_ooxml_xml(xml_bytes: bytes) -> bytes:
    try:
        text = xml_bytes.decode("utf-8")
    except UnicodeDecodeError:
        return xml_bytes
    if _LEGACY_OOXML_PREFIX not in text:
        return xml_bytes
    for legacy, standard in _OOXML_URI_REPLACEMENTS:
        text = text.replace(legacy, standard)
    return text.encode("utf-8")


def _ensure_standard_docx(template_path: Path) -> tuple[Path, bool]:
    """将非标准 OOXML 命名空间转为 python-docx 可识别的格式，必要时写临时文件。"""
    if not template_path.exists():
        return template_path, False
    with zipfile.ZipFile(template_path, "r") as zin:
        fixed: dict[str, bytes] = {}
        needs_fix = False
        for name in zin.namelist():
            original = zin.read(name)
            if not (name.endswith(".xml") or name.endswith(".rels")):
                continue
            normalized = _normalize_ooxml_xml(original)
            fixed[name] = normalized
            if normalized != original:
                needs_fix = True
        if not needs_fix:
            return template_path, False
        tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
        tmp_path = Path(tmp.name)
        tmp.close()
        with zipfile.ZipFile(tmp_path, "w", compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                zout.writestr(item, fixed.get(item.filename, zin.read(item.filename)))
    return tmp_path, True


def _replace_docx_in_tables(doc: Document, mapping: dict[str, Any]) -> None:
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _replace_paragraph_text(p, mapping)
                # 合并单元格时 inner table
                for inner in cell.tables:
                    for inner_row in inner.rows:
                        for inner_cell in inner_row.cells:
                            for p in inner_cell.paragraphs:
                                _replace_paragraph_text(p, mapping)


def _replace_docx_in_story(story, mapping: dict[str, Any]) -> None:
    for p in story.paragraphs:
        _replace_paragraph_text(p, mapping)
    for table in story.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _replace_paragraph_text(p, mapping)


def _sectpr_ref_types(sect_pr, local_name: str) -> set[str]:
    types: set[str] = set()
    if sect_pr is None:
        return types
    for el in sect_pr:
        if el.tag == qn(f"w:{local_name}"):
            types.add(el.get(qn("w:type")) or "default")
    return types


def _replace_docx_headers_footers(doc: Document, mapping: dict[str, Any]) -> None:
    """只替换模板里已有的页眉/页脚，避免访问时新建 footer 把单页撑破。"""
    seen: set[int] = set()

    def visit(story) -> None:
        part = getattr(story, "part", None)
        marker = id(part if part is not None else story)
        if marker in seen:
            return
        seen.add(marker)
        _replace_docx_in_story(story, mapping)

    for section in doc.sections:
        header_types = _sectpr_ref_types(section._sectPr, "headerReference")
        footer_types = _sectpr_ref_types(section._sectPr, "footerReference")
        if header_types:
            visit(section.header)
            if "first" in header_types:
                visit(section.first_page_header)
            if "even" in header_types:
                visit(section.even_page_header)
        if footer_types:
            visit(section.footer)
            if "first" in footer_types:
                visit(section.first_page_footer)
            if "even" in footer_types:
                visit(section.even_page_footer)


def render_docx_template(
    template_path: Path,
    output_path: Path,
    mapping: dict[str, Any],
    append_blocks: list[str] | None = None,
) -> None:
    resolved_path, is_temp = (
        _ensure_standard_docx(template_path)
        if template_path.exists()
        else (template_path, False)
    )
    try:
        doc = Document(str(resolved_path)) if template_path.exists() else Document()
        _replace_docx_in_paragraphs(doc, mapping)
        _replace_docx_in_tables(doc, mapping)
        _replace_docx_headers_footers(doc, mapping)
        if append_blocks:
            for block in append_blocks:
                doc.add_paragraph(block)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))
    finally:
        if is_temp:
            resolved_path.unlink(missing_ok=True)


def _replace_xlsx_scalars(ws: Worksheet, mapping: dict[str, Any]) -> None:
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and "{{" in cell.value and "}}" in cell.value:
                cell.value = _render_text(cell.value, mapping)


def _find_item_template_row(ws: Worksheet) -> int | None:
    for row_idx in range(1, ws.max_row + 1):
        for col_idx in range(1, ws.max_column + 1):
            value = ws.cell(row_idx, col_idx).value
            if isinstance(value, str) and "{{item." in value:
                return row_idx
    return None


def _render_item_row_values(template_values: list[Any], item: dict[str, Any]) -> list[Any]:
    row_values: list[Any] = []
    for v in template_values:
        if isinstance(v, str):
            text = v
            for k, vv in item.items():
                text = text.replace(f"{{{{item.{k}}}}}", str(vv if vv is not None else ""))
            row_values.append(text)
        else:
            row_values.append(v)
    return row_values


def _copy_cell_style(source, dest) -> None:
    dest.font = copy(source.font)
    dest.border = copy(source.border)
    dest.fill = copy(source.fill)
    dest.alignment = copy(source.alignment)
    dest.number_format = source.number_format
    dest.protection = copy(source.protection)


def _expand_item_template_rows(ws: Worksheet, template_row: int, items: list[dict[str, Any]]) -> None:
    """在明细模板行下方插入行，复制边框/字体/行高，并把下方合并区整体下移。"""
    max_col = max(ws.max_column, 1)
    template_values = [ws.cell(template_row, c).value for c in range(1, max_col + 1)]
    template_height = ws.row_dimensions[template_row].height
    extra = len(items) - 1
    if extra > 0:
        insert_at = template_row + 1
        original_merges = [
            (rng.min_row, rng.min_col, rng.max_row, rng.max_col) for rng in ws.merged_cells.ranges
        ]
        original_heights = {idx: ws.row_dimensions[idx].height for idx in range(1, ws.max_row + 1)}
        for rng in list(ws.merged_cells.ranges):
            ws.unmerge_cells(str(rng))
        ws.insert_rows(insert_at, extra)
        for idx, height in original_heights.items():
            dest = idx + extra if idx >= insert_at else idx
            ws.row_dimensions[dest].height = height
        for min_row, min_col, max_row, max_col in original_merges:
            if min_row >= insert_at:
                min_row += extra
                max_row += extra
            elif max_row >= insert_at:
                max_row += extra
            ws.merge_cells(start_row=min_row, start_column=min_col, end_row=max_row, end_column=max_col)
        for offset in range(1, extra + 1):
            dest_row = template_row + offset
            ws.row_dimensions[dest_row].height = template_height
            for col in range(1, max_col + 1):
                _copy_cell_style(ws.cell(template_row, col), ws.cell(dest_row, col))
    for index, item in enumerate(items):
        values = _render_item_row_values(template_values, item)
        dest_row = template_row + index
        for col, value in enumerate(values, start=1):
            ws.cell(dest_row, col).value = value


def render_xlsx_template(
    template_path: Path,
    output_path: Path,
    mapping: dict[str, Any],
    items: list[dict[str, Any]] | None = None,
    fallback_headers: list[str] | None = None,
) -> None:
    wb = load_workbook(template_path) if template_path.exists() else Workbook()
    ws = wb.active
    _replace_xlsx_scalars(ws, mapping)

    item_template_row = _find_item_template_row(ws)
    if items:
        if item_template_row:
            _expand_item_template_rows(ws, item_template_row, items)
        else:
            # 没有item占位行则追加标准表格
            start_row = ws.max_row + 2
            if fallback_headers:
                for i, h in enumerate(fallback_headers, start=1):
                    ws.cell(start_row, i).value = h
                start_row += 1
            for item in items:
                for i, value in enumerate(item.values(), start=1):
                    ws.cell(start_row, i).value = value
                start_row += 1

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
