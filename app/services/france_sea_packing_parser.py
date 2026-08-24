"""法国客户 SEA 箱单 Excel 解析（只读，不落盘）。"""
from __future__ import annotations

import io
import re
from datetime import date, datetime
from typing import Any

try:
    from openpyxl import load_workbook
except Exception:  # pragma: no cover
    load_workbook = None

REV_RE = re.compile(r"(?i)\s*Rev\s*(\d+)")
NA_VALUES = {"", "#N/A", "N/A", "NA", "NONE", "-"}


def _normalize_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    if text.upper() in NA_VALUES:
        return ""
    return text


def _safe_int(value: Any, default: int | None = 0) -> int | None:
    text = _normalize_str(value)
    if text == "":
        return default
    try:
        return int(float(text))
    except (TypeError, ValueError):
        return default


def _safe_float(value: Any, default: float = 0.0) -> float:
    text = _normalize_str(value)
    if text == "":
        return default
    try:
        return float(text)
    except (TypeError, ValueError):
        return default


def _round2(value: float) -> float:
    return round(float(value or 0), 2)


def _cell_date(value: Any) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = _normalize_str(value)
    if not text:
        return ""
    compact = text.replace("/", "-").replace(".", "-")
    if len(compact) == 8 and compact.isdigit():
        return f"{compact[:4]}-{compact[4:6]}-{compact[6:8]}"
    return compact[:10]


def dn_key(value: Any) -> str:
    text = _normalize_str(value).upper().replace(" ", "")
    text = text.replace("DN", "").replace("NB", "")
    match = re.match(r"(\d+(?:\.\d+)?)", text)
    if not match:
        return text
    number = float(match.group(1))
    if number.is_integer():
        return str(int(number))
    return match.group(1)


def drawing_rev_number(drawing_no: str) -> int:
    match = REV_RE.search(drawing_no or "")
    if not match:
        return -1
    try:
        return int(match.group(1))
    except ValueError:
        return -1


def calc_total_weight(unit_weight: float, quantity: int) -> float:
    return _round2(float(unit_weight or 0) * int(quantity or 0))


def _pick_sheet(wb: Any):
    for name in wb.sheetnames:
        if str(name).strip().upper().startswith("CONTAINER"):
            return wb[name]
    return wb[wb.sheetnames[0]]


def parse_sea_workbook(content: bytes) -> dict[str, Any]:
    if load_workbook is None:
        raise RuntimeError("缺少openpyxl依赖，无法读取Excel。")
    wb = load_workbook(io.BytesIO(content), data_only=True)
    ws = _pick_sheet(wb)
    departure = _cell_date(ws["C1"].value)
    arrival = _cell_date(ws["C2"].value)
    items: list[dict[str, Any]] = []
    seq = 0
    max_row = int(ws.max_row or 0)
    for row_no in range(4, max_row + 1):
        material_no = _normalize_str(ws.cell(row_no, 1).value)
        po = _normalize_str(ws.cell(row_no, 8).value)
        if not material_no:
            continue
        seq += 1
        quantity = int(_safe_int(ws.cell(row_no, 4).value, 0) or 0)
        unit_weight = _round2(_safe_float(ws.cell(row_no, 6).value, 0.0))
        excel_total = _round2(_safe_float(ws.cell(row_no, 7).value, 0.0))
        total_weight = calc_total_weight(unit_weight, quantity)
        weight_mismatch = abs(excel_total - total_weight) > 0.05 if excel_total else 0
        items.append(
            {
                "seq": seq,
                "material_no": material_no,
                "dn": _normalize_str(ws.cell(row_no, 2).value),
                "units_per_pallet": _safe_float(ws.cell(row_no, 3).value, 0.0),
                "quantity": quantity,
                "nb_of_pallets": _safe_float(ws.cell(row_no, 5).value, 0.0),
                "unit_weight": unit_weight,
                "total_weight": total_weight,
                "excel_total_weight": excel_total,
                "weight_mismatch": 1 if weight_mismatch else 0,
                "po": po,
                "pos": _safe_int(ws.cell(row_no, 9).value, None),
            }
        )
    if not items:
        raise ValueError("读取失败：未找到有效明细行（Item 物料号为空或整表为空）。")
    pos = sorted({_normalize_str(r.get("po")) for r in items if _normalize_str(r.get("po"))})
    return {
        "departure_date": departure,
        "arrival_date": arrival,
        "items": items,
        "po_list": pos,
    }
