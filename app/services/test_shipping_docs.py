import tempfile
import unittest
import zipfile
import xml.etree.ElementTree as ET
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.core.paths import get_templates_dir
from app.services.license_service import TRANSPORT_AIR, TRANSPORT_SEA
from app.services.shipping_docs_service import (
    format_shipping_date,
    plan_shipping_jobs,
    render_shipping_job,
)


def _customer(material_no, material, qty, unit_weight, tax_unit, tax_total):
    return {
        "material_no": material_no,
        "material": material,
        "quantity": qty,
        "unit_weight": unit_weight,
        "product_unit_price": tax_unit,
        "tax_unit_price": tax_unit,
        "tax_total_amount": tax_total,
    }


def _box(box_no, material_no, material, qty, unit_weight, length, width, height):
    return {
        "box_no": box_no,
        "material_no": material_no,
        "material": material,
        "quantity": qty,
        "unit_weight": unit_weight,
        "box_length": length,
        "box_width": width,
        "box_height": height,
    }


CUSTOMER = [
    _customer("FDK-A", "A105", 10, "7.48", "117.35", "1173.50"),
    _customer("FDK-B", "A105", 4, "3.10", "20.00", "80.00"),
    _customer("FDK-S", "304", 4, "4.86", "240.41", "961.64"),
]
PACKING = [
    _box(1, "FDK-A", "A105", 10, "7.48", 1200, 800, 200),
    _box(1, "FDK-B", "A105", 4, "3.10", 1200, 800, 200),
    _box(2, "FDK-S", "304", 4, "4.86", 250, 250, 90),
]


class ShippingDocsTest(unittest.TestCase):
    def test_date_has_no_zero_padding(self):
        self.assertEqual(format_shipping_date(date(2026, 5, 5)), "2026-5-5")

    def test_sea_filenames_and_totals(self):
        jobs = plan_shipping_jobs("2000980249", TRANSPORT_SEA, date(2026, 5, 5), CUSTOMER, PACKING)
        names = [job["filename"] for job in jobs]
        self.assertEqual(
            names,
            [
                "2000980249 invoice（CS）.xlsx",
                "2000980249-2 invoice（SS）.xlsx",
                "2000980249Packing list.xlsx",
                "2000980249截单数据.xlsx",
            ],
        )
        invoice = jobs[0]["mapping"]
        self.assertEqual(invoice["invoice_no"], "XMZ2000980249")
        self.assertEqual(invoice["box_summary"], "1 boxes,0.19CBM")
        self.assertEqual(invoice["weight_sum"], "87.20")
        self.assertEqual(invoice["gross_weight"], "107.20")
        self.assertEqual(invoice["qty_sum"], "14")
        packing = jobs[2]
        self.assertEqual(packing["mapping"]["invoice_no_cell"], "INVOICE NO.：XMZ2000980249 / XMZ2000980249-2")
        self.assertEqual(packing["group_lengths"], [2, 1])
        cutoff = jobs[3]["mapping"]
        self.assertEqual(cutoff["box_range"], "1-2")
        self.assertEqual(cutoff["cs_cbm"], "0.19")
        self.assertEqual(cutoff["ss_cbm"], "0.01")
        self.assertEqual(cutoff["ss_qty"], "4")

    def test_air_splits_packing_list(self):
        jobs = plan_shipping_jobs("2000988002", TRANSPORT_AIR, date(2026, 6, 25), CUSTOMER, PACKING)
        self.assertEqual(
            [job["filename"] for job in jobs],
            [
                "2000988002 invoice.xlsx",
                "2000988002Packing list.xlsx",
                "2000988002-2 invoice.xlsx",
                "2000988002-2Packing list.xlsx",
            ],
        )

    def test_mixed_grade_box_is_rejected(self):
        packing = [
            _box(1, "FDK-A", "A105", 10, "7.48", 100, 100, 100),
            _box(1, "FDK-S", "304", 4, "4.86", 100, 100, 100),
        ]
        with self.assertRaisesRegex(ValueError, "拆箱"):
            plan_shipping_jobs("2000980249", TRANSPORT_SEA, date(2026, 5, 5), CUSTOMER, packing)

    def test_box_numbers_must_be_continuous(self):
        packing = [
            _box(1, "FDK-A", "A105", 10, "7.48", 100, 100, 100),
            _box(3, "FDK-S", "304", 4, "4.86", 100, 100, 100),
        ]
        with self.assertRaisesRegex(ValueError, "连续编号"):
            plan_shipping_jobs("2000980249", TRANSPORT_SEA, date(2026, 5, 5), CUSTOMER, packing)

    def test_duplicate_unit_weight_is_rejected(self):
        customer = [
            _customer("FDK-A", "A105", 10, "7.48", "1", "10"),
            _customer("FDK-A", "A105", 2, "8.00", "1", "2"),
        ]
        packing = [_box(1, "FDK-A", "A105", 12, "7.48", 100, 100, 100)]
        with self.assertRaisesRegex(ValueError, "单重不一致"):
            plan_shipping_jobs("2000980249", TRANSPORT_SEA, date(2026, 5, 5), customer, packing)

    def test_render_keeps_stamp_and_merges_commodity(self):
        jobs = plan_shipping_jobs("2000980249", TRANSPORT_SEA, date(2026, 5, 5), CUSTOMER, PACKING)
        invoice = jobs[0]
        template = get_templates_dir() / "20-FR-xkz" / "03-shipping" / "BySea" / "invoice.xlsx"
        before = load_workbook(template)
        origin = int(before.active._images[0].anchor._from.row)
        before.close()
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / invoice["filename"]
            render_shipping_job(invoice, target)
            wb = load_workbook(target)
            ws = wb.active
            self.assertEqual(len(ws._images), 1)
            self.assertEqual(int(ws._images[0].anchor._from.row), origin + 1)
            self.assertEqual(ws["F7"].value, "ORDER NO.:2000980249")
            self.assertEqual(ws["F8"].value, "INVOICE NO.:XMZ2000980249")
            self.assertEqual(ws["G9"].value, "DATE:2026-5-5")
            self.assertIn("CARBON STEEL FLANGE", str(ws["B13"].value))
            self.assertIsNone(ws["B14"].value)
            merged = {str(item) for item in ws.merged_cells.ranges}
            self.assertIn("B13:B14", merged)
            self.assertNotIn("{{", str(ws["A13"].value))
            _assert_saved_detail_style(self, target, 13, 14, {2})
            wb.close()

        packing = jobs[2]
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / packing["filename"]
            render_shipping_job(packing, target)
            wb = load_workbook(target)
            ws = wb.active
            self.assertEqual(ws["A12"].value, "1")
            self.assertIsNone(ws["A13"].value)
            self.assertEqual(ws["B13"].value, "FDK-B")
            self.assertEqual(ws["H12"].value, "107.20")
            self.assertIsNone(ws["H13"].value)
            merged = {str(item) for item in ws.merged_cells.ranges}
            self.assertIn("A12:A13", merged)
            self.assertIn("H12:H13", merged)
            self.assertEqual(len(ws._images), 1)
            _assert_saved_detail_style(self, target, 12, 13, {1, 4, 5, 6, 8})
            wb.close()


def _names(order_no, transport, customer, packing, when=date(2026, 5, 5)):
    jobs = plan_shipping_jobs(order_no, transport, when, customer, packing)
    return [job["filename"] for job in jobs], jobs


_SHEET_NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def _border_style(cell, side):
    edge = getattr(cell.border, side)
    return None if edge is None else edge.style


def _saved_sheet_cells(path):
    archive = zipfile.ZipFile(path)
    sheet = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
    styles = ET.fromstring(archive.read("xl/styles.xml"))
    fonts = list(styles.find("m:fonts", _SHEET_NS))
    borders = list(styles.find("m:borders", _SHEET_NS))
    xfs = list(styles.find("m:cellXfs", _SHEET_NS))
    cells = {}
    heights = {}
    for row in sheet.find("m:sheetData", _SHEET_NS):
        heights[row.attrib.get("r")] = row.attrib.get("ht")
        for cell in row:
            xf = xfs[int(cell.attrib.get("s", "0"))]
            font = fonts[int(xf.attrib["fontId"])]
            border = borders[int(xf.attrib["borderId"])]
            align = xf.find("m:alignment", _SHEET_NS)
            edges = {}
            for edge_name in ("left", "right", "top", "bottom"):
                node = border.find(f"m:{edge_name}", _SHEET_NS)
                edges[edge_name] = None if node is None else node.attrib.get("style")
            cells[cell.attrib["r"]] = {
                "font": font.find("m:name", _SHEET_NS).attrib.get("val"),
                "size": font.find("m:sz", _SHEET_NS).attrib.get("val"),
                "bold": font.find("m:b", _SHEET_NS) is not None,
                "horizontal": None if align is None else align.attrib.get("horizontal"),
                "vertical": None if align is None else align.attrib.get("vertical"),
                "number_format": xf.attrib.get("numFmtId"),
                "borders": edges,
            }
    return cells, heights


def _assert_saved_detail_style(test, path, first_row, extra_row, merged_columns):
    cells, heights = _saved_sheet_cells(path)
    test.assertEqual(heights[str(extra_row)], heights[str(first_row)])
    for col in range(1, 9):
        letter = get_column_letter(col)
        first = cells[f"{letter}{first_row}"]
        cell = cells[f"{letter}{extra_row}"]
        test.assertEqual(cell["font"], first["font"], letter)
        test.assertEqual(cell["size"], first["size"], letter)
        test.assertEqual(cell["bold"], first["bold"], letter)
        test.assertEqual(cell["horizontal"], first["horizontal"], letter)
        test.assertEqual(cell["vertical"], first["vertical"], letter)
        test.assertEqual(cell["number_format"], first["number_format"], letter)
        for side in ("left", "right", "top", "bottom"):
            if col in merged_columns and side in ("top", "bottom"):
                continue
            test.assertEqual(cell["borders"][side], "thin", f"{letter}{extra_row}.{side}")
            test.assertEqual(first["borders"][side], "thin", f"{letter}{first_row}.{side}")
        if col in merged_columns:
            test.assertIsNone(first["borders"]["bottom"], f"{letter}{first_row}.bottom")
            test.assertEqual(first["borders"]["top"], "thin", f"{letter}{first_row}.top")
            test.assertIsNone(cell["borders"]["top"], f"{letter}{extra_row}.top")
            test.assertEqual(cell["borders"]["bottom"], "thin", f"{letter}{extra_row}.bottom")


def _assert_filled(ws):
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and "{{" in str(cell.value):
                raise AssertionError(f"{cell.coordinate} 仍有占位符 {cell.value}")


class ShippingDocsCaseTest(unittest.TestCase):
    def test_rejects_missing_source_and_bad_transport(self):
        when = date(2026, 5, 5)
        with self.assertRaisesRegex(ValueError, "空运或海运"):
            plan_shipping_jobs("2000980249", "truck", when, CUSTOMER, PACKING)
        with self.assertRaisesRegex(ValueError, "没有法国客户箱单"):
            plan_shipping_jobs("2000980249", TRANSPORT_SEA, when, [], PACKING)
        with self.assertRaisesRegex(ValueError, "还没有箱号"):
            plan_shipping_jobs("2000980249", TRANSPORT_SEA, when, CUSTOMER, [])

    def test_rejects_bad_rows(self):
        when = date(2026, 5, 5)
        customer = [_customer("", "A105", 1, "1", "1", "1")]
        packing = [_box(1, "FDK-A", "A105", 1, "1", 10, 10, 10)]
        with self.assertRaisesRegex(ValueError, "空物料号"):
            plan_shipping_jobs("2000980249", TRANSPORT_SEA, when, customer, packing)

        customer = [_customer("FDK-A", "Q235", 1, "1", "1", "1")]
        with self.assertRaisesRegex(ValueError, "无法归类"):
            plan_shipping_jobs("2000980249", TRANSPORT_SEA, when, customer, packing)

        customer = [_customer("FDK-A", "A105", 1, "1", "1", "1")]
        packing = [_box(1, "", "A105", 1, "1", 10, 10, 10)]
        with self.assertRaisesRegex(ValueError, "装箱表存在空物料号"):
            plan_shipping_jobs("2000980249", TRANSPORT_SEA, when, customer, packing)

        packing = [_box(1, "FDK-X", "A105", 1, "1", 10, 10, 10)]
        with self.assertRaisesRegex(ValueError, "不在法国客户箱单"):
            plan_shipping_jobs("2000980249", TRANSPORT_SEA, when, customer, packing)

        packing = [_box(0, "FDK-A", "A105", 1, "1", 10, 10, 10)]
        with self.assertRaisesRegex(ValueError, "箱号无效"):
            plan_shipping_jobs("2000980249", TRANSPORT_SEA, when, customer, packing)

        packing = [_box(2, "FDK-A", "A105", 1, "1", 10, 10, 10)]
        with self.assertRaisesRegex(ValueError, "连续编号"):
            plan_shipping_jobs("2000980249", TRANSPORT_SEA, when, customer, packing)

    def test_316_is_stainless_and_316_wins_over_304(self):
        customer = [_customer("FDK-316", "316L", 2, "1.25", "8", "16")]
        packing = [_box(1, "FDK-316", "", 2, "0", 1000, 1000, 195)]
        names, jobs = _names("2000988002", TRANSPORT_AIR, customer, packing, date(2026, 10, 8))
        self.assertEqual(names, ["2000988002-2 invoice.xlsx", "2000988002-2Packing list.xlsx"])
        invoice = jobs[0]
        self.assertEqual(invoice["mapping"]["batch_date"], "2026-10-8")
        self.assertEqual(invoice["mapping"]["invoice_no"], "XMZ2000988002-2")
        self.assertEqual(invoice["mapping"]["total_label"], "Total 01 of  STAINLESS STEEL FLANGE")
        self.assertIn("HS Code:73072100", invoice["mapping"]["commodity"])
        self.assertEqual(invoice["items"][0]["unit_weight"], "1.25")
        self.assertEqual(invoice["items"][0]["total_weight"], "2.50")
        self.assertEqual(invoice["mapping"]["box_summary"], "1 boxes,0.20CBM")
        self.assertEqual(jobs[1]["items"][0]["net_weight"], "2.50")
        self.assertEqual(jobs[1]["items"][0]["gross_weight"], "22.50")

        mixed_name = [_customer("FDK-M", "304/316", 1, "1", "1", "1")]
        packing = [_box(1, "FDK-M", "304/316", 1, "1", 10, 10, 10)]
        names, _jobs = _names("2000980249", TRANSPORT_SEA, mixed_name, packing)
        self.assertEqual(names[0], "2000980249-2 invoice（SS）.xlsx")

    def test_carbon_only_and_stainless_only_file_sets(self):
        carbon = [_customer("FDK-A", "A105", 3, "2", "10", "30")]
        carbon_box = [_box(1, "FDK-A", "A105", 3, "2", 100, 100, 100)]
        names, jobs = _names("2000980249", TRANSPORT_SEA, carbon, carbon_box)
        self.assertEqual(
            names,
            ["2000980249 invoice（CS）.xlsx", "2000980249Packing list.xlsx", "2000980249截单数据.xlsx"],
        )
        self.assertEqual(jobs[1]["mapping"]["invoice_no_cell"], "INVOICE NO.：XMZ2000980249")
        cutoff = jobs[2]["mapping"]
        self.assertEqual(cutoff["cs_qty"], "3")
        self.assertEqual(cutoff["ss_qty"], "0")
        self.assertEqual(cutoff["ss_net"], "0.00")
        self.assertEqual(cutoff["ss_amount"], "0.00")
        self.assertEqual(cutoff["ss_cbm"], "0.00")
        self.assertEqual(cutoff["boxes_sum"], "1")

        stainless = [_customer("FDK-S", "304", 5, "1", "4", "20")]
        stainless_box = [_box(1, "FDK-S", "304", 5, "1", 100, 100, 100)]
        names, jobs = _names("2000980249", TRANSPORT_SEA, stainless, stainless_box)
        self.assertEqual(
            names,
            ["2000980249-2 invoice（SS）.xlsx", "2000980249Packing list.xlsx", "2000980249截单数据.xlsx"],
        )
        self.assertEqual(jobs[1]["mapping"]["invoice_no_cell"], "INVOICE NO.：XMZ2000980249-2")
        self.assertEqual(jobs[2]["mapping"]["cs_boxes"], "0")
        self.assertEqual(jobs[2]["mapping"]["cs_gross"], "0.00")

        names, _jobs = _names("2000988002", TRANSPORT_AIR, carbon, carbon_box)
        self.assertEqual(names, ["2000988002 invoice.xlsx", "2000988002Packing list.xlsx"])
        names, _jobs = _names("2000988002", TRANSPORT_AIR, stainless, stainless_box)
        self.assertEqual(names, ["2000988002-2 invoice.xlsx", "2000988002-2Packing list.xlsx"])

    def test_merged_invoice_line_split_boxes_and_tax_fallback(self):
        customer = [
            _customer("FDK-A", "A105", 10, "1.00", "10", "100"),
            _customer("FDK-A", "A105", 10, "1.00", "20", "200"),
            _customer("FDK-S", "304", 2, "3.00", "0", "0"),
        ]
        customer[2]["product_unit_price"] = "10.50"
        customer[2]["tax_unit_price"] = "0"
        customer[2]["tax_total_amount"] = "0"
        packing = [
            _box(2, "FDK-S", "", 2, "0", 500, 400, 300),
            _box(1, "FDK-A", "A105", 12, "1.00", 1000, 500, 200),
            _box(3, "FDK-A", "A105", 8, "1.00", 1000, 500, 200),
        ]
        _names_ignored, jobs = _names("2000980249", TRANSPORT_SEA, customer, packing)
        cs = jobs[0]
        self.assertEqual([item["material_no"] for item in cs["items"]], ["FDK-A"])
        self.assertEqual(cs["items"][0]["qty"], "20")
        self.assertEqual(cs["items"][0]["tax_unit_price"], "15.00")
        self.assertEqual(cs["items"][0]["tax_total_amount"], "300.00")
        self.assertEqual(cs["items"][0]["total_weight"], "20.00")
        self.assertEqual(cs["mapping"]["box_summary"], "2 boxes,0.20CBM")
        ss = jobs[1]
        self.assertEqual(ss["items"][0]["tax_unit_price"], "10.50")
        self.assertEqual(ss["items"][0]["tax_total_amount"], "21.00")
        self.assertEqual(ss["items"][0]["line_no"], "1")
        packing_job = jobs[2]
        self.assertEqual(
            [item["material_no"] for item in packing_job["items"]],
            ["FDK-A", "FDK-S", "FDK-A"],
        )
        self.assertEqual([item["box_no"] for item in packing_job["items"]], ["1", "2", "3"])
        self.assertEqual(packing_job["items"][1]["net_weight"], "6.00")
        self.assertEqual(packing_job["items"][1]["gross_weight"], "26.00")
        self.assertEqual(packing_job["group_lengths"], [1, 1, 1])

    def test_unpacked_material_keeps_invoice_with_zero_boxes(self):
        customer = [
            _customer("FDK-A", "A105", 6, "2.50", "8", "48"),
            _customer("FDK-S", "304", 1, "1", "1", "1"),
        ]
        packing = [_box(1, "FDK-S", "304", 1, "1", 200, 200, 200)]
        _ignored, jobs = _names("2000988002", TRANSPORT_AIR, customer, packing)
        carbon = jobs[0]["mapping"]
        self.assertEqual(carbon["qty_sum"], "6")
        self.assertEqual(carbon["weight_sum"], "15.00")
        self.assertEqual(carbon["box_summary"], "0 boxes,0.00CBM")
        self.assertEqual(carbon["gross_weight"], "15.00")
        self.assertEqual(jobs[1]["items"], [])
        self.assertEqual(jobs[1]["mapping"]["qty_sum"], "0")
        self.assertEqual(jobs[1]["mapping"]["gross_sum"], "0.00")

    def test_three_box_volume_rounds_after_sum(self):
        customer = [_customer("FDK-A", "A105", 9, "1", "1", "9")]
        packing = [
            _box(1, "FDK-A", "A105", 3, "1", 1200, 800, 200),
            _box(2, "FDK-A", "A105", 3, "1", 1200, 800, 200),
            _box(3, "FDK-A", "A105", 3, "1", 1200, 800, 200),
        ]
        _ignored, jobs = _names("2000980249", TRANSPORT_SEA, customer, packing)
        self.assertEqual(jobs[0]["mapping"]["box_summary"], "3 boxes,0.58CBM")
        self.assertEqual(jobs[1]["mapping"]["gross_sum"], "69.00")

    def test_render_all_templates(self):
        sea = plan_shipping_jobs("2000980249", TRANSPORT_SEA, date(2026, 5, 5), CUSTOMER, PACKING)
        air = plan_shipping_jobs("2000988002", TRANSPORT_AIR, date(2026, 6, 25), CUSTOMER, PACKING)
        carbon = [_customer("FDK-A", "A105", 1, "1", "1", "1")]
        only_box = [_box(1, "FDK-A", "A105", 1, "1", 10, 10, 10)]
        sea_cs = plan_shipping_jobs("2000981001", TRANSPORT_SEA, date(2026, 12, 9), carbon, only_box)
        air_empty = plan_shipping_jobs(
            "2000988111",
            TRANSPORT_AIR,
            date(2026, 6, 25),
            [
                _customer("FDK-A", "A105", 1, "1", "1", "1"),
                _customer("FDK-S", "304", 1, "1", "1", "1"),
            ],
            [_box(1, "FDK-S", "304", 1, "1", 10, 10, 10)],
        )
        with tempfile.TemporaryDirectory() as folder:
            folder_path = Path(folder)
            rendered = {}
            books = []
            for job in sea + air + sea_cs + air_empty:
                target = folder_path / job["filename"]
                render_shipping_job(job, target)
                wb = load_workbook(target)
                _assert_filled(wb.active)
                books.append(wb)
                rendered[job["filename"]] = wb.active

            sea_invoice = rendered["2000980249 invoice（CS）.xlsx"]
            self.assertEqual(sea_invoice["A8"].value, "Port of Destination:Antwerp Port, Belgium")
            self.assertEqual(sea_invoice["A13"].value, "1")
            self.assertEqual(sea_invoice["C13"].value, "FDK-A")
            self.assertEqual(sea_invoice["A14"].value, "2")
            self.assertEqual(sea_invoice["A15"].value, "Total 01 of  CARBON STEEL FLANGE")
            self.assertEqual(sea_invoice["B16"].value, "1 boxes,0.19CBM")
            self.assertEqual(sea_invoice["F17"].value, "107.20")

            ss_invoice = rendered["2000980249-2 invoice（SS）.xlsx"]
            self.assertEqual(ss_invoice["F8"].value, "INVOICE NO.:XMZ2000980249-2")
            self.assertEqual(ss_invoice["C13"].value, "FDK-S")
            self.assertEqual(ss_invoice["B15"].value, "1 boxes,0.01CBM")
            self.assertEqual(ss_invoice["F16"].value, "39.44")
            self.assertNotIn("B13:B14", {str(item) for item in ss_invoice.merged_cells.ranges})
            for col in range(1, 9):
                cell = ss_invoice.cell(13, col)
                for side in ("left", "right", "top", "bottom"):
                    self.assertEqual(_border_style(cell, side), "thin", f"{cell.coordinate}.{side}")

            cutoff = rendered["2000980249截单数据.xlsx"]
            self.assertEqual(cutoff["A2"].value, "2000980249")
            self.assertEqual(cutoff["B2"].value, "1-2")
            self.assertEqual(cutoff["C2"].value, "碳钢法兰")
            self.assertEqual(cutoff["C3"].value, "不锈钢法兰")
            self.assertEqual(cutoff["D2"].value, "14")
            self.assertEqual(cutoff["D3"].value, "4")
            self.assertEqual(cutoff["B4"].value, "2")
            self.assertEqual(cutoff["H4"].value, "2215.14")
            self.assertIn("7307910000", str(cutoff["A5"].value))
            self.assertIn("7307210000", str(cutoff["A6"].value))
            self.assertIn("碳钢制", str(cutoff["A9"].value))
            self.assertIn("不锈钢制", str(cutoff["A11"].value))

            air_invoice = rendered["2000988002 invoice.xlsx"]
            self.assertEqual(air_invoice["G7"].value, "ORDER NO.:2000988002")
            self.assertEqual(air_invoice["A8"].value, "Port of Destination: SXB(STRASBOURG),FRANCE")
            self.assertEqual(air_invoice["G9"].value, "DATE:2026-6-25")
            air_ss_pack = rendered["2000988002-2Packing list.xlsx"]
            self.assertEqual(air_ss_pack["G7"].value, "ORDER NO.:2000988002")
            self.assertEqual(air_ss_pack["F8"].value, "INVOICE NO.：XMZ2000988002-2")
            self.assertIn("STRASBOURG", str(air_ss_pack["A8"].value))
            self.assertEqual(air_ss_pack["A12"].value, "2")

            cs_only = rendered["2000981001截单数据.xlsx"]
            self.assertEqual(cs_only["D3"].value, "0")
            self.assertEqual(cs_only["E3"].value, "0")
            self.assertEqual(cs_only["F3"].value, "0.00")
            self.assertEqual(cs_only["I3"].value, "0.00")
            self.assertEqual(cs_only["I2"].value, "0.00")
            self.assertEqual(rendered["2000981001 invoice（CS）.xlsx"]["G9"].value, "DATE:2026-12-9")

            empty_pack = rendered["2000988111Packing list.xlsx"]
            self.assertIsNone(empty_pack["A12"].value)
            self.assertIsNone(empty_pack["B12"].value)
            self.assertEqual(empty_pack["C13"].value, "0")
            self.assertEqual(empty_pack["A13"].value, "合计")
            for book in books:
                book.close()


if __name__ == "__main__":
    unittest.main()
