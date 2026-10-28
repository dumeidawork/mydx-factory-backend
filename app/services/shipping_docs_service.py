"""法国出货资料：商业发票、出口装箱单、海运截单。"""
from __future__ import annotations

import os
import shutil
import tempfile
import zipfile
from collections import OrderedDict
from copy import copy
from datetime import date
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Side

from app.core.database import fetch_all
from app.core.paths import get_templates_dir
from app.services.license_service import (
    GRADE_A105,
    GRADE_304,
    GRADE_316,
    TRANSPORT_AIR,
    TRANSPORT_SEA,
    _po_variants,
    classify_grade,
)
from app.services.template_renderer import (
    _expand_item_template_rows,
    _find_item_template_row,
    _replace_xlsx_scalars,
)

CLASS_CS = "CS"
CLASS_SS = "SS"
BOX_TARE_KG = Decimal("20")
_Q2 = Decimal("0.01")
_CBM_DIVISOR = Decimal("1000000000")

COMMODITY = {
    CLASS_CS: "CARBON STEEL FLANGE\nHS Code:73079100",
    CLASS_SS: "STAINLESS STEEL FLANGE\nHS Code:73072100",
}
TOTAL_LABEL = {
    CLASS_CS: "Total 01 of  CARBON STEEL FLANGE",
    CLASS_SS: "Total 01 of  STAINLESS STEEL FLANGE",
}


def format_shipping_date(value: date) -> str:
    return f"{value.year}-{value.month}-{value.day}"


def material_class_of(material: str) -> str | None:
    grade = classify_grade(material)
    if grade == GRADE_A105:
        return CLASS_CS
    if grade in (GRADE_304, GRADE_316):
        return CLASS_SS
    return None


def invoice_no_for(order_no: str, material_class: str) -> str:
    if material_class == CLASS_SS:
        return f"XMZ{order_no}-2"
    return f"XMZ{order_no}"


def invoice_filename(transport: str, order_no: str, material_class: str) -> str:
    if transport == TRANSPORT_AIR:
        if material_class == CLASS_SS:
            return f"{order_no}-2 invoice.xlsx"
        return f"{order_no} invoice.xlsx"
    if material_class == CLASS_SS:
        return f"{order_no}-2 invoice（SS）.xlsx"
    return f"{order_no} invoice（CS）.xlsx"


def packing_filename(transport: str, order_no: str, material_class: str | None) -> str:
    if transport == TRANSPORT_AIR and material_class == CLASS_SS:
        return f"{order_no}-2Packing list.xlsx"
    return f"{order_no}Packing list.xlsx"


def _q2(value: Decimal) -> Decimal:
    return value.quantize(_Q2, rounding=ROUND_HALF_UP)


def _fmt2(value: Decimal) -> str:
    return f"{_q2(value):.2f}"


def _dec(value: Any) -> Decimal:
    if value is None or value == "":
        return Decimal("0")
    return Decimal(str(value))


def _qty(value: Any) -> int:
    return int(_dec(value))


def _tax_pair(row: dict[str, Any]) -> tuple[Decimal, Decimal]:
    qty = _qty(row.get("quantity"))
    product = _dec(row.get("product_unit_price"))
    tax_unit = _dec(row.get("tax_unit_price"))
    tax_total = _dec(row.get("tax_total_amount"))
    if tax_unit == 0 and product != 0:
        tax_unit = _q2(product)
    if tax_total == 0 and tax_unit != 0 and qty:
        tax_total = _q2(tax_unit * Decimal(qty))
    return _q2(tax_unit), _q2(tax_total)


def _class_of_row(row: dict[str, Any], customer_by_no: dict[str, dict[str, Any]]) -> str | None:
    material_no = str(row.get("material_no") or "")
    customer = customer_by_no.get(material_no)
    if customer:
        found = material_class_of(str(customer.get("material") or ""))
        if found:
            return found
    return material_class_of(str(row.get("material") or ""))


def aggregate_invoice_lines(rows: list[dict[str, Any]], material_class: str) -> list[dict[str, Any]]:
    grouped: OrderedDict[str, dict[str, Any]] = OrderedDict()
    for row in rows:
        if _class_of_row(row, {}) != material_class:
            continue
        material_no = str(row.get("material_no") or "").strip()
        if not material_no:
            raise ValueError("法国客户箱单存在空物料号，不能生成出货资料")
        qty = _qty(row.get("quantity"))
        unit_weight = _q2(_dec(row.get("unit_weight")))
        tax_unit, tax_total = _tax_pair(row)
        current = grouped.get(material_no)
        if current is None:
            grouped[material_no] = {
                "material_no": material_no,
                "qty": qty,
                "unit_weight": unit_weight,
                "tax_unit": tax_unit,
                "amount": tax_total,
                "merged": False,
            }
            continue
        if current["unit_weight"] != unit_weight:
            raise ValueError(f"物料 {material_no} 在法国客户箱单中的单重不一致")
        current["qty"] += qty
        current["amount"] = _q2(current["amount"] + tax_total)
        current["merged"] = True
    lines: list[dict[str, Any]] = []
    for index, item in enumerate(grouped.values(), start=1):
        qty = int(item["qty"])
        unit_weight = item["unit_weight"]
        amount = _q2(item["amount"])
        if item["merged"] and qty:
            unit_price = _q2(amount / Decimal(qty))
        else:
            unit_price = item["tax_unit"]
        total_weight = _q2(unit_weight * Decimal(qty))
        lines.append(
            {
                "line_no": str(index),
                "material_no": item["material_no"],
                "qty": str(qty),
                "unit_weight": _fmt2(unit_weight),
                "total_weight": _fmt2(total_weight),
                "tax_unit_price": _fmt2(unit_price),
                "tax_total_amount": _fmt2(amount),
                "qty_num": qty,
                "weight_num": total_weight,
                "amount_num": amount,
            }
        )
    return lines


def _customer_index(rows: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    indexed: dict[str, dict[str, Any]] = {}
    for row in rows:
        material_no = str(row.get("material_no") or "").strip()
        if material_no and material_no not in indexed:
            indexed[material_no] = row
    return indexed


def _validate_boxes(packing_rows: list[dict[str, Any]], customer_by_no: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    prepared: list[dict[str, Any]] = []
    classes_by_box: dict[int, set[str]] = {}
    for row in packing_rows:
        material_no = str(row.get("material_no") or "").strip()
        if not material_no:
            raise ValueError("装箱表存在空物料号，不能生成出货资料")
        if material_no not in customer_by_no:
            raise ValueError(f"装箱物料 {material_no} 不在法国客户箱单中")
        material_class = _class_of_row(row, customer_by_no)
        if not material_class:
            raise ValueError(f"物料 {material_no} 材质无法归类为碳钢或不锈钢")
        box_no = int(_dec(row.get("box_no")))
        if box_no <= 0:
            raise ValueError("装箱表箱号无效，不能生成出货资料")
        classes_by_box.setdefault(box_no, set()).add(material_class)
        customer = customer_by_no[material_no]
        unit_weight = _dec(row.get("unit_weight"))
        if unit_weight == 0:
            unit_weight = _dec(customer.get("unit_weight"))
        prepared.append(
            {
                "box_no": box_no,
                "material_no": material_no,
                "material_class": material_class,
                "qty": _qty(row.get("quantity")),
                "unit_weight": _q2(unit_weight),
                "length": int(_dec(row.get("box_length"))),
                "width": int(_dec(row.get("box_width"))),
                "height": int(_dec(row.get("box_height"))),
                "_seq": len(prepared),
            }
        )
    for box_no, classes in classes_by_box.items():
        if len(classes) > 1:
            raise ValueError(f"箱号 {box_no} 同时有碳钢和不锈钢，请先拆箱")
    prepared.sort(key=lambda row: (row["box_no"], row["_seq"]))
    for row in prepared:
        row.pop("_seq", None)
    numbers = sorted(classes_by_box)
    if numbers != list(range(1, len(numbers) + 1)):
        shown = "、".join(str(number) for number in numbers)
        raise ValueError(f"箱号须从 1 连续编号，当前箱号为 {shown}")
    return prepared


def _box_groups(rows: list[dict[str, Any]]) -> list[list[dict[str, Any]]]:
    groups: list[list[dict[str, Any]]] = []
    current_no: int | None = None
    for row in rows:
        if current_no != row["box_no"]:
            groups.append([])
            current_no = row["box_no"]
        groups[-1].append(row)
    return groups


def _cbm_of_boxes(groups: list[list[dict[str, Any]]]) -> Decimal:
    total = Decimal("0")
    for group in groups:
        first = group[0]
        total += Decimal(first["length"]) * Decimal(first["width"]) * Decimal(first["height"]) / _CBM_DIVISOR
    return _q2(total)


def _packing_items(groups: list[list[dict[str, Any]]]) -> tuple[list[dict[str, Any]], list[int]]:
    items: list[dict[str, Any]] = []
    lengths: list[int] = []
    for group in groups:
        weights = [_q2(row["unit_weight"] * Decimal(row["qty"])) for row in group]
        gross = _q2(sum(weights, Decimal("0")) + BOX_TARE_KG)
        first = group[0]
        for index, row in enumerate(group):
            items.append(
                {
                    "box_no": "" if index else str(row["box_no"]),
                    "material_no": row["material_no"],
                    "qty": str(row["qty"]),
                    "length": "" if index else str(first["length"]),
                    "width": "" if index else str(first["width"]),
                    "height": "" if index else str(first["height"]),
                    "net_weight": _fmt2(weights[index]),
                    "gross_weight": "" if index else _fmt2(gross),
                }
            )
        lengths.append(len(group))
    return items, lengths


def _sum_lines(lines: list[dict[str, Any]]) -> tuple[int, Decimal, Decimal]:
    qty = sum(int(line["qty_num"]) for line in lines)
    weight = _q2(sum((line["weight_num"] for line in lines), Decimal("0")))
    amount = _q2(sum((line["amount_num"] for line in lines), Decimal("0")))
    return qty, weight, amount


def _invoice_mapping(
    order_no: str,
    material_class: str,
    batch_text: str,
    lines: list[dict[str, Any]],
    box_count: int,
    cbm: Decimal,
) -> dict[str, str]:
    qty, weight, amount = _sum_lines(lines)
    gross = _q2(weight + BOX_TARE_KG * Decimal(box_count))
    return {
        "order_no": order_no,
        "invoice_no": invoice_no_for(order_no, material_class),
        "batch_date": batch_text,
        "commodity": COMMODITY[material_class],
        "total_label": TOTAL_LABEL[material_class],
        "qty_sum": str(qty),
        "weight_sum": _fmt2(weight),
        "amount_sum": _fmt2(amount),
        "box_summary": f"{box_count} boxes,{_fmt2(cbm)}CBM",
        "gross_weight": _fmt2(gross),
    }


def _class_stats(
    lines: list[dict[str, Any]],
    groups: list[list[dict[str, Any]]],
) -> dict[str, Any]:
    qty, weight, amount = _sum_lines(lines)
    box_count = len(groups)
    cbm = _cbm_of_boxes(groups)
    gross = _q2(weight + BOX_TARE_KG * Decimal(box_count))
    return {
        "qty": qty,
        "boxes": box_count,
        "net": weight,
        "gross": gross,
        "amount": amount,
        "cbm": cbm,
    }


def plan_shipping_jobs(
    order_no: str,
    transport: str,
    batch_date_value: date,
    customer_rows: list[dict[str, Any]],
    packing_rows: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    if transport not in (TRANSPORT_AIR, TRANSPORT_SEA):
        raise ValueError("请选择空运或海运")
    if not customer_rows:
        raise ValueError("没有法国客户箱单，不能生成出货资料")
    if not packing_rows:
        raise ValueError("装箱表还没有箱号，不能生成出货资料")

    customer_by_no = _customer_index(customer_rows)
    for row in customer_rows:
        material_no = str(row.get("material_no") or "").strip()
        if not material_no:
            raise ValueError("法国客户箱单存在空物料号，不能生成出货资料")
        if material_class_of(str(row.get("material") or "")) is None:
            raise ValueError(f"物料 {material_no} 材质无法归类为碳钢或不锈钢")

    prepared = _validate_boxes(packing_rows, customer_by_no)
    batch_text = format_shipping_date(batch_date_value)
    present = [CLASS_CS, CLASS_SS]
    present = [item for item in present if aggregate_invoice_lines(customer_rows, item)]
    if not present:
        raise ValueError("法国客户箱单没有可归类的碳钢或不锈钢物料")

    jobs: list[dict[str, Any]] = []
    stats: dict[str, dict[str, Any]] = {}
    all_groups = _box_groups(prepared)

    def groups_for(material_class: str) -> list[list[dict[str, Any]]]:
        selected = [row for row in prepared if row["material_class"] == material_class]
        return _box_groups(selected)

    for material_class in present:
        lines = aggregate_invoice_lines(customer_rows, material_class)
        groups = groups_for(material_class)
        stats[material_class] = _class_stats(lines, groups)
        mapping = _invoice_mapping(
            order_no,
            material_class,
            batch_text,
            lines,
            stats[material_class]["boxes"],
            stats[material_class]["cbm"],
        )
        jobs.append(
            {
                "filename": invoice_filename(transport, order_no, material_class),
                "template": "invoice.xlsx",
                "transport": transport,
                "kind": "invoice",
                "mapping": mapping,
                "items": [
                    {key: line[key] for key in (
                        "line_no", "material_no", "qty", "unit_weight", "total_weight",
                        "tax_unit_price", "tax_total_amount",
                    )}
                    for line in lines
                ],
                "group_lengths": [],
            }
        )
        if transport == TRANSPORT_AIR:
            packing_items, group_lengths = _packing_items(groups)
            qty_sum = sum(int(row["qty"]) for group in groups for row in group)
            net_sum = _q2(sum(
                (_q2(row["unit_weight"] * Decimal(row["qty"])) for group in groups for row in group),
                Decimal("0"),
            ))
            gross_sum = _q2(net_sum + BOX_TARE_KG * Decimal(len(groups)))
            jobs.append(
                {
                    "filename": packing_filename(transport, order_no, material_class),
                    "template": "packing_list.xlsx",
                    "transport": transport,
                    "kind": "packing",
                    "mapping": {
                        "order_no": order_no,
                        "invoice_no": invoice_no_for(order_no, material_class),
                        "batch_date": batch_text,
                        "qty_sum": str(qty_sum),
                        "net_sum": _fmt2(net_sum),
                        "gross_sum": _fmt2(gross_sum),
                    },
                    "items": packing_items,
                    "group_lengths": group_lengths,
                }
            )

    if transport == TRANSPORT_SEA:
        packing_items, group_lengths = _packing_items(all_groups)
        qty_sum = sum(row["qty"] for row in prepared)
        net_sum = _q2(sum((_q2(row["unit_weight"] * Decimal(row["qty"])) for row in prepared), Decimal("0")))
        gross_sum = _q2(net_sum + BOX_TARE_KG * Decimal(len(all_groups)))
        invoice_cell = "INVOICE NO.：" + " / ".join(invoice_no_for(order_no, item) for item in present)
        jobs.append(
            {
                "filename": packing_filename(transport, order_no, None),
                "template": "packing_list.xlsx",
                "transport": transport,
                "kind": "packing",
                "mapping": {
                    "order_no": order_no,
                    "invoice_no_cell": invoice_cell,
                    "batch_date": batch_text,
                    "qty_sum": str(qty_sum),
                    "net_sum": _fmt2(net_sum),
                    "gross_sum": _fmt2(gross_sum),
                },
                "items": packing_items,
                "group_lengths": group_lengths,
            }
        )
        empty = {"qty": 0, "boxes": 0, "net": Decimal("0"), "gross": Decimal("0"), "amount": Decimal("0"), "cbm": Decimal("0")}
        cs = stats.get(CLASS_CS, empty)
        ss = stats.get(CLASS_SS, empty)
        jobs.append(
            {
                "filename": f"{order_no}截单数据.xlsx",
                "template": "cutoff.xlsx",
                "transport": transport,
                "kind": "cutoff",
                "mapping": {
                    "order_no": order_no,
                    "box_range": f"1-{len(all_groups)}",
                    "box_count": str(len(all_groups)),
                    "cs_qty": str(cs["qty"]),
                    "cs_boxes": str(cs["boxes"]),
                    "cs_net": _fmt2(cs["net"]),
                    "cs_gross": _fmt2(cs["gross"]),
                    "cs_amount": _fmt2(cs["amount"]),
                    "cs_cbm": _fmt2(cs["cbm"]),
                    "ss_qty": str(ss["qty"]),
                    "ss_boxes": str(ss["boxes"]),
                    "ss_net": _fmt2(ss["net"]),
                    "ss_gross": _fmt2(ss["gross"]),
                    "ss_amount": _fmt2(ss["amount"]),
                    "ss_cbm": _fmt2(ss["cbm"]),
                    "qty_sum": str(cs["qty"] + ss["qty"]),
                    "boxes_sum": str(cs["boxes"] + ss["boxes"]),
                    "net_sum": _fmt2(cs["net"] + ss["net"]),
                    "gross_sum": _fmt2(cs["gross"] + ss["gross"]),
                    "amount_sum": _fmt2(cs["amount"] + ss["amount"]),
                    "cbm_sum": _fmt2(cs["cbm"] + ss["cbm"]),
                },
                "items": [],
                "group_lengths": [],
            }
        )
    return jobs


def _template_path(transport: str, name: str) -> Path:
    mode = "ByAir" if transport == TRANSPORT_AIR else "BySea"
    path = get_templates_dir() / "20-FR-xkz" / "03-shipping" / mode / name
    if not path.is_file():
        raise FileNotFoundError(f"出货资料模板缺失: {path.name}")
    return path


def _blank_leftover_placeholders(ws, row_index: int | None) -> None:
    if not row_index:
        return
    for col in range(1, ws.max_column + 1):
        value = ws.cell(row_index, col).value
        if isinstance(value, str) and "{{" in value:
            ws.cell(row_index, col).value = None


def _merge_commodity(ws, start_row: int, count: int) -> None:
    if count <= 0:
        return
    cell = ws.cell(start_row, 2)
    cell.alignment = Alignment(
        horizontal=cell.alignment.horizontal,
        vertical=cell.alignment.vertical or "center",
        wrap_text=True,
    )
    if count == 1:
        return
    for offset in range(1, count):
        ws.cell(start_row + offset, 2).value = None
    ws.merge_cells(start_row=start_row, start_column=2, end_row=start_row + count - 1, end_column=2)


def _merge_packing_groups(ws, start_row: int, group_lengths: list[int]) -> None:
    row_index = start_row
    for length in group_lengths:
        if length > 1:
            for column in (1, 4, 5, 6, 8):
                for offset in range(1, length):
                    ws.cell(row_index + offset, column).value = None
                ws.merge_cells(
                    start_row=row_index,
                    start_column=column,
                    end_row=row_index + length - 1,
                    end_column=column,
                )
        row_index += length


def _detail_last_column(ws, start_row: int) -> int:
    last = 1
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(start_row, col)
        if cell.value not in (None, ""):
            last = col
            continue
        border = cell.border
        if any(getattr(getattr(border, name, None), "style", None) for name in ("left", "right", "top", "bottom")):
            last = col
    return last


def _border_side(border, name: str) -> Side:
    current = getattr(border, name)
    if current is not None and current.style:
        return Side(style=current.style, color=current.color)
    for other in ("left", "right", "top", "bottom"):
        side = getattr(border, other)
        if side is not None and side.style:
            return Side(style=side.style, color=side.color)
    return Side(style="thin")


def _match_detail_row_styles(ws, start_row: int, count: int) -> None:
    """明细每一格沿用首行字体和格式，并补齐可见边框。合并区只保留外框。"""
    if count <= 0:
        return
    end_row = start_row + count - 1
    last_col = _detail_last_column(ws, start_row)
    merges = [
        rng
        for rng in ws.merged_cells.ranges
        if not (rng.max_row < start_row or rng.min_row > end_row or rng.max_col < 1 or rng.min_col > last_col)
    ]

    def visible_sides(row: int, col: int) -> tuple[str, ...]:
        for rng in merges:
            if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
                sides: list[str] = []
                if row == rng.min_row:
                    sides.append("top")
                if row == rng.max_row:
                    sides.append("bottom")
                if col == rng.min_col:
                    sides.append("left")
                if col == rng.max_col:
                    sides.append("right")
                return tuple(sides)
        return ("left", "right", "top", "bottom")

    prototypes = []
    for col in range(1, last_col + 1):
        source = ws.cell(start_row, col)
        prototypes.append(
            {
                "font": copy(source.font),
                "fill": copy(source.fill),
                "alignment": copy(source.alignment),
                "number_format": source.number_format,
                "protection": copy(source.protection),
                "border": copy(source.border),
            }
        )
    for offset in range(count):
        row = start_row + offset
        for col, source in enumerate(prototypes, start=1):
            dest = ws.cell(row, col)
            dest.font = copy(source["font"])
            dest.fill = copy(source["fill"])
            dest.alignment = copy(source["alignment"])
            dest.number_format = source["number_format"]
            dest.protection = copy(source["protection"])
            wanted = visible_sides(row, col)
            dest.border = Border(
                left=_border_side(source["border"], "left") if "left" in wanted else Side(),
                right=_border_side(source["border"], "right") if "right" in wanted else Side(),
                top=_border_side(source["border"], "top") if "top" in wanted else Side(),
                bottom=_border_side(source["border"], "bottom") if "bottom" in wanted else Side(),
            )


def render_shipping_job(job: dict[str, Any], output_path: Path) -> None:
    template = _template_path(job["transport"], job["template"])
    wb = load_workbook(template)
    ws = wb.active
    _replace_xlsx_scalars(ws, job["mapping"])
    item_row = _find_item_template_row(ws)
    items = job["items"]
    if items:
        if not item_row:
            raise ValueError(f"出货资料模板缺少明细占位行: {template.name}")
        _expand_item_template_rows(ws, item_row, items)
        if job["kind"] == "invoice":
            _merge_commodity(ws, item_row, len(items))
        elif job["kind"] == "packing":
            _merge_packing_groups(ws, item_row, job["group_lengths"])
        _match_detail_row_styles(ws, item_row, len(items))
    else:
        _blank_leftover_placeholders(ws, item_row)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def _fetch_customer_rows(order_no: str) -> list[dict[str, Any]]:
    raw, digits, prefixed = _po_variants(order_no)
    return fetch_all(
        """
        SELECT id, list_id, seq, material_no, material, quantity, unit_weight,
               product_unit_price, tax_unit_price, tax_total_amount
        FROM france_customer_packing_list_items
        WHERE TRIM(po) IN (%s, %s, %s)
        ORDER BY list_id, seq, id
        """,
        (raw, digits, prefixed),
    )


def _fetch_packing_rows(order_no: str) -> list[dict[str, Any]]:
    raw, digits, prefixed = _po_variants(order_no)
    return fetch_all(
        """
        SELECT id, box_no, material_no, material, quantity, unit_weight,
               box_length, box_width, box_height
        FROM packing_details
        WHERE TRIM(order_no) IN (%s, %s, %s)
          AND box_no IS NOT NULL
          AND box_no > 0
        ORDER BY box_no, id
        """,
        (raw, digits, prefixed),
    )


def build_shipping_doc_files(
    *,
    order_no: str,
    transport: str,
    batch_date_value: date,
) -> tuple[Path, list[str]]:
    jobs = plan_shipping_jobs(
        order_no,
        transport,
        batch_date_value,
        _fetch_customer_rows(order_no),
        _fetch_packing_rows(order_no),
    )
    work = Path(tempfile.mkdtemp(prefix="shipping_docs_"))
    produced: list[tuple[str, Path]] = []
    try:
        for job in jobs:
            path = work / job["filename"]
            render_shipping_job(job, path)
            produced.append((job["filename"], path))
        fd, zip_name = tempfile.mkstemp(suffix=".zip")
        os.close(fd)
        zip_path = Path(zip_name)
        with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
            for name, path in produced:
                zf.write(path, arcname=name)
        return zip_path, [name for name, _path in produced]
    finally:
        shutil.rmtree(work, ignore_errors=True)
