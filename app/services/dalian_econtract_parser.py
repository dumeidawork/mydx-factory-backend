"""大连电子合同 Excel 解析（只读，不落盘）。按表头名匹配，不依赖列顺序或工作表名。"""
from __future__ import annotations

import io
import re
from typing import Any

try:
    from openpyxl import load_workbook
except Exception:  # pragma: no cover
    load_workbook = None

NA_VALUES = {"", "#N/A", "N/A", "NA", "NONE", "-"}


def _normalize_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    if text.upper() in NA_VALUES:
        return ""
    return text


def _safe_int(value: Any, default: int | None = 0) -> int | None:
    text = _normalize_str(value)
    if text == "":
        return default
    try:
        return int(float(text))
    except (TypeError, ValueError):
        return default


def _safe_float(value: Any, default: float = 0.0) -> float:
    text = _normalize_str(value)
    if text == "":
        return default
    try:
        return float(text)
    except (TypeError, ValueError):
        return default

HEADER_ALIASES = {
    "料号": "part_no",
    "物料号": "material_no",
    "dn大小": "dn",
    "dn": "dn",
    "订单号": "order_no",
    "单价": "excel_unit_price",
    "数量": "quantity",
    "备注": "remark1",
}

REQUIRED_FIELDS = ("part_no", "material_no", "dn", "order_no", "excel_unit_price", "quantity", "remark1")


def _norm_header(value: Any) -> str:
    text = "" if value is None else str(value)
    text = text.replace("\n", "").replace("\r", "").replace("\t", "")
    text = re.sub(r"\s+", "", text).strip().lower()
    return text


def _round2(value: float) -> float:
    return round(float(value or 0), 2)


def _scan_header_row(ws: Any) -> tuple[int, dict[str, int]] | None:
    max_row = min(int(ws.max_row or 0), 30)
    max_col = int(ws.max_column or 0)
    if max_row <= 0 or max_col <= 0:
        return None
    for row_no in range(1, max_row + 1):
        mapping: dict[str, int] = {}
        for col in range(1, max_col + 1):
            key = _norm_header(ws.cell(row_no, col).value)
            field = HEADER_ALIASES.get(key)
            if field and field not in mapping:
                mapping[field] = col
        if all(name in mapping for name in REQUIRED_FIELDS):
            return row_no, mapping
    return None


def _cell(ws: Any, row_no: int, col: int) -> Any:
    return ws.cell(row_no, col).value


def parse_dalian_econtract_workbook(content: bytes) -> dict[str, Any]:
    if load_workbook is None:
        raise RuntimeError("缺少openpyxl依赖，无法读取Excel。")
    wb = load_workbook(io.BytesIO(content), data_only=True)
    found: tuple[Any, int, dict[str, int]] | None = None
    for name in wb.sheetnames:
        ws = wb[name]
        hit = _scan_header_row(ws)
        if hit:
            found = (ws, hit[0], hit[1])
            break
    if not found:
        raise ValueError("读取失败：未找到表头（需含料号、物料号、DN大小、订单号、单价、数量、备注）。")
    ws, header_row, cols = found
    items: list[dict[str, Any]] = []
    seq = 0
    max_row = int(ws.max_row or 0)
    for row_no in range(header_row + 1, max_row + 1):
        part_no = _normalize_str(_cell(ws, row_no, cols["part_no"]))
        material_no = _normalize_str(_cell(ws, row_no, cols["material_no"]))
        order_no = _normalize_str(_cell(ws, row_no, cols["order_no"]))
        if not material_no and not order_no and not part_no:
            continue
        seq += 1
        quantity = int(_safe_int(_cell(ws, row_no, cols["quantity"]), 0) or 0)
        excel_unit_price = _round2(_safe_float(_cell(ws, row_no, cols["excel_unit_price"]), 0.0))
        items.append(
            {
                "seq": seq,
                "part_no": part_no,
                "material_no": material_no,
                "dn": _normalize_str(_cell(ws, row_no, cols["dn"])),
                "order_no": order_no,
                "excel_unit_price": excel_unit_price,
                "quantity": quantity,
                "remark1": _normalize_str(_cell(ws, row_no, cols["remark1"])),
            }
        )
    if not items:
        raise ValueError("读取失败：未找到有效明细行。")
    order_nos = sorted({_normalize_str(r.get("order_no")) for r in items if _normalize_str(r.get("order_no"))})
    return {
        "items": items,
        "order_nos": order_nos,
    }
