"""解析法国 / 大连价格拆分 Excel（A105 与 304&316）。"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from app.services.price_split_calc import apply_derived, empty_row
from app.services.price_split_schema import MATERIAL_GROUPS

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover
    load_workbook = None

PRICE_MISMATCH_TOLERANCE = 0.05

_HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "drawing_no": ("drawing no and version", "drawing no,", "drawing no", "图纸号"),
    "drawing_code": ("drawing",),
    "drawing_rev": ("version",),
    "part_no": ("ap1",),
    "material_no": ("r3p no.", "r3p no", "r3p"),
    "description": ("description", "描述"),
    "flange_type": ("flange type", "类型", "type"),
    "dn": ("dn size", "dn"),
    "material": ("材质", "material"),
    "steel_price": ("钢材价格", "steel price"),
    "blanking_weight": ("不含热处理", "gross weight of blanking", "下料毛重"),
    "ht_blanking_weight": ("增加热处理）下料", "increase heat treatment gross weight of blanking"),
    "forged_weight": ("锻造后毛重", "gross weight after forging"),
    "ht_loss": ("热处理 火耗", "increase heat treatment"),
    "gas_loss": ("天然气加热火耗", "火耗"),
    "heat_loss": ("火耗", "loss in heat"),
    "net_weight": ("净重", "net weight"),
    "recover_weight": ("回收料", "recover material weight"),
    "scrap_price": ("废料回收价格", "recover material price"),
    "material_cost": ("材料成本", "material cost"),
    "blanking_cost": ("下料费用", "blanking cost"),
    "forging_cost": ("锻造费用", "forging cost"),
    "heat_treatment_cost": ("热处理费用", "heat treatment cost"),
    "ht_freight_cost": ("外协热处理", "heat treatment freight"),
    "ring_rolling_cost": ("碾环", "ring-rolling"),
    "machining_cost": ("机加工费用", "车机加工", "machining cost"),
    "ht_machining_extra": ("热处理机加工增加", "additional cost of heat treatment"),
    "drilling_cost": ("钻孔费用", "钻法兰孔", "drilling cost"),
    "grounding_hole_cost": ("打孔接地",),
    "hoisting_hole_cost": ("吊装孔", "大规格吊装", "through/hoisting", "通孔/吊装"),
    "through_hole_cost": ("通孔费用",),
    "ptfe_hole_cost": ("ptfe钻孔", "ptfe孔", "ptfe hole cost"),
    "m6_hole_cost": ("m6", "侧孔"),
    "packing_cost": ("包装费用", "packing cost"),
    "other_cost": ("其他成本", "other costs"),
    "total_process_cost": ("总加工成本", "total process cost"),
    "profit_rate": ("利润率", "profit rate"),
    "transport_cost": ("运输费用", "transport"),
    "port_surcharge": ("港杂", "port surcharge"),
    "vat_rate": ("增值税", "vat"),
    "net_price": ("最终报价", "最终单价", "final price", "net price", "jan1", "客户"),
    "award_yn": ("award",),
    "sales_mode": ("模式",),
    "year_usage": ("year usage",),
    "has_m16": ("m16",),
    "has_through_hole": ("through hole",),
    "has_ptfe_hole": ("ptfe hole",),
    "scope": ("scope",),
}


def _norm_header(value: Any) -> str:
    text = str(value or "").replace("\n", " ").replace("\r", " ").lower()
    return re.sub(r"\s+", " ", text).strip()


def _to_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _to_float(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    if not text:
        return 0.0
    try:
        return float(text)
    except ValueError:
        return 0.0


def _to_int(value: Any) -> int:
    return int(round(_to_float(value)))


def _split_drawing(raw: str) -> tuple[str, str, str]:
    text = _to_str(raw)
    if not text:
        return "", "", ""
    match = re.match(r"^(?P<code>.+?)\s+(?P<rev>Rev\.?\s*\S+)$", text, re.I)
    if match:
        code = match.group("code").strip()
        rev = match.group("rev").strip()
        return text, code, rev
    return text, text, ""


def _detect_sheet_kind(sheet_name: str) -> str | None:
    name = sheet_name.strip()
    if "价格变化" in name or "说明" in name:
        return None
    lowered = name.lower()
    if "不锈钢" in name:
        return "SS"
    if "304" in name or "316" in name:
        return "SS"
    if "a105" in lowered:
        return "A105"
    return None


def _header_row_for(region: str, group: str) -> int:
    if region == "dalian" and group == "SS":
        return 7
    if region == "dalian":
        return 1
    return 2


def _build_col_map(ws, header_row: int, group: str) -> dict[str, int]:
    """表头 -> 列号。不锈钢身份字段取最后一次匹配（右侧产品栏）。"""
    prefer_last = {"net_price"}
    if group == "SS":
        prefer_last.update(
            {
                "drawing_code",
                "drawing_rev",
                "part_no",
                "material_no",
                "description",
                "flange_type",
                "dn",
                "material",
                "year_usage",
                "scope",
            }
        )
    col_map: dict[str, int] = {}
    max_col = ws.max_column or 1
    for col in range(1, max_col + 1):
        header = _norm_header(ws.cell(header_row, col).value)
        if not header:
            continue
        for field, aliases in _HEADER_ALIASES.items():
            if field in col_map and field not in prefer_last:
                continue
            if not any(alias in header for alias in aliases):
                continue
            if field == "drawing_code" and ("drawing no" in header or "图纸" in header):
                continue
            if field == "blanking_weight" and ("增加热处理" in header or "increase heat treatment" in header):
                continue
            if field == "ht_loss" and "gross weight" in header:
                continue
            if field == "heat_loss" and ("热处理" in header or "increase heat treatment" in header):
                continue
            if field == "gas_loss":
                if group == "SS":
                    continue
                # 「不含热处理」里也含「热处理」，不能据此排除天然气火耗列
                if "热处理 火耗" in header or "increase heat treatment" in header:
                    continue
            if field == "heat_loss" and group == "A105":
                continue
            if field == "material" and (
                "cost" in header or "材料成本" in header or "price" in header or "weight" in header or "recover" in header
            ):
                continue
            if field == "flange_type" and header.strip() in {"type"} and group == "SS" and col > 35:
                col_map["scope"] = col
                continue
            if field == "has_m16" and ("m6" in header or "侧孔" in header):
                continue
            if field == "hoisting_hole_cost" and "通孔费用" in header and "吊装" not in header:
                continue
            if field == "has_through_hole" and "通孔费用" in header:
                continue
            col_map[field] = col
            break
    return col_map


def _cell(ws, row: int, col_map: dict[str, int], field: str) -> Any:
    col = col_map.get(field)
    if not col:
        return None
    return ws.cell(row, col).value


def _parse_sheet(ws, region: str, group: str, source_file: str) -> tuple[list[dict], list[dict], int]:
    header_row = _header_row_for(region, group)
    col_map = _build_col_map(ws, header_row, group)
    rows: list[dict] = []
    warnings: list[dict] = []
    skipped = 0
    seen: set[str] = set()

    for r in range(header_row + 1, (ws.max_row or header_row) + 1):
        left_drawing = _to_str(_cell(ws, r, col_map, "drawing_no"))
        left_r3p = _to_str(_cell(ws, r, col_map, "material_no"))
        right_code = _to_str(_cell(ws, r, col_map, "drawing_code"))
        # 不锈钢：身份优先右侧；A105 无右侧
        if group == "SS":
            material_no = left_r3p
            # 右侧 R3P 与左侧共用 material_no 映射时，prefer_last 已指向右侧
            drawing_full, drawing_code, drawing_rev = _split_drawing(left_drawing)
            if right_code:
                drawing_code = right_code
                drawing_rev = _to_str(_cell(ws, r, col_map, "drawing_rev")) or drawing_rev
                drawing_full = f"{drawing_code} {drawing_rev}".strip()
            part_no = _to_str(_cell(ws, r, col_map, "part_no"))
            material = _to_str(_cell(ws, r, col_map, "material"))
            description = _to_str(_cell(ws, r, col_map, "description"))
            flange_type = _to_str(_cell(ws, r, col_map, "flange_type"))
            dn = _to_float(_cell(ws, r, col_map, "dn"))
        else:
            drawing_full, drawing_code, drawing_rev = _split_drawing(left_drawing)
            material_no = left_r3p
            part_no = _to_str(_cell(ws, r, col_map, "part_no"))
            material = _to_str(_cell(ws, r, col_map, "material")) or ("A105" if group == "A105" else "")
            description = _to_str(_cell(ws, r, col_map, "description"))
            flange_type = _to_str(_cell(ws, r, col_map, "flange_type"))
            dn = _to_float(_cell(ws, r, col_map, "dn"))

        if not material_no and not drawing_full and not right_code:
            skipped += 1
            continue
        if not material_no:
            skipped += 1
            continue

        item = empty_row(group)
        item["drawing_no"] = drawing_full
        item["drawing_code"] = drawing_code
        item["drawing_rev"] = drawing_rev
        item["material_no"] = material_no
        item["part_no"] = part_no
        item["description"] = description
        item["flange_type"] = flange_type
        item["dn"] = dn
        item["material"] = material or ("A105" if group == "A105" else "")
        item["scope"] = _to_str(_cell(ws, r, col_map, "scope"))
        item["award_yn"] = _to_str(_cell(ws, r, col_map, "award_yn"))
        item["sales_mode"] = _to_str(_cell(ws, r, col_map, "sales_mode"))
        item["year_usage"] = _to_int(_cell(ws, r, col_map, "year_usage"))
        item["has_m16"] = _to_str(_cell(ws, r, col_map, "has_m16"))
        item["has_through_hole"] = _to_str(_cell(ws, r, col_map, "has_through_hole"))
        item["has_ptfe_hole"] = _to_str(_cell(ws, r, col_map, "has_ptfe_hole"))

        if group == "SS":
            # 左侧对照件：用未 prefer_last 的列很难再取，按大连表固定列回填
            if region == "dalian":
                left_ap1 = _to_str(ws.cell(r, 3).value)
                left_r3p_raw = _to_str(ws.cell(r, 5).value)
                if left_ap1 and left_ap1 != part_no:
                    item["ref_part_no"] = left_ap1
                if left_r3p_raw and left_r3p_raw != material_no:
                    item["ref_material_no"] = left_r3p_raw
            item["steel_price"] = _to_float(_cell(ws, r, col_map, "steel_price"))
            item["blanking_weight"] = _to_float(_cell(ws, r, col_map, "blanking_weight"))
            item["heat_loss"] = _to_float(_cell(ws, r, col_map, "heat_loss"))
            item["net_weight"] = _to_float(_cell(ws, r, col_map, "net_weight"))
            item["scrap_price"] = _to_float(_cell(ws, r, col_map, "scrap_price"))
        else:
            item["steel_price"] = _to_float(_cell(ws, r, col_map, "steel_price"))
            item["blanking_weight"] = _to_float(_cell(ws, r, col_map, "blanking_weight"))
            item["gas_loss"] = _to_float(_cell(ws, r, col_map, "gas_loss"))
            item["net_weight"] = _to_float(_cell(ws, r, col_map, "net_weight"))
            item["scrap_price"] = _to_float(_cell(ws, r, col_map, "scrap_price"))
            item["heat_treatment_cost"] = _to_float(_cell(ws, r, col_map, "heat_treatment_cost"))
            item["ht_freight_cost"] = _to_float(_cell(ws, r, col_map, "ht_freight_cost"))
            item["ht_machining_extra"] = _to_float(_cell(ws, r, col_map, "ht_machining_extra"))
            item["grounding_hole_cost"] = _to_float(_cell(ws, r, col_map, "grounding_hole_cost"))

        for field in (
            "blanking_cost",
            "forging_cost",
            "ring_rolling_cost",
            "machining_cost",
            "drilling_cost",
            "hoisting_hole_cost",
            "through_hole_cost",
            "ptfe_hole_cost",
            "m6_hole_cost",
            "packing_cost",
            "other_cost",
            "transport_cost",
            "port_surcharge",
            "profit_rate",
            "vat_rate",
        ):
            value = _cell(ws, r, col_map, field)
            if value is not None and value != "":
                item[field] = _to_float(value)
        if not item.get("profit_rate"):
            item["profit_rate"] = 1.07

        excel_net = _to_float(_cell(ws, r, col_map, "net_price"))
        item["excel_net_price"] = excel_net if excel_net else None
        item["source_file"] = source_file
        item["source_sheet"] = ws.title
        item["source_row"] = r

        apply_derived(item, region)

        if item["excel_net_price"] is not None:
            delta = abs(float(item["net_price"]) - float(item["excel_net_price"]))
            if delta > PRICE_MISMATCH_TOLERANCE:
                warnings.append(
                    {
                        "row": r,
                        "material_no": material_no,
                        "message": f"引擎报价 {item['net_price']:.4f} 与 Excel {item['excel_net_price']:.4f} 相差 {delta:.4f}",
                    }
                )

        key = f"{group}|{material_no}"
        if key in seen:
            warnings.append({"row": r, "material_no": material_no, "message": "文件内物料号重复，后出现的行将覆盖"})
        seen.add(key)
        rows.append(item)

    return rows, warnings, skipped


def parse_price_split_workbook(file_path: Path, region: str) -> dict[str, Any]:
    if load_workbook is None:
        raise RuntimeError("缺少 openpyxl 依赖，无法读取 Excel")
    wb = load_workbook(file_path, data_only=True)
    grouped: dict[str, list[dict]] = {g: [] for g in MATERIAL_GROUPS}
    warnings: list[dict] = []
    skipped = 0
    sheets_used: list[str] = []
    try:
        for name in wb.sheetnames:
            kind = _detect_sheet_kind(name)
            if not kind:
                continue
            ws = wb[name]
            sheet_rows, sheet_warnings, sheet_skipped = _parse_sheet(ws, region, kind, file_path.name)
            grouped[kind].extend(sheet_rows)
            warnings.extend({"sheet": name, **w} for w in sheet_warnings)
            skipped += sheet_skipped
            sheets_used.append(name)
    finally:
        wb.close()

    if not sheets_used:
        raise ValueError("未识别到 A105 或 304&316 工作表，请确认上传的是对应区域的价格拆分表")

    return {
        "sheets": sheets_used,
        "rows_a105": grouped["A105"],
        "rows_ss": grouped["SS"],
        "count_a105": len(grouped["A105"]),
        "count_ss": len(grouped["SS"]),
        "skipped": skipped,
        "warnings": warnings[:200],
        "warning_count": len(warnings),
    }
