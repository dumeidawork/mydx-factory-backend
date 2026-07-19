"""质保书占位符：从 xlsx 映射表读取 G/H/I 列规则。"""
from __future__ import annotations

import os
import re
from functools import lru_cache
from pathlib import Path
from typing import Any, Callable

from app.core.paths import get_backend_dir, get_project_root

BACKEND_DIR = get_backend_dir()
PROJECT_ROOT = get_project_root()
MAPPING_CANDIDATES = (
    BACKEND_DIR / "resources/qc_placeholder_mapping.xlsx",
    PROJECT_ROOT / "docs/03-4-3-质保书占位符映射表.xlsx",
    PROJECT_ROOT / "docs/ph1/03-4-3-质保书占位符映射表（待填写）.xlsx",
)
VALID_REGIONS = frozenset({"10-DL-qc", "11-FR-qc"})
SKIP_NOTE_MARKERS = ("暂时不改", "模板中固定")
IGNORE_SAMPLE_MARKERS = (
    "该样本文件中没填写",
    "实际生成时按页面",
    "空值时填入",
)

MODULE1_FIELD_MAP = {
    "证书编号": "certificate_no",
    "客户名称": "customer",
    "合同号": "contract_no",
    "出厂日期": "date",
    "工厂编号": "works_no",
    "交货状态": "state_of_delivery",
    "材料": "material",
    "物料名称": "article",
    "标记要求": "marking",
    "签发地点与日期": "place_date",
}

MODULE2_FIELD_MAP = {
    "数量": "qty",
    "图纸编号": "drawing_no",
    "零件号": "part_no",
    "规格": "specifications",
    "炉号": "raw_material_no",
    "热处理批号": "batch_no",
}

MECH_FIELD_MAP = {
    "炉号": "heat_no",
    "零件号": "part_no",
    "规格": "spec",
    "测试温度": "test_temp",
    "屈服强度": "yield_strength",
    "抗拉强度": "tensile_strength",
    "伸长率": "elongation",
    "断面收缩率": "reduction_area",
    "硬度": "hardness",
    "硬度1": "hardness_1",
    "硬度2": "hardness_2",
    "硬度3": "hardness_3",
    "冲击试验": "impact_test",
    "-20℃冲击试验": "impact_test",
    "硬度1/硬度2/硬度3": "hardness_combo",
}

CHEM_ELEMENT_MAP = {
    "碳 C": "C",
    "锰 Mn": "Mn",
    "磷 P": "P",
    "硫 S": "S",
    "硅 Si": "Si",
    "铜 Cu": "Cu",
    "镍 Ni": "Ni",
    "铬 Cr": "Cr",
    "钼 Mo": "Mo",
    "钒 V": "V",
    "碳当量 CEQ": "CEQ",
}


def _norm(v: Any) -> str:
    if v is None:
        return ""
    return str(v).strip()


def resolve_mapping_xlsx() -> Path:
    env_path = os.getenv("QC_MAPPING_XLSX", "").strip()
    if env_path:
        custom = Path(env_path).expanduser().resolve()
        if custom.is_file():
            return custom
    for path in MAPPING_CANDIDATES:
        if path.is_file():
            return path
    tried = ", ".join(str(p) for p in MAPPING_CANDIDATES)
    raise FileNotFoundError(f"未找到质保书占位符映射表 xlsx，已尝试: {tried}")


def _should_skip_note(note: str) -> bool:
    n = _norm(note)
    return any(m in n for m in SKIP_NOTE_MARKERS)


def _sample_usable(sample: str) -> bool:
    s = _norm(sample)
    if not s or s in ("-", "—", "None"):
        return False
    return not any(m in s for m in IGNORE_SAMPLE_MARKERS)


def _placeholder_key(ph: str) -> str:
    text = _norm(ph)
    m = re.match(r"^\{\{(.+?)\}\}$", text)
    return m.group(1) if m else text.strip("{}")


@lru_cache(maxsize=1)
def _load_directory_same_as_a105() -> dict[tuple[str, str], bool]:
    import openpyxl

    path = resolve_mapping_xlsx()
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["八套模板目录清单"] if "八套模板目录清单" in wb.sheetnames else wb.active
    flags: dict[tuple[str, str], bool] = {}
    for r in ws.iter_rows(values_only=True):
        region, cert = _norm(r[0]), _norm(r[2])
        if region in VALID_REGIONS and cert:
            flags[(region, cert)] = _norm(r[8]).upper() == "Y"
    wb.close()
    return flags


def cert_same_as_a105(region_dir: str, cert_type: str) -> bool:
    return _load_directory_same_as_a105().get((region_dir, cert_type), False)


@lru_cache(maxsize=1)
def _load_all_rules() -> tuple[dict[tuple[str, str], list[dict]], set[str]]:
    import openpyxl

    path = resolve_mapping_xlsx()
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["占位符映射"] if "占位符映射" in wb.sheetnames else wb.active
    rules: dict[tuple[str, str], list[dict]] = {}
    skip_ph: set[str] = set()
    for r in ws.iter_rows(values_only=True):
        if not r or len(r) < 8:
            continue
        region, cert = _norm(r[1]), _norm(r[2])
        ph = _norm(r[6])
        if region not in VALID_REGIONS or not cert or not ph.startswith("{{"):
            continue
        note = _norm(r[9]) if len(r) > 9 else ""
        item = {
            "ph": ph,
            "key": _placeholder_key(ph),
            "sample": _norm(r[7]),
            "fill_source": _norm(r[8]) if len(r) > 8 else "",
            "note": note,
            "skip_template": _should_skip_note(note),
        }
        rules.setdefault((region, cert), []).append(item)
        if item["skip_template"]:
            skip_ph.add(item["key"])
    wb.close()
    return rules, skip_ph


def merge_rules_for_cert(
    region: str,
    cert: str,
    same_as_a105: bool = False,
) -> list[dict]:
    rules, _ = _load_all_rules()
    own = rules.get((region, cert), [])
    base = rules.get((region, "A105"), []) if same_as_a105 and cert != "A105" else []
    by_key: dict[str, dict] = {}
    for item in base + own:
        key = item["key"]
        prev = by_key.get(key)
        if not prev:
            by_key[key] = dict(item)
            continue
        if _sample_usable(item["sample"]) or item["fill_source"]:
            by_key[key] = dict(item)
    return list(by_key.values())


def merge_template_replace_rules(
    region: str,
    cert: str,
    same_as_a105: bool = False,
) -> list[dict]:
    """模板生成用：H→G，跳过模板固定项与无效样例原文。"""
    out = [
        x
        for x in merge_rules_for_cert(region, cert, same_as_a105)
        if not x["skip_template"] and _sample_usable(x["sample"])
    ]
    out.sort(key=lambda x: len(x["sample"]), reverse=True)
    return out


def _parse_mech_row_index(fill_source: str) -> int | None:
    if "第一行" in fill_source:
        return 1
    if "第二行" in fill_source:
        return 2
    return None


def _parse_mech_field(fill_source: str) -> str | None:
    m = re.search(r"【([^】]+)】\s*$", fill_source)
    if not m:
        return None
    return m.group(1).strip()


def _parse_chem_parts(fill_source: str) -> tuple[str | None, str | None]:
    m = re.search(r"【模块四：化学分析】_【([^】]+)】_【([^】]+)】", fill_source)
    if not m:
        return None, None
    element_label, value_type = m.group(1).strip(), m.group(2).strip()
    element = CHEM_ELEMENT_MAP.get(element_label)
    if value_type == "原材化验值":
        return element, "raw"
    if value_type == "自检化验值":
        return element, "self"
    return element, None


def _mech_list_row_values(rows: list[dict] | None, field_name: str) -> str:
    labels = MECH_FIELD_MAP.get(field_name, field_name)
    if isinstance(labels, str):
        names = (labels, field_name)
    else:
        names = (field_name,)
    for row in rows or []:
        f = _norm(row.get("field", ""))
        if f in names or f == field_name:
            return _norm(row.get("value", ""))
    return ""


def _dalian_mech_group(record: dict, row_index: int) -> dict:
    mechanical = record.get("mechanical_tests") or {}
    if not isinstance(mechanical, dict):
        return {}
    if row_index == 1:
        return mechanical.get("min", {}) or {}
    if row_index == 2:
        return mechanical.get("max", {}) or {}
    return {}


def _resolve_fill_source(
    fill_source: str,
    note: str,
    *,
    record: dict,
    region: str,
    cert_type: str,
    cert_text: Callable[[Any], str],
    cert_number: Callable[[Any], str],
    spec_display_fn: Callable[[str], str],
) -> str:
    src = _norm(fill_source)
    note_text = _norm(note)
    base = record.get("base_info") or {}
    delivery_rows = record.get("delivery_content") or []
    delivery = delivery_rows[0] if delivery_rows else {}

    if not src or src in ("空值 无需占位符", "现阶段暂时 为空"):
        return ""
    if "UT" in src and "MT" in src:
        return ""

    if src.startswith("字符＋") and "出厂日期" in src:
        date_val = cert_text(base.get("date", ""))
        return f"ZHANG QIU/{date_val}" if date_val and date_val != "/" else "/"

    m1 = re.match(r"^【模块一：基础与合同信息】_【([^】]+)】$", src)
    if m1:
        field_label = m1.group(1)
        key = MODULE1_FIELD_MAP.get(field_label)
        if not key:
            return ""
        if note_text == "订单号" and field_label in ("合同号", "工厂编号"):
            return cert_text(record.get("order_no", ""))
        return cert_text(base.get(key, ""))

    m2 = re.match(r"^【模块二：交付结果】_【([^】]+)】$", src)
    if m2:
        field_label = m2.group(1)
        key = MODULE2_FIELD_MAP.get(field_label)
        if not key:
            return ""
        if note_text == "物料号" or field_label == "零件号" and "物料号" in note_text:
            return cert_text(record.get("material_no", delivery.get("part_no", "")))
        if key == "specifications":
            if "standard" in note_text or "【standard】" in note_text:
                spec_part = _norm(delivery.get("specifications", ""))
                std_part = _norm(delivery.get("standard", ""))
                combined = f"{spec_part} {std_part}".strip()
                return cert_text(spec_display_fn(combined))
            return cert_text(spec_display_fn(delivery.get("specifications", "")))
        if note_text == "炉号":
            return cert_text(delivery.get("raw_material_no", ""))
        if note_text == "热处理批号":
            return cert_text(delivery.get("batch_no", ""))
        return cert_text(delivery.get(key, ""))

    if "模块三：机械测试" in src:
        field_label = _parse_mech_field(src) or ""
        row_index = _parse_mech_row_index(src)
        is_dalian = region == "10-DL-qc"
        if is_dalian and row_index:
            group = _dalian_mech_group(record, row_index)
            if field_label == "硬度1/硬度2/硬度3":
                parts = [
                    group.get("hardness_1", ""),
                    group.get("hardness_2", ""),
                    group.get("hardness_3", ""),
                ]
                joined = "/".join(_norm(p) for p in parts if _norm(p))
                return cert_text(joined or group.get("hardness", ""))
            key = MECH_FIELD_MAP.get(field_label, "")
            return cert_text(group.get(key, ""))
        rows = record.get("mechanical_tests")
        if isinstance(rows, list):
            if field_label == "硬度1/硬度2/硬度3":
                h1 = _mech_list_row_values(rows, "硬度1")
                h2 = _mech_list_row_values(rows, "硬度2")
                h3 = _mech_list_row_values(rows, "硬度3")
                joined = "/".join(x for x in (h1, h2, h3) if x)
                if not joined:
                    joined = _mech_list_row_values(rows, "硬度")
                return cert_text(joined)
            return cert_text(_mech_list_row_values(rows, field_label))
        return ""

    element, chem_kind = _parse_chem_parts(src)
    if element and chem_kind:
        chem_rows = record.get("chemical_analysis") or []
        chem = chem_rows[0] if chem_rows else {}
        bucket = chem.get(chem_kind, {}) if isinstance(chem, dict) else {}
        if note_text and "空值" in note_text:
            return cert_number(bucket.get(element, ""))
        return cert_number(bucket.get(element, ""))

    return ""


_CHEM_PLACEHOLDER_SUFFIX = {
    "raw_bar": "raw",
    "self_inspection": "self",
}


def _resolve_chem_by_placeholder_key(
    key: str,
    record: dict,
    cert_number: Callable[[Any], str],
) -> str | None:
    """按占位符命名（如 C_raw_bar）从化学分析取值，避免映射表笔误。"""
    for suffix, kind in _CHEM_PLACEHOLDER_SUFFIX.items():
        token = f"_{suffix}"
        if not key.endswith(token):
            continue
        element = key[: -len(token)]
        if not element:
            return None
        chem_rows = record.get("chemical_analysis") or []
        chem = chem_rows[0] if chem_rows else {}
        bucket = chem.get(kind, {}) if isinstance(chem, dict) else {}
        return cert_number(bucket.get(element, ""))
    return None


def build_certificate_mapping(
    record: dict,
    *,
    region: str,
    cert_type: str,
    cert_text: Callable[[Any], str],
    cert_number: Callable[[Any], str],
    spec_display_fn: Callable[[str], str],
    same_as_a105: bool = False,
    extra: dict[str, Any] | None = None,
) -> dict[str, str]:
    """按映射表 I 列（数据来源）从质检记录生成占位符填充字典。"""
    region_dir = "10-DL-qc" if _norm(region).lower() in ("dalian", "10-dl-qc", "大连") else "11-FR-qc"
    if not same_as_a105:
        same_as_a105 = cert_same_as_a105(region_dir, cert_type)
    rules = merge_rules_for_cert(region_dir, cert_type, same_as_a105)
    mapping: dict[str, str] = {}
    for rule in rules:
        if rule["skip_template"]:
            continue
        key = rule["key"]
        if key in mapping:
            continue
        chem_value = _resolve_chem_by_placeholder_key(key, record, cert_number)
        if chem_value is not None:
            value = chem_value
        else:
            value = _resolve_fill_source(
                rule["fill_source"],
                rule["note"],
                record=record,
                region=region_dir,
                cert_type=cert_type,
                cert_text=cert_text,
                cert_number=cert_number,
                spec_display_fn=spec_display_fn,
            )
        mapping[key] = value

    if cert_type == "A105_UTMTPT":
        mapping["ut"] = ""
        mapping["mt"] = ""
        mapping["pt"] = ""

    if extra:
        for k, v in extra.items():
            mapping[k] = cert_text(v)

    # 大连 316 模板使用法国单行机械键名，将 min 行数据同步到无后缀键。
    if region_dir == "10-DL-qc" and cert_type == "316":
        for src, dst in (
            ("mech_yield_1", "mech_yield"),
            ("mech_tensile_1", "mech_tensile"),
            ("mech_elongation_1", "mech_elongation"),
            ("mech_reduction_1", "mech_reduction"),
        ):
            if src in mapping and dst not in mapping:
                mapping[dst] = mapping[src]

    mapping.setdefault("certificate_no", mapping.get("cert_no", cert_text((record.get("base_info") or {}).get("certificate_no", ""))))
    mapping.setdefault("customer", mapping.get("customer_name", cert_text((record.get("base_info") or {}).get("customer", ""))))
    return mapping
