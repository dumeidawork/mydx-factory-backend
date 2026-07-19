"""校验 backend/templates 中各业务模板是否可被正常加载/生成。"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))

from openpyxl import load_workbook

from app.core.paths import get_templates_dir
from app.services.template_renderer import render_docx_template
from qc_test_common import (
    CERT_CONFIGS,
    build_full_record,
    build_render_mapping,
    extract_template_placeholders,
    validate_rendered_docx,
)

OUT_DIR = BACKEND / "storage" / "_template_smoke"


def _check_qc_templates() -> list[str]:
    lines: list[str] = []
    templates_root = get_templates_dir()
    all_ok = True
    for idx, (region_dir, cert, region_api, material, use_utmtpt) in enumerate(CERT_CONFIGS, start=1):
        label = f"{region_dir}/{cert}"
        tpl = templates_root / region_dir / "word" / cert / "quality_certificate.docx"
        if not tpl.exists():
            lines.append(f"[FAIL] 质保书 {label}: 模板不存在")
            all_ok = False
            continue
        try:
            keys = extract_template_placeholders(tpl)
            record = build_full_record(region_dir, cert, region_api, material, seq=idx)
            mapping = build_render_mapping(record, region_api, cert, use_utmtpt=use_utmtpt)
            out = OUT_DIR / f"{region_dir}_{cert}.docx"
            render_docx_template(tpl, out, mapping)
            validation = validate_rendered_docx(out, keys, mapping)
            status = "PASS" if validation["passed"] else "FAIL"
            if not validation["passed"]:
                all_ok = False
            lines.append(
                f"[{status}] 质保书 {label}: 占位符 {len(keys)} 个, "
                f"残留 {len(validation['leftover_placeholders'])}"
            )
        except Exception as exc:
            lines.append(f"[FAIL] 质保书 {label}: {exc}")
            all_ok = False
    lines.append(f"质保书合计: {'全部通过' if all_ok else '存在失败'}")
    return lines


def _check_packing_template() -> list[str]:
    tpl = get_templates_dir() / "装箱单模板.xlsx"
    if not tpl.exists():
        return ["[FAIL] 装箱单: 模板不存在"]
    try:
        wb = load_workbook(tpl)
        ws = wb.active
        OUT_DIR.mkdir(parents=True, exist_ok=True)
        out = OUT_DIR / "packing_smoke.xlsx"
        ws["A1"] = "订单号：TEST-001 装箱单"
        wb.save(out)
        if out.exists() and out.stat().st_size > 0:
            return [f"[PASS] 装箱单: 已加载并写出 ({tpl.name})"]
        return ["[FAIL] 装箱单: 写出文件为空"]
    except Exception as exc:
        return [f"[FAIL] 装箱单: {exc}"]


def _check_dalian_mark_template() -> list[str]:
    from app.api.v1.workflow import _resolve_dalian_shipping_mark_template

    tpl = _resolve_dalian_shipping_mark_template()
    if not tpl.is_file():
        return ["[FAIL] 大连唛头: 模板不存在"]
    try:
        wb = load_workbook(tpl)
        ws = wb.active
        OUT_DIR.mkdir(parents=True, exist_ok=True)
        out = OUT_DIR / "dalian_mark_smoke.xlsx"
        ws["A1"] = "SMOKE TEST"
        wb.save(out)
        if out.exists() and out.stat().st_size > 0:
            return [f"[PASS] 大连唛头: 已加载并写出 ({tpl.name})"]
        return ["[FAIL] 大连唛头: 写出文件为空"]
    except Exception as exc:
        return [f"[FAIL] 大连唛头: {exc}"]


def _check_fallback_qc_template() -> list[str]:
    tpl = get_templates_dir() / "质保书模板.docx"
    if not tpl.exists():
        return ["[SKIP] 兜底质保书模板.docx: 未配置"]
    try:
        keys = extract_template_placeholders(tpl)
        record = build_full_record("11-FR-qc", "A105", "france", "ASTM A105", seq=99)
        mapping = build_render_mapping(record, "france", "A105")
        mapping.setdefault("customer_name", mapping.get("customer", ""))
        out = OUT_DIR / "fallback_qc_smoke.docx"
        render_docx_template(tpl, out, mapping)
        validation = validate_rendered_docx(out, keys, mapping)
        status = "PASS" if validation["passed"] else "FAIL"
        return [f"[{status}] 兜底质保书模板.docx: 占位符 {len(keys)} 个"]
    except Exception as exc:
        return [f"[FAIL] 兜底质保书模板.docx: {exc}"]


def main() -> int:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    lines = ["# 模板冒烟测试", f"templates_dir: {get_templates_dir()}", ""]
    sections = [
        _check_qc_templates(),
        _check_packing_template(),
        _check_dalian_mark_template(),
        _check_fallback_qc_template(),
    ]
    failed = False
    for section in sections:
        lines.extend(section)
        lines.append("")
        if any(line.startswith("[FAIL]") for line in section):
            failed = True

    report = OUT_DIR / "smoke_report.txt"
    text = "\n".join(lines)
    report.write_text(text + "\n", encoding="utf-8")
    print(text)
    print(f"\n报告: {report}")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
