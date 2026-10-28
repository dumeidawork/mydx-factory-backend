"""从装箱汇总表抽出物料-数量/托对照。

  python scripts/extract_france_auto_packing_qty_seed.py
"""
from __future__ import annotations

import sys
from collections import Counter, defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

from openpyxl import load_workbook  # noqa: E402

from app.data.france_auto_packing_qty_seed import SeedRow, write_precomputed  # noqa: E402

EXCEL_CANDIDATES = (
    Path(
        r"c:\Users\dumei\OneDrive\デスクトップ\铭远数字化\00-前期要件"
        r"\铭远数字化\01-前期要件整理\法国流程资料\步骤6.装箱单汇总\装箱汇总表.xlsx"
    ),
    ROOT.parent / "00-temp" / "装箱汇总表.xlsx",
)

EURO_PALLET_BY_DIM = {
    (1200, 800, 200): "欧式200",
    (1200, 800, 400): "欧式400",
}


def _as_int(value: object) -> int | None:
    if value is None or value == "":
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def _norm(value: object) -> str:
    return str(value or "").strip()


def _map_pallet(remark: str, length_mm: int | None, width_mm: int | None, height_mm: int | None) -> str:
    if remark == "欧式托盘" and length_mm is not None and width_mm is not None and height_mm is not None:
        named = EURO_PALLET_BY_DIM.get((length_mm, width_mm, height_mm))
        if named:
            return named
    return remark


def _extract_file(path: Path) -> list[SeedRow]:
    wb = load_workbook(path, data_only=True, read_only=True)
    rows_out: list[SeedRow] = []
    try:
        sheet = wb[wb.sheetnames[0]]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row:
                continue
            values = list(row) + [None] * 7
            material_no = _norm(values[1])
            qty = _as_int(values[2])
            length_mm = _as_int(values[3])
            width_mm = _as_int(values[4])
            height_mm = _as_int(values[5])
            remark = _norm(values[6])
            if not material_no or qty is None or qty <= 0 or not remark:
                continue
            pallet_type = _map_pallet(remark, length_mm, width_mm, height_mm)
            rows_out.append((material_no, qty, pallet_type, length_mm, width_mm, height_mm, ""))
    finally:
        wb.close()
    return rows_out


def _same_payload(left: SeedRow, right: SeedRow) -> bool:
    return left[:6] == right[:6]


def dedupe(rows: list[SeedRow]) -> list[SeedRow]:
    by_key: dict[tuple[str, str], SeedRow] = {}
    extras: dict[tuple[str, str], list[SeedRow]] = defaultdict(list)
    identical = 0
    for row in rows:
        key = (row[0], row[2])
        current = by_key.get(key)
        if current is None:
            by_key[key] = row
            continue
        if _same_payload(current, row):
            identical += 1
            continue
        extras[key].append(row)
    out: list[SeedRow] = []
    for key, row in by_key.items():
        note = row[6]
        extra_rows = extras.get(key) or []
        if extra_rows:
            bits = [
                f"{item[1]}件/托，{item[3]}×{item[4]}×{item[5]}"
                for item in extra_rows
            ]
            extra = "附件另有：" + "；".join(bits)
            note = f"{note} {extra}".strip() if note else extra
        out.append((row[0], row[1], row[2], row[3], row[4], row[5], note))
    print(f"raw={len(rows)} unique={len(out)} identical_dropped={identical} conflicts={len(extras)}")
    return out


def main() -> None:
    source = next((path for path in EXCEL_CANDIDATES if path.is_file()), None)
    if not source:
        raise SystemExit("未找到装箱汇总表.xlsx")
    print("parse", source)
    collected = _extract_file(source)
    rows = tuple(dedupe(collected))
    dest = write_precomputed(rows)
    print("wrote", dest, "count", len(rows))
    print("pallet types", dict(Counter(item[2] for item in rows)))
    noted = sum(1 for item in rows if item[6])
    print(f"带备注：{noted} 条")


if __name__ == "__main__":
    main()
