"""直接调用 generate_dalian_quality_certificate 做冒烟测试。"""
from pathlib import Path

from app.api.v1 import workflow as wf
from app.services.in_memory_store import quality_records
from docx import Document

quality_records[9001] = {
    "id": 9001,
    "order_no": "4511727401",
    "material_no": "FDK-130-618",
    "base_info": {
        "certificate_type": "dalian",
        "certificate_no": "ZYXMZ251227-9",
        "customer": "Siemens Sensors＆Communication Ltd.",
        "contract_no": "4511727401",
        "works_no": "4511727401",
        "date": "2025-12-27",
        "material": "ASTM A105",
        "state_of_delivery": "调质 Quenching +Tempering",
        "article": "锻造法兰 Forged Flange",
        "place_date": "ZHANGQIU/2025-12-27",
    },
    "delivery_content": [
        {
            "qty": 2,
            "drawing_no": "A5E02972437A Rev007",
            "part_no": "FDK-130-618",
            "name": "101166287",
            "specifications": "DN1800 PN10",
            "raw_material_no": "42402125",
        }
    ],
    "mechanical_tests": {
        "min": {
            "heat_no": "42402125",
            "part_no": "P-min",
            "spec": "318",
            "yield_strength": "318",
            "tensile_strength": "493",
            "elongation": "33.5",
            "reduction_area": "69",
            "hardness": "142",
            "impact_test": "T:98,95,88/R:78,65,68",
        },
        "max": {
            "heat_no": "42402125",
            "part_no": "P-max",
            "spec": "493",
            "yield_strength": "320",
            "tensile_strength": "500",
            "elongation": "34",
            "reduction_area": "70",
            "hardness": "145",
            "impact_test": "T:99,95,88/R:78,65,68",
        },
    },
    "chemical_analysis": [
        {
            "raw": {"C": 0.18, "Mn": 0.99, "P": 0.008, "S": 0.002, "Si": 0.22, "Cu": 0.017, "Ni": 0.025, "Cr": 0.021, "Mo": 0.006, "V": 0.002, "CEQ": 0.35},
            "self": {"C": 0.18, "Mn": 0.93, "P": 0.02, "S": 0.017, "Si": 0.22, "Cu": 0.022, "Ni": 0.044, "Cr": 0.041, "Mo": 0.02, "V": 0.003, "CEQ": 0.35},
        }
    ],
}

if __name__ == "__main__":
    tpl = wf._resolve_dalian_certificate_template()
    assert tpl.exists(), tpl
    resp = wf.generate_dalian_quality_certificate(record_id=9001)
    src_path = Path(getattr(resp, "path", ""))
    assert src_path.exists(), src_path
    out = Path(__file__).resolve().parents[1] / "storage" / "test_dalian_api_cert.docx"
    out.write_bytes(src_path.read_bytes())
    doc = Document(str(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            text += "\n" + " | ".join(c.text for c in row.cells)
    assert "ZYXMZ251227-9" in text
    assert "318" in text and "320" in text
    assert "P-min" in text and "P-max" in text
    assert "{{" not in text
    print("generate ok", out, out.stat().st_size)
