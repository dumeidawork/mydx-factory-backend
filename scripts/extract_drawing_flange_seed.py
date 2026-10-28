"""从附件 Excel 抽出图纸号-法兰类型对照。已核对数据请用 deploy_drawing_flange_seed.py --from-db。"""
from __future__ import annotations

import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook  # noqa: E402

from app.utils.drawing_no import drawing_no_base  # noqa: E402

FLANGE_TYPES = ("滑动", "平板", "盲板", "高径")
_FLANGE_ALIASES = {
    "滑动": "滑动",
    "so": "滑动",
    "slip": "滑动",
    "slip-on": "滑动",
    "slip on": "滑动",
    "slipon": "滑动",
    "sop": "滑动",
    "soh": "滑动",
    "平板": "平板",
    "pl": "平板",
    "plate": "平板",
    "flat": "平板",
    "盲板": "盲板",
    "bl": "盲板",
    "blind": "盲板",
    "高径": "高径",
    "高颈": "高径",
    "wn": "高径",
    "lwn": "高径",
    "weld neck": "高径",
    "weldneck": "高径",
    "welding neck": "高径",
}


def normalize_flange_type(raw: str) -> str:
    text = (raw or "").strip()
    if not text:
        return ""
    if text in FLANGE_TYPES:
        return text
    compact = re.sub(r"[\s_\-]+", " ", text).strip().lower().replace("/", " ")
    if compact in _FLANGE_ALIASES:
        return _FLANGE_ALIASES[compact]
    first = compact.split()[0] if compact else ""
    return _FLANGE_ALIASES.get(first, "")

CANDIDATES = [
    Path(r"c:\Users\dumei\OneDrive\デスクトップ\铭远数字化\00-前期要件\大连+法国1-3号更新-1xlsx(3)(1).xlsx"),
    Path(r"c:\Users\dumei\OneDrive\デスクトップ\铭远数字化\00-前期要件"),
    Path(r"c:\00-MYDX\00_mingyuanerp\00-temp"),
]

DRAWING_HEADERS = (
    "图纸号",
    "drawing no and version no",
    "drawing no and version",
    "drawing no,",
    "drawing no",
    "drawing",
)
TYPE_HEADERS = ("法兰类型", "flange type", "类型", "type", "备注")


def _norm_header(value: object) -> str:
    text = str(value or "").replace("\n", " ").replace("\r", " ").lower()
    return re.sub(r"\s+", " ", text).strip().strip(":")


def _pick_col(headers: dict[int, str], aliases: tuple[str, ...]) -> int | None:
    for alias in aliases:
        for col, name in headers.items():
            if name == alias or name.startswith(alias):
                return col
    return None


def _iter_files() -> list[Path]:
    found: list[Path] = []
    for item in CANDIDATES:
        try:
            if item.is_file() and item.suffix.lower() in {".xlsx", ".xls"}:
                found.append(item)
            elif item.is_dir():
                for path in item.rglob("*"):
                    if not path.is_file() or path.suffix.lower() not in {".xlsx", ".xls"}:
                        continue
                    if path.name.startswith("~$"):
                        continue
                    name = path.name.lower()
                    if item.name in {"00-前期要件", "00-temp", "00_mingyuanerp"} or any(
                        token in path.name or token in name
                        for token in ("法兰", "1-3", "price", "大连", "法国", "更新")
                    ):
                        found.append(path)
        except OSError as exc:
            print(f"skip {item}: {exc}")
    unique: list[Path] = []
    seen: set[str] = set()
    for path in found:
        key = str(path).lower()
        if key not in seen:
            seen.add(key)
            unique.append(path)
    return unique


def extract_file(path: Path) -> list[tuple[str, str, str]]:
    wb = load_workbook(path, data_only=True, read_only=True)
    rows_out: list[tuple[str, str, str]] = []
    for sheet in wb.worksheets:
        header_row = None
        headers: dict[int, str] = {}
        for row in sheet.iter_rows(min_row=1, max_row=20, values_only=False):
            current = {cell.column: _norm_header(cell.value) for cell in row if cell.value not in (None, "")}
            if _pick_col(current, DRAWING_HEADERS) and _pick_col(current, TYPE_HEADERS):
                header_row = next((getattr(cell, "row", None) for cell in row if getattr(cell, "row", None)), None)
                headers = current
                break
        if not header_row:
            continue
        drawing_col = _pick_col(headers, DRAWING_HEADERS)
        type_col = _pick_col(headers, TYPE_HEADERS)
        if not drawing_col or not type_col:
            continue
        print(f"  sheet={sheet.title!r} drawing_col={drawing_col} type_col={type_col} header_row={header_row}")
        for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
            if not row:
                continue
            drawing_raw = row[drawing_col - 1] if drawing_col - 1 < len(row) else None
            type_raw = row[type_col - 1] if type_col - 1 < len(row) else None
            drawing_no = drawing_no_base(str(drawing_raw or "").strip())
            flange_type = normalize_flange_type(str(type_raw or "").strip())
            if not drawing_no or not flange_type:
                continue
            if drawing_no.lower() in {"drawing", "drawing no", "图纸号"}:
                continue
            rows_out.append((drawing_no, flange_type, ""))
    wb.close()
    return rows_out


def dedupe(rows: list[tuple[str, str, str]]) -> list[tuple[str, str, str]]:
    by_key: dict[str, tuple[str, str, str]] = {}
    conflicts: dict[str, set[str]] = {}
    for drawing_no, flange_type, remark in rows:
        key = drawing_no.upper()
        if key not in by_key:
            by_key[key] = (drawing_no, flange_type, remark)
            conflicts[key] = {flange_type}
            continue
        conflicts[key].add(flange_type)
    out: list[tuple[str, str, str]] = []
    conflict_count = 0
    for key, (drawing_no, flange_type, remark) in by_key.items():
        types = conflicts[key]
        note = remark
        if len(types) > 1:
            conflict_count += 1
            extra = "、".join(sorted(types))
            note = f"附件中存在多种类型：{extra}" if not note else f"{note}；多种类型：{extra}"
        out.append((drawing_no, flange_type, note))
    out.sort(key=lambda item: item[0].upper())
    print(f"unique={len(out)} conflicts={conflict_count}")
    return out


def write_seed(rows: list[tuple[str, str, str]]) -> Path:
    dest = ROOT / "app" / "data" / "drawing_flange_type_seed_data.py"
    lines = [
        '"""由 extract_drawing_flange_seed.py 生成的预置对照。"""',
        "",
        "PRECOMPUTED_MAPS = (",
    ]
    for drawing_no, flange_type, remark in rows:
        lines.append(f"    ({drawing_no!r}, {flange_type!r}, {remark!r}),")
    lines.append(")")
    lines.append("")
    dest.write_text("\n".join(lines), encoding="utf-8")
    return dest


def main() -> None:
    files = _iter_files()
    print("files:", [str(p) for p in files])
    if not files:
        raise SystemExit("未找到附件 Excel")
    collected: list[tuple[str, str, str]] = []
    for path in files:
        print("parse", path)
        collected.extend(extract_file(path))
    rows = dedupe(collected)
    dest = write_seed(rows)
    print("wrote", dest, "count", len(rows))
    counts: dict[str, int] = {}
    for _, flange_type, _ in rows:
        counts[flange_type] = counts.get(flange_type, 0) + 1
    print("types", counts)


if __name__ == "__main__":
    main()
